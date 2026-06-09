from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
import pandas as pd
import numpy as np
import io, re, json, warnings
import openpyxl
warnings.filterwarnings("ignore")

app = FastAPI(title="Velytics API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Smart Clean Engine ────────────────────────────────────────────────────────
def smart_clean(df_raw: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    fixes = []
    df = df_raw.copy()

    # Empty sheet — nothing to clean (don't crash on index access below).
    if df.shape[0] == 0 or df.shape[1] == 0:
        return df.reset_index(drop=True), ["This sheet is empty — no data to analyse"]

    def header_score(row):
        score = 0
        for val in row:
            v = str(val).strip()
            if v and v.lower() not in ["nan","none",""]:
                if not re.match(r"^[\d,.\-\+\(\)₹$%\s]+$", v):
                    score += 1
        return score

    best_row = 0
    for i in range(1, min(10, len(df))):
        if header_score(df.iloc[i]) > header_score(df.iloc[best_row]):
            best_row = i

    df.columns = df.iloc[best_row].astype(str).str.strip()
    df = df.iloc[best_row + 1:].reset_index(drop=True)

    # Make column names unique & non-empty, so df[col] always returns a Series.
    # (Real files often have blank or repeated headers like "AMOUNT", "AMOUNT", ...)
    new_cols, seen = [], {}
    for i, c in enumerate(df.columns):
        name = str(c).strip()
        if name == "" or name.lower() in ("nan", "none", "unnamed", "0"):
            name = f"Column {i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 0
        new_cols.append(name)
    df.columns = new_cols

    # Drop summary/total rows that pivot & export tools append — they skew every metric.
    if len(df.columns):
        first = df.iloc[:, 0].astype(str).str.strip().str.lower()
        total_mask = first.str.match(r"^(grand\s*total|sub\s*total|total)\b", na=False)
        if total_mask.any():
            df = df[~total_mask]
            fixes.append(f"Removed {int(total_mask.sum())} total/summary row(s)")

    before = len(df)
    df.dropna(how="all", inplace=True)
    df.drop_duplicates(inplace=True)
    removed = before - len(df)
    if removed: fixes.append(f"Removed {removed} blank/duplicate rows")

    for col in df.select_dtypes(include="object").columns:
        try:
            df[col] = df[col].astype(str).str.strip().str.replace(r"\s+", " ", regex=True)
            df[col] = df[col].replace({"nan": np.nan, "none": np.nan, "None": np.nan,
                                        "NULL": np.nan, "N/A": np.nan, "": np.nan})
            nuniq = df[col].nunique()
            if 0 < nuniq < 30:
                df[col] = df[col].str.title()
        except Exception:
            pass

    # Drop spacer / fully-empty columns (common in exported spreadsheets).
    empty_cols = [c for c in df.columns if df[c].isna().all()]
    if empty_cols:
        df = df.drop(columns=empty_cols)
        fixes.append(f"Removed {len(empty_cols)} empty column(s)")

    def _strip_sym(s):
        return s.astype(str).str.replace(r"[₹$€£,.\s%\(\)']", "", regex=True)  # for detection only

    def _to_num_us(s):   # 1,250,000.50  (US/UK)
        return pd.to_numeric(s.astype(str).str.replace(r"[₹$€£\s%\(\)]", "", regex=True).str.replace(",", "", regex=False), errors="coerce")

    def _to_num_eu(s):   # 1.250.000,50  (DE/FR)
        cleaned = s.astype(str).str.replace(r"[₹$€£\s%\(\)]", "", regex=True).str.replace(".", "", regex=False).str.replace(",", ".", regex=False)
        return pd.to_numeric(cleaned, errors="coerce")

    num_converted = 0
    for col in df.columns:
        if df[col].dtype == object:
            sample = df[col].dropna().head(30)
            if len(sample) == 0: continue
            # quick check this column is mostly numeric-ish at all
            if pd.to_numeric(_strip_sym(sample), errors="coerce").notna().mean() < 0.8:
                continue
            # pick the locale (US vs EU) that parses more values correctly
            us_ok = _to_num_us(sample).notna().mean()
            eu_ok = _to_num_eu(sample).notna().mean()
            df[col] = _to_num_eu(df[col]) if eu_ok > us_ok else _to_num_us(df[col])
            num_converted += 1

    if num_converted: fixes.append(f"Converted {num_converted} columns to numeric")

    # We do NOT fabricate numbers. Missing numeric cells are left blank so totals
    # stay honest (sums skip blanks, averages ignore them) — trust over prettiness.
    missing = int(df.select_dtypes(include="number").isnull().sum().sum())
    if missing:
        fixes.append(f"Found {missing} blank numeric cell(s) — left blank, not invented")

    df.reset_index(drop=True, inplace=True)
    return df, fixes


# ── Column Detector ───────────────────────────────────────────────────────────
def detect(df: pd.DataFrame, *keywords: str):
    for kw in keywords:
        col = next((c for c in df.columns if kw.lower() in c.lower()), None)
        if col: return col
    return None


# ── Sales Analysis ────────────────────────────────────────────────────────────
def analyze_sales(df: pd.DataFrame, region: str, category: str, period: str) -> dict:
    date_col    = detect(df, "date")
    rev_col     = detect(df, "total revenue", "revenue")
    region_col  = detect(df, "region")
    product_col = detect(df, "product name", "product")
    person_col  = detect(df, "salesperson", "sales person")
    qty_col     = detect(df, "quantity", "qty")
    target_col  = detect(df, "target")
    cat_col     = detect(df, "category")
    ctype_col   = detect(df, "customer type")
    payment_col = detect(df, "payment")

    # Apply filters
    dff = df.copy()
    if region != "All Regions" and region_col:
        dff = dff[dff[region_col].astype(str).str.title() == region]
    if category != "All Categories" and cat_col:
        dff = dff[dff[cat_col].astype(str).str.title() == category]

    if date_col:
        dff[date_col] = pd.to_datetime(dff[date_col], dayfirst=True, errors="coerce")
        if period == "Last 6 Months" and dff[date_col].notna().any():
            cutoff = dff[date_col].max() - pd.DateOffset(months=6)
            dff = dff[dff[date_col] >= cutoff]
        elif period == "Last 3 Months" and dff[date_col].notna().any():
            cutoff = dff[date_col].max() - pd.DateOffset(months=3)
            dff = dff[dff[date_col] >= cutoff]
        elif period == "This Month" and dff[date_col].notna().any():
            latest = dff[date_col].max()
            dff = dff[(dff[date_col].dt.month == latest.month) & (dff[date_col].dt.year == latest.year)]

    total_rev    = float(pd.to_numeric(dff[rev_col], errors="coerce").sum()) if rev_col else 0
    total_qty    = float(pd.to_numeric(dff[qty_col], errors="coerce").sum()) if qty_col else 0
    total_target = float(pd.to_numeric(dff[target_col], errors="coerce").sum()) if target_col else 0
    achievement  = round(total_rev / total_target * 100, 1) if total_target > 0 else 0
    aov          = round(total_rev / len(dff), 0) if len(dff) > 0 else 0

    # Monthly trend
    monthly_trend = []
    if date_col and rev_col:
        dff2 = dff.copy()
        dff2["_month"] = dff2[date_col].dt.to_period("M")
        mt = dff2.groupby("_month")[rev_col].sum().reset_index()
        mt.columns = ["month", "revenue"]
        monthly_trend = [{"month": str(r["month"]), "revenue": round(float(r["revenue"]), 0)} for _, r in mt.iterrows()]

    # By region
    by_region = []
    if region_col and rev_col:
        rr = dff.groupby(region_col)[rev_col].sum().sort_values(ascending=False).reset_index()
        total_r = rr[rev_col].sum()
        by_region = [{"region": str(r[region_col]), "revenue": round(float(r[rev_col]), 0),
                      "pct": round(float(r[rev_col]) / total_r * 100, 1) if total_r > 0 else 0}
                     for _, r in rr.iterrows()]

    # By category
    by_category = []
    if cat_col and rev_col:
        cc = dff.groupby(cat_col)[rev_col].sum().sort_values(ascending=False).reset_index()
        by_category = [{"category": str(r[cat_col]), "revenue": round(float(r[rev_col]), 0)}
                       for _, r in cc.iterrows()]

    # Top products
    top_products = []
    if product_col and rev_col:
        pp = dff.groupby(product_col)[rev_col].sum().sort_values(ascending=False).head(10).reset_index()
        max_rev = pp[rev_col].max()
        top_products = [{"product": str(r[product_col]),
                         "revenue": round(float(r[rev_col]), 0),
                         "pct": round(float(r[rev_col]) / max_rev * 100, 1) if max_rev > 0 else 0}
                        for _, r in pp.iterrows()]

    # Salesperson performance
    salesperson_perf = []
    if person_col and rev_col:
        sp_rev = dff.groupby(person_col)[rev_col].sum()
        sp_target = dff.groupby(person_col)[target_col].sum() if target_col else None
        sp_df = sp_rev.reset_index()
        sp_df.columns = [person_col, "revenue"]
        if sp_target is not None:
            sp_df = sp_df.merge(sp_target.reset_index().rename(columns={target_col: "target"}), on=person_col, how="left")
            sp_df["achievement"] = (sp_df["revenue"] / sp_df["target"] * 100).round(1)
        else:
            sp_df["achievement"] = 0
        sp_df = sp_df.sort_values("achievement" if "achievement" in sp_df.columns else "revenue", ascending=False)
        salesperson_perf = [{"name": str(r[person_col]),
                              "revenue": round(float(r["revenue"]), 0),
                              "achievement": round(float(r.get("achievement", 0)), 1)}
                            for _, r in sp_df.iterrows()]

    # Payment breakdown
    payment_breakdown = []
    if payment_col and rev_col:
        pm = dff.groupby(payment_col)[rev_col].sum().sort_values(ascending=False).reset_index()
        total_pm = pm[rev_col].sum()
        payment_breakdown = [{"method": str(r[payment_col]),
                               "revenue": round(float(r[rev_col]), 0),
                               "pct": round(float(r[rev_col]) / total_pm * 100, 1) if total_pm > 0 else 0}
                             for _, r in pm.iterrows()]

    # Customer type
    customer_type = []
    if ctype_col and rev_col:
        ct = dff.groupby(ctype_col)[rev_col].sum().sort_values(ascending=False).reset_index()
        total_ct = ct[rev_col].sum()
        customer_type = [{"type": str(r[ctype_col]),
                           "revenue": round(float(r[rev_col]), 0),
                           "pct": round(float(r[rev_col]) / total_ct * 100, 1) if total_ct > 0 else 0}
                        for _, r in ct.iterrows()]

    # Smart alerts
    alerts = []
    if by_region and target_col and rev_col:
        avg_rev = total_rev / max(len(by_region), 1)
        for r in by_region:
            if r["revenue"] < avg_rev * 0.75:
                alerts.append({"type": "Critical", "text": f"{r['region']} region is significantly below average ({_fmt(r['revenue'])} vs avg {_fmt(avg_rev)}). Review strategy."})
    if salesperson_perf:
        low = [s for s in salesperson_perf if s["achievement"] > 0 and s["achievement"] < 75]
        if low:
            alerts.append({"type": "Critical", "text": f"{len(low)} salesperson(s) below 75% of target: {', '.join(s['name'] for s in low[:3])}. Schedule 1:1 coaching."})
        high = [s for s in salesperson_perf if s["achievement"] >= 110]
        if high:
            alerts.append({"type": "Opportunity", "text": f"{high[0]['name']} exceeded target by {high[0]['achievement']-100:.0f}%. Document and replicate this approach."})
    if achievement < 80 and achievement > 0:
        alerts.append({"type": "Warning", "text": f"Overall target achievement is {achievement}% — below the 80% threshold. Review pipeline urgently."})
    if achievement >= 100:
        alerts.append({"type": "Opportunity", "text": f"Team achieved {achievement}% of target. Consider raising Q2 targets by 10-15%."})
    if not alerts:
        alerts.append({"type": "Info", "text": "No critical alerts. Business is performing within normal parameters."})

    # Forecast (simple linear)
    forecast = []
    if monthly_trend and len(monthly_trend) >= 3:
        from sklearn.linear_model import LinearRegression
        vals = [m["revenue"] for m in monthly_trend[-6:]]
        X = np.array(range(len(vals))).reshape(-1, 1)
        y = np.array(vals)
        model = LinearRegression().fit(X, y)
        for i in range(1, 4):
            pred = float(model.predict([[len(vals) + i - 1]])[0])
            forecast.append({"month": f"M+{i}", "revenue": round(max(pred, 0), 0), "forecast": True})

    return {
        "kpis": {
            "total_revenue": round(total_rev, 0),
            "total_revenue_fmt": _fmt(total_rev),
            "target_achievement": achievement,
            "units_sold": int(total_qty),
            "avg_order_value": round(aov, 0),
            "avg_order_value_fmt": _fmt(aov),
            "total_rows": len(dff),
        },
        "monthly_trend": monthly_trend,
        "by_region": by_region,
        "by_category": by_category,
        "top_products": top_products,
        "salesperson_perf": salesperson_perf,
        "payment_breakdown": payment_breakdown,
        "customer_type": customer_type,
        "alerts": alerts,
        "forecast": forecast,
        "columns_detected": {
            "date": date_col, "revenue": rev_col, "region": region_col,
            "product": product_col, "salesperson": person_col,
            "quantity": qty_col, "target": target_col, "category": cat_col,
        }
    }


def _grp(df, group_col, val_col, top=10):
    """Group by col, sum val_col, return top N as list of dicts."""
    if not group_col or not val_col: return []
    g = df.groupby(group_col)[val_col].sum().sort_values(ascending=False).head(top).reset_index()
    mx = g[val_col].max()
    return [{"name": str(r[group_col]), "value": round(float(r[val_col]), 0),
             "pct": round(float(r[val_col]) / mx * 100, 1) if mx > 0 else 0}
            for _, r in g.iterrows()]

def _monthly(df, date_col, val_col):
    if not date_col or not val_col: return []
    d = df.copy()
    d[date_col] = pd.to_datetime(d[date_col], dayfirst=True, errors="coerce")
    d["_m"] = d[date_col].dt.to_period("M")
    m = d.groupby("_m")[val_col].sum().reset_index()
    return [{"month": str(r["_m"]), "value": round(float(r[val_col]), 0)} for _, r in m.iterrows()]

def _alert(cond, kind, text):
    return {"type": kind, "text": text} if cond else None

# ── Inventory ─────────────────────────────────────────────────────────────────
def analyze_inventory(df):
    prod_col    = detect(df, "product name", "product")
    cat_col     = detect(df, "category")
    stock_col   = detect(df, "current stock", "stock")
    reorder_col = detect(df, "reorder level", "reorder")
    cost_col    = detect(df, "unit cost", "cost")
    price_col   = detect(df, "selling price", "price")
    demand_col  = detect(df, "monthly demand", "demand")
    sold90_col  = detect(df, "last 90", "units sold last 90")
    last_sold   = detect(df, "last sold")

    dff = df.copy()
    total_skus = dff[prod_col].nunique() if prod_col else len(dff)

    inv_value = 0
    if stock_col and cost_col:
        dff["_inv_val"] = pd.to_numeric(dff[stock_col], errors="coerce").fillna(0) * pd.to_numeric(dff[cost_col], errors="coerce").fillna(0)
        inv_value = float(dff["_inv_val"].sum())

    oos = int((pd.to_numeric(dff[stock_col], errors="coerce").fillna(0) == 0).sum()) if stock_col else 0
    low_stock = 0
    if stock_col and reorder_col:
        s = pd.to_numeric(dff[stock_col], errors="coerce").fillna(0)
        r = pd.to_numeric(dff[reorder_col], errors="coerce").fillna(0)
        low_stock = int(((s > 0) & (s <= r)).sum())

    by_cat = _grp(dff, cat_col, "_inv_val") if cat_col and "_inv_val" in dff.columns else []
    abc = []
    if prod_col and sold90_col:
        g = dff.groupby(prod_col)[sold90_col].sum().sort_values(ascending=False).reset_index()
        total_s = g[sold90_col].sum()
        g["cum_pct"] = g[sold90_col].cumsum() / total_s * 100 if total_s > 0 else 0
        g["class"] = g["cum_pct"].apply(lambda x: "A" if x <= 80 else ("B" if x <= 95 else "C"))
        abc = [{"product": str(r[prod_col]), "units": int(r[sold90_col]), "class": r["class"]} for _, r in g.head(15).iterrows()]

    alerts = [a for a in [
        _alert(oos > 0, "Critical", f"{oos} products are OUT OF STOCK — immediate reorder required."),
        _alert(low_stock > 0, "Warning", f"{low_stock} products below reorder level — place orders now."),
        _alert(inv_value > 0 and oos / max(total_skus, 1) > 0.2, "Critical", f"Over 20% of SKUs out of stock — supply chain risk."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Inventory health looks good — no critical alerts."})

    return {
        "kpis": {"total_skus": total_skus, "inventory_value": round(inv_value, 0),
                 "inventory_value_fmt": _fmt(inv_value), "out_of_stock": oos, "low_stock": low_stock},
        "by_category": by_cat, "abc_analysis": abc, "alerts": alerts,
        "columns_detected": {"product": prod_col, "category": cat_col, "stock": stock_col, "cost": cost_col}
    }

# ── HR & Payroll ──────────────────────────────────────────────────────────────
def analyze_hr(df):
    name_col    = detect(df, "employee name", "name")
    dept_col    = detect(df, "department", "dept")
    salary_col  = detect(df, "salary", "ctc", "pay")
    perf_col    = detect(df, "performance score", "rating", "score")
    gender_col  = detect(df, "gender")
    ot_col      = detect(df, "overtime")
    leave_col   = detect(df, "leave")
    status_col  = detect(df, "status")

    total_emp  = len(df)
    total_pay  = float(pd.to_numeric(df[salary_col], errors="coerce").sum()) if salary_col else 0
    avg_salary = float(pd.to_numeric(df[salary_col], errors="coerce").mean()) if salary_col else 0
    avg_perf   = float(pd.to_numeric(df[perf_col], errors="coerce").mean()) if perf_col else 0
    ot_risk    = int((pd.to_numeric(df[ot_col], errors="coerce") > 20).sum()) if ot_col else 0

    by_dept    = _grp(df, dept_col, salary_col) if dept_col and salary_col else []
    by_gender  = _grp(df, gender_col, salary_col) if gender_col and salary_col else []

    # Attrition risk
    at_risk = []
    if salary_col and dept_col:
        d = df.copy()
        d["_sal"] = pd.to_numeric(d[salary_col], errors="coerce")
        dept_avg = d.groupby(dept_col)["_sal"].transform("mean")
        d["_risk"] = (d["_sal"] < dept_avg * 0.8).astype(int)
        if ot_col:
            d["_risk"] += (pd.to_numeric(d[ot_col], errors="coerce") > 20).astype(int)
        if perf_col:
            d["_risk"] += (pd.to_numeric(d[perf_col], errors="coerce") < pd.to_numeric(d[perf_col], errors="coerce").quantile(0.25)).astype(int)
        high = d[d["_risk"] >= 2]
        if name_col:
            at_risk = [{"name": str(r[name_col]), "dept": str(r[dept_col]) if dept_col else "-", "salary": round(float(r["_sal"]), 0)} for _, r in high.head(10).iterrows()]

    alerts = [a for a in [
        _alert(ot_risk > 0, "Warning", f"{ot_risk} employees working 20+ overtime hours — burnout risk."),
        _alert(len(at_risk) > 0, "Critical", f"{len(at_risk)} employees at high attrition risk — underpaid or overworked."),
        _alert(avg_perf > 0 and avg_perf < 6, "Warning", f"Average performance score is {avg_perf:.1f}/10 — below healthy threshold."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "HR metrics look healthy — no critical alerts."})

    return {
        "kpis": {"total_employees": total_emp, "total_payroll": round(total_pay, 0),
                 "total_payroll_fmt": _fmt(total_pay), "avg_salary": round(avg_salary, 0),
                 "avg_salary_fmt": _fmt(avg_salary), "avg_performance": round(avg_perf, 1), "overtime_risk": ot_risk},
        "by_department": by_dept, "by_gender": by_gender, "attrition_risk": at_risk, "alerts": alerts,
        "columns_detected": {"name": name_col, "dept": dept_col, "salary": salary_col, "performance": perf_col}
    }

# ── Finance ───────────────────────────────────────────────────────────────────
def analyze_finance(df):
    date_col   = detect(df, "date", "month", "period")
    type_col   = detect(df, "type", "category", "nature")
    amount_col = detect(df, "amount", "value", "actual")
    budget_col = detect(df, "budget", "target", "plan")
    dept_col   = detect(df, "department", "dept", "division")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    total_amount = float(pd.to_numeric(df[amount_col], errors="coerce").sum()) if amount_col else 0
    total_budget = float(pd.to_numeric(df[budget_col], errors="coerce").sum()) if budget_col else 0
    variance = total_amount - total_budget if budget_col else 0
    variance_pct = round(variance / total_budget * 100, 1) if total_budget else 0

    by_type = _grp(df, type_col, amount_col) if type_col and amount_col else []
    by_dept = _grp(df, dept_col, amount_col) if dept_col and amount_col else []
    monthly = _monthly(df, date_col, amount_col) if date_col and amount_col else []

    alerts = [a for a in [
        _alert(variance > 0 and total_budget > 0, "Warning", f"Spending is {_fmt(variance)} over budget ({variance_pct:+.1f}%). Review department costs."),
        _alert(variance < 0 and total_budget > 0, "Opportunity", f"Under budget by {_fmt(abs(variance))} ({abs(variance_pct):.1f}%). Consider reinvesting surplus."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Finance looks healthy — no significant variances detected."})

    return {
        "kpis": {"total_amount": round(total_amount, 0), "total_amount_fmt": _fmt(total_amount),
                 "total_budget": round(total_budget, 0), "budget_variance": round(variance, 0),
                 "budget_variance_pct": variance_pct, "total_records": len(df)},
        "by_type": by_type, "by_department": by_dept, "monthly_trend": monthly, "alerts": alerts,
        "columns_detected": {"date": date_col, "type": type_col, "amount": amount_col, "budget": budget_col}
    }

# ── Retail ────────────────────────────────────────────────────────────────────
def analyze_retail(df):
    date_col    = detect(df, "date")
    prod_col    = detect(df, "product name", "product")
    cat_col     = detect(df, "category")
    rev_col     = detect(df, "revenue")
    qty_col     = detect(df, "quantity", "qty")
    return_col  = detect(df, "returned", "return")
    rating_col  = detect(df, "rating")
    channel_col = detect(df, "channel")
    city_col    = detect(df, "city")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    total_rev   = float(pd.to_numeric(df[rev_col], errors="coerce").sum()) if rev_col else 0
    total_orders= len(df)
    aov         = total_rev / total_orders if total_orders else 0
    return_rate = float((df[return_col].astype(str).str.upper() == "YES").mean() * 100) if return_col else 0
    avg_rating  = float(pd.to_numeric(df[rating_col], errors="coerce").mean()) if rating_col else 0

    by_cat     = _grp(df, cat_col, rev_col) if cat_col and rev_col else []
    by_channel = _grp(df, channel_col, rev_col) if channel_col and rev_col else []
    by_product = _grp(df, prod_col, rev_col) if prod_col and rev_col else []
    by_city    = _grp(df, city_col, rev_col) if city_col and rev_col else []
    monthly    = _monthly(df, date_col, rev_col) if date_col and rev_col else []

    alerts = [a for a in [
        _alert(return_rate > 15, "Critical", f"Return rate is {return_rate:.1f}% — above 15% threshold. Investigate product quality."),
        _alert(5 < return_rate <= 15, "Warning", f"Return rate is {return_rate:.1f}% — monitor closely."),
        _alert(avg_rating > 0 and avg_rating < 3.5, "Warning", f"Average rating is {avg_rating:.1f}/5 — below acceptable threshold."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Retail metrics look healthy."})

    return {
        "kpis": {"total_revenue": round(total_rev, 0), "total_revenue_fmt": _fmt(total_rev),
                 "total_orders": total_orders, "avg_order_value": round(aov, 0),
                 "return_rate": round(return_rate, 1), "avg_rating": round(avg_rating, 1)},
        "by_category": by_cat, "by_channel": by_channel, "top_products": by_product,
        "by_city": by_city, "monthly_trend": monthly, "alerts": alerts,
        "columns_detected": {"date": date_col, "revenue": rev_col, "product": prod_col, "category": cat_col}
    }

# ── Logistics ─────────────────────────────────────────────────────────────────
def analyze_logistics(df):
    date_col    = detect(df, "date")
    cost_col    = detect(df, "freight cost", "cost", "amount")
    route_col   = detect(df, "route")
    driver_col  = detect(df, "driver")
    vehicle_col = detect(df, "vehicle")
    ontime_col  = detect(df, "on time", "delivery status", "ontime")
    delay_col   = detect(df, "delay", "days late")
    weight_col  = detect(df, "weight", "cargo")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    total_cost  = float(pd.to_numeric(df[cost_col], errors="coerce").sum()) if cost_col else 0
    total_shipments = len(df)
    ontime_rate = float((df[ontime_col].astype(str).str.upper() == "YES").mean() * 100) if ontime_col else 0

    by_route   = _grp(df, route_col, cost_col) if route_col and cost_col else []
    by_driver  = _grp(df, driver_col, cost_col) if driver_col and cost_col else []
    by_vehicle = _grp(df, vehicle_col, cost_col) if vehicle_col and cost_col else []
    monthly    = _monthly(df, date_col, cost_col) if date_col and cost_col else []

    alerts = [a for a in [
        _alert(ontime_rate > 0 and ontime_rate < 80, "Critical", f"On-time delivery rate is only {ontime_rate:.1f}% — well below 90% target."),
        _alert(80 <= ontime_rate < 90, "Warning", f"On-time rate is {ontime_rate:.1f}% — slightly below 90% target."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Logistics operations look healthy."})

    return {
        "kpis": {"total_cost": round(total_cost, 0), "total_cost_fmt": _fmt(total_cost),
                 "total_shipments": total_shipments, "ontime_rate": round(ontime_rate, 1)},
        "by_route": by_route, "by_driver": by_driver, "by_vehicle": by_vehicle,
        "monthly_trend": monthly, "alerts": alerts,
        "columns_detected": {"date": date_col, "cost": cost_col, "route": route_col, "driver": driver_col}
    }

# ── Restaurant ────────────────────────────────────────────────────────────────
def analyze_restaurant(df):
    date_col   = detect(df, "date")
    item_col   = detect(df, "menu", "item", "dish")
    rev_col    = detect(df, "revenue", "sales", "amount")
    cost_col   = detect(df, "cost", "food cost")
    waste_col  = detect(df, "wastage", "waste")
    rating_col = detect(df, "rating")
    shift_col  = detect(df, "shift")
    qty_col    = detect(df, "quantity", "orders", "covers")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    total_rev    = float(pd.to_numeric(df[rev_col], errors="coerce").sum()) if rev_col else 0
    total_orders = int(pd.to_numeric(df[qty_col], errors="coerce").sum()) if qty_col else len(df)
    avg_rating   = float(pd.to_numeric(df[rating_col], errors="coerce").mean()) if rating_col else 0
    total_waste  = float(pd.to_numeric(df[waste_col], errors="coerce").sum()) if waste_col else 0
    food_cost_pct = 0
    if cost_col and rev_col:
        tc = pd.to_numeric(df[cost_col], errors="coerce").sum()
        tr = pd.to_numeric(df[rev_col], errors="coerce").sum()
        food_cost_pct = round(tc / tr * 100, 1) if tr > 0 else 0

    by_item  = _grp(df, item_col, rev_col) if item_col and rev_col else []
    by_shift = _grp(df, shift_col, rev_col) if shift_col and rev_col else []
    monthly  = _monthly(df, date_col, rev_col) if date_col and rev_col else []

    alerts = [a for a in [
        _alert(food_cost_pct > 45, "Critical", f"Food cost is {food_cost_pct}% of revenue — above 45% danger threshold."),
        _alert(35 < food_cost_pct <= 45, "Warning", f"Food cost is {food_cost_pct}% — target is below 35%."),
        _alert(avg_rating > 0 and avg_rating < 3.5, "Warning", f"Average customer rating is {avg_rating:.1f}/5 — investigate service quality."),
        _alert(total_waste > 0, "Warning", f"Total wastage is {total_waste:.0f} units — review portion control and ordering."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Restaurant operations look healthy."})

    return {
        "kpis": {"total_revenue": round(total_rev, 0), "total_revenue_fmt": _fmt(total_rev),
                 "total_orders": total_orders, "avg_rating": round(avg_rating, 1),
                 "food_cost_pct": food_cost_pct, "total_waste": round(total_waste, 0)},
        "by_item": by_item, "by_shift": by_shift, "monthly_trend": monthly, "alerts": alerts,
        "columns_detected": {"date": date_col, "revenue": rev_col, "item": item_col, "rating": rating_col}
    }

# ── Healthcare ────────────────────────────────────────────────────────────────
def analyze_healthcare(df):
    date_col    = detect(df, "date")
    doc_col     = detect(df, "doctor", "physician")
    dept_col    = detect(df, "department", "specialty")
    rev_col     = detect(df, "revenue", "billing", "amount")
    diag_col    = detect(df, "diagnosis", "condition")
    readmit_col = detect(df, "readmit")
    rating_col  = detect(df, "rating", "satisfaction")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    total_rev     = float(pd.to_numeric(df[rev_col], errors="coerce").sum()) if rev_col else 0
    total_patients= len(df)
    readmit_rate  = float((df[readmit_col].astype(str).str.upper() == "YES").mean() * 100) if readmit_col else 0
    avg_rating    = float(pd.to_numeric(df[rating_col], errors="coerce").mean()) if rating_col else 0

    by_dept = _grp(df, dept_col, rev_col) if dept_col and rev_col else []
    by_doc  = _grp(df, doc_col, rev_col) if doc_col and rev_col else []
    by_diag = _grp(df, diag_col, rev_col) if diag_col and rev_col else []
    monthly = _monthly(df, date_col, rev_col) if date_col and rev_col else []

    alerts = [a for a in [
        _alert(readmit_rate > 10, "Critical", f"Readmission rate is {readmit_rate:.1f}% — above 10% threshold. Review discharge protocols."),
        _alert(avg_rating > 0 and avg_rating < 3.5, "Warning", f"Patient satisfaction is {avg_rating:.1f}/5 — needs improvement."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Healthcare metrics look healthy."})

    return {
        "kpis": {"total_revenue": round(total_rev, 0), "total_revenue_fmt": _fmt(total_rev),
                 "total_patients": total_patients, "readmission_rate": round(readmit_rate, 1), "avg_rating": round(avg_rating, 1)},
        "by_department": by_dept, "by_doctor": by_doc, "by_diagnosis": by_diag,
        "monthly_trend": monthly, "alerts": alerts,
        "columns_detected": {"date": date_col, "revenue": rev_col, "department": dept_col, "doctor": doc_col}
    }

# ── Manufacturing ─────────────────────────────────────────────────────────────
def analyze_manufacturing(df):
    date_col     = detect(df, "date")
    machine_col  = detect(df, "machine")
    planned_col  = detect(df, "planned output", "planned")
    actual_col   = detect(df, "actual output", "actual")
    defect_col   = detect(df, "defective", "defect")
    downtime_col = detect(df, "downtime")
    dcost_col    = detect(df, "downtime cost")
    shift_col    = detect(df, "shift")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    actual_total  = float(pd.to_numeric(df[actual_col], errors="coerce").sum()) if actual_col else 0
    planned_total = float(pd.to_numeric(df[planned_col], errors="coerce").sum()) if planned_col else 0
    efficiency    = round(actual_total / planned_total * 100, 1) if planned_total > 0 else 0
    defect_total  = float(pd.to_numeric(df[defect_col], errors="coerce").sum()) if defect_col else 0
    defect_rate   = round(defect_total / actual_total * 100, 2) if actual_total > 0 else 0
    downtime      = float(pd.to_numeric(df[downtime_col], errors="coerce").sum()) if downtime_col else 0
    downtime_cost = float(pd.to_numeric(df[dcost_col], errors="coerce").sum()) if dcost_col else 0

    by_machine = _grp(df, machine_col, downtime_col) if machine_col and downtime_col else []
    by_shift   = _grp(df, shift_col, actual_col) if shift_col and actual_col else []
    monthly    = _monthly(df, date_col, actual_col) if date_col and actual_col else []

    alerts = [a for a in [
        _alert(efficiency > 0 and efficiency < 80, "Critical", f"Production efficiency is {efficiency}% — below 80% threshold."),
        _alert(defect_rate > 2, "Warning", f"Defect rate is {defect_rate}% — above 2% acceptable limit."),
        _alert(downtime_cost > 0, "Warning", f"Total downtime cost is {_fmt(downtime_cost)} — review maintenance schedule."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Manufacturing operations look healthy."})

    return {
        "kpis": {"actual_output": round(actual_total, 0), "planned_output": round(planned_total, 0),
                 "efficiency_pct": efficiency, "defect_rate": defect_rate,
                 "total_downtime_hrs": round(downtime, 1), "downtime_cost": round(downtime_cost, 0),
                 "downtime_cost_fmt": _fmt(downtime_cost)},
        "by_machine": by_machine, "by_shift": by_shift, "monthly_trend": monthly, "alerts": alerts,
        "columns_detected": {"date": date_col, "machine": machine_col, "actual": actual_col, "planned": planned_col}
    }

# ── Marketing ─────────────────────────────────────────────────────────────────
def analyze_marketing(df):
    date_col    = detect(df, "date")
    channel_col = detect(df, "channel")
    spend_col   = detect(df, "spend", "cost", "budget")
    revenue_col = detect(df, "revenue generated", "revenue")
    leads_col   = detect(df, "leads")
    conv_col    = detect(df, "conversions")
    cac_col     = detect(df, "cac")
    roas_col    = detect(df, "roas")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    total_spend   = float(pd.to_numeric(df[spend_col], errors="coerce").sum()) if spend_col else 0
    total_revenue = float(pd.to_numeric(df[revenue_col], errors="coerce").sum()) if revenue_col else 0
    total_leads   = float(pd.to_numeric(df[leads_col], errors="coerce").sum()) if leads_col else 0
    total_conv    = float(pd.to_numeric(df[conv_col], errors="coerce").sum()) if conv_col else 0
    roas          = total_revenue / total_spend if total_spend > 0 else 0
    conv_rate     = round(total_conv / total_leads * 100, 1) if total_leads > 0 else 0

    by_channel = _grp(df, channel_col, revenue_col) if channel_col and revenue_col else []
    by_spend   = _grp(df, channel_col, spend_col) if channel_col and spend_col else []
    monthly    = _monthly(df, date_col, revenue_col) if date_col and revenue_col else []

    alerts = [a for a in [
        _alert(roas > 0 and roas < 2, "Critical", f"ROAS is {roas:.1f}x — below 2x minimum. Review channel spend allocation."),
        _alert(conv_rate > 0 and conv_rate < 5, "Warning", f"Conversion rate is {conv_rate}% — below 5% benchmark."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Marketing performance looks healthy."})

    return {
        "kpis": {"total_spend": round(total_spend, 0), "total_spend_fmt": _fmt(total_spend),
                 "total_revenue": round(total_revenue, 0), "total_revenue_fmt": _fmt(total_revenue),
                 "roas": round(roas, 2), "total_leads": int(total_leads),
                 "conversion_rate": conv_rate},
        "by_channel": by_channel, "spend_by_channel": by_spend,
        "monthly_trend": monthly, "alerts": alerts,
        "columns_detected": {"date": date_col, "channel": channel_col, "spend": spend_col, "revenue": revenue_col}
    }

# ── Education ─────────────────────────────────────────────────────────────────
def analyze_education(df):
    date_col    = detect(df, "date")
    class_col   = detect(df, "class", "grade", "standard")
    subject_col = detect(df, "subject")
    score_col   = detect(df, "avg score", "score", "marks")
    attend_col  = detect(df, "present", "attendance")
    total_col   = detect(df, "total students", "students")
    fee_col     = detect(df, "fee charged", "fee paid", "fee")
    pass_col    = detect(df, "pass")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    avg_score    = float(pd.to_numeric(df[score_col], errors="coerce").mean()) if score_col else 0
    total_students = int(pd.to_numeric(df[total_col], errors="coerce").sum()) if total_col else len(df)
    total_fees   = float(pd.to_numeric(df[fee_col], errors="coerce").sum()) if fee_col else 0
    pass_rate    = float((df[pass_col].astype(str).str.upper() == "YES").mean() * 100) if pass_col else 0
    attend_rate  = float(pd.to_numeric(df[attend_col], errors="coerce").mean() / pd.to_numeric(df[total_col], errors="coerce").mean() * 100) if attend_col and total_col else 0

    by_class   = _grp(df, class_col, score_col) if class_col and score_col else []
    by_subject = _grp(df, subject_col, score_col) if subject_col and score_col else []

    alerts = [a for a in [
        _alert(avg_score > 0 and avg_score < 50, "Critical", f"Average score is {avg_score:.1f} — below 50% passing threshold."),
        _alert(attend_rate > 0 and attend_rate < 75, "Warning", f"Attendance rate is {attend_rate:.1f}% — below 75% minimum."),
        _alert(pass_rate > 0 and pass_rate < 80, "Warning", f"Pass rate is {pass_rate:.1f}% — below 80% target."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Academic performance looks healthy."})

    return {
        "kpis": {"avg_score": round(avg_score, 1), "total_students": total_students,
                 "total_fees": round(total_fees, 0), "total_fees_fmt": _fmt(total_fees),
                 "pass_rate": round(pass_rate, 1), "attendance_rate": round(attend_rate, 1)},
        "by_class": by_class, "by_subject": by_subject, "alerts": alerts,
        "columns_detected": {"date": date_col, "class": class_col, "subject": subject_col, "score": score_col}
    }

# ── Banking ───────────────────────────────────────────────────────────────────
def analyze_banking(df):
    date_col    = detect(df, "date")
    type_col    = detect(df, "loan type", "type")
    amount_col  = detect(df, "loan amount", "amount")
    emi_col     = detect(df, "monthly emi", "emi")
    status_col  = detect(df, "status")
    overdue_col = detect(df, "days overdue", "overdue")
    out_col     = detect(df, "outstanding amount", "outstanding")
    branch_col  = detect(df, "branch")
    agent_col   = detect(df, "agent")

    if date_col: df[date_col] = pd.to_datetime(df[date_col], dayfirst=True, errors="coerce")

    total_disbursed = float(pd.to_numeric(df[amount_col], errors="coerce").sum()) if amount_col else 0
    total_outstanding = float(pd.to_numeric(df[out_col], errors="coerce").sum()) if out_col else 0
    npa_count = int((df[status_col].astype(str).str.upper() == "NPA").sum()) if status_col else 0
    npa_rate  = round(npa_count / max(len(df), 1) * 100, 1)
    npa_value = float(df[df[status_col].astype(str).str.upper() == "NPA"][out_col].apply(pd.to_numeric, errors="coerce").sum()) if status_col and out_col else 0

    by_type   = _grp(df, type_col, amount_col) if type_col and amount_col else []
    by_branch = _grp(df, branch_col, amount_col) if branch_col and amount_col else []
    by_agent  = _grp(df, agent_col, amount_col) if agent_col and amount_col else []

    alerts = [a for a in [
        _alert(npa_rate > 10, "Critical", f"NPA rate is {npa_rate}% — critically high. Immediate recovery action required."),
        _alert(5 < npa_rate <= 10, "Warning", f"NPA rate is {npa_rate}% — above 5% threshold. Strengthen credit assessment."),
        _alert(npa_value > total_disbursed * 0.1, "Critical", f"NPA value {_fmt(npa_value)} exceeds 10% of portfolio — systemic risk."),
    ] if a]
    if not alerts: alerts.append({"type": "Info", "text": "Loan portfolio looks healthy — no critical alerts."})

    return {
        "kpis": {"total_disbursed": round(total_disbursed, 0), "total_disbursed_fmt": _fmt(total_disbursed),
                 "total_outstanding": round(total_outstanding, 0), "total_outstanding_fmt": _fmt(total_outstanding),
                 "npa_count": npa_count, "npa_rate": npa_rate, "npa_value": round(npa_value, 0),
                 "npa_value_fmt": _fmt(npa_value), "total_loans": len(df)},
        "by_type": by_type, "by_branch": by_branch, "by_agent": by_agent, "alerts": alerts,
        "columns_detected": {"date": date_col, "amount": amount_col, "status": status_col, "branch": branch_col}
    }

# ── Generic fallback ──────────────────────────────────────────────────────────
def analyze_generic(df):
    num_cols = df.select_dtypes(include="number").columns.tolist()
    summaries = []
    for col in num_cols[:6]:
        summaries.append({"column": col, "sum": round(float(df[col].sum()), 0),
                          "mean": round(float(df[col].mean()), 2),
                          "min": round(float(df[col].min()), 2),
                          "max": round(float(df[col].max()), 2)})
    return {
        "kpis": {"total_rows": len(df), "total_columns": len(df.columns), "numeric_columns": len(num_cols)},
        "column_summaries": summaries,
        "columns": list(df.columns),
        "alerts": [{"type": "Info", "text": "Generic analysis complete. Select a specific module for deeper insights."}]
    }


# Detected per request (set in /analyze). Western default; sync analysis = safe.
_CURRENCY = "$"

def detect_currency(df_raw: pd.DataFrame) -> str:
    """Sniff the currency from raw cells — symbol first, then ISO code / word in
    headers (e.g. 'Revenue (INR)'); default to $ (Western)."""
    try:
        text = " ".join(df_raw.head(200).astype(str).values.ravel().tolist())
        counts = {s: text.count(s) for s in ["€", "£", "₹", "$"]}
        best = max(counts, key=counts.get)
        if counts[best] > 0:
            return best
        up = text.upper()
        codes = [("₹", ["INR", "RUPEE"]), ("€", ["EUR", "EURO"]),
                 ("£", ["GBP", "STERLING", "POUND"]), ("$", ["USD", "DOLLAR"])]
        cc = {sym: sum(up.count(w) for w in words) for sym, words in codes}
        bestc = max(cc, key=cc.get)
        return bestc if cc[bestc] > 0 else "$"
    except Exception:
        return "$"

def _fmt(v: float) -> str:
    """Universal money format: $1.23B / €4.5M / £12K (locale-neutral, K/M/B)."""
    if v is None:
        return "—"
    c = _CURRENCY
    n = abs(float(v))
    if n >= 1e9: return f"{c}{v/1e9:.2f}B"
    if n >= 1e6: return f"{c}{v/1e6:.2f}M"
    if n >= 1e3: return f"{c}{v/1e3:.1f}K"
    return f"{c}{v:,.0f}"


# ════════════════════════════════════════════════════════════════════════════
# UNIVERSAL ADAPTIVE ENGINE — works on ANY sheet/module, driven by the data
# itself (not hardcoded column names). Profiles columns → builds dynamic filters
# → auto-generates KPIs, breakdowns & trends from whatever columns exist.
# ════════════════════════════════════════════════════════════════════════════

_ID_HINT   = re.compile(r"(\bid\b|\bno\.?\b|number|code|sku|ref|phone|mobile|pin|zip|barcode|serial|account)", re.I)
_MONEY_HINT = re.compile(r"(amount|amt|price|cost|revenue|sales|value|spend|budget|salary|payroll|fee|total|income|profit|tax|turnover|€|£|\$|₹|usd|eur|gbp|inr|umsatz|betrag|chiffre|montant|ventas|importe|precio|preis|prix)", re.I)


def _is_money(name: str) -> bool:
    return bool(_MONEY_HINT.search(str(name)))


def _to_datetime(s: pd.Series):
    """Parse a series to datetime if it plausibly is one, else None. Tries both
    month-first (US) and day-first (EU/India) and keeps whichever parses more —
    so 24/09/2024 is understood, not silently dropped."""
    if pd.api.types.is_datetime64_any_dtype(s):
        return s
    nonnull = s.dropna()
    if len(nonnull) == 0 or pd.api.types.is_numeric_dtype(s):
        return None
    sample = nonnull.astype(str).head(50)
    r_mf = pd.to_datetime(sample, errors="coerce").notna().mean()              # month-first
    r_df = pd.to_datetime(sample, errors="coerce", dayfirst=True).notna().mean()  # day-first
    if max(r_mf, r_df) < 0.8:
        return None
    return pd.to_datetime(s, errors="coerce", dayfirst=(r_df > r_mf))


def profile_columns(df: pd.DataFrame) -> list[dict]:
    """Classify every column by ROLE (date / number / category / id / text) from
    its actual values — so we adapt to any sheet instead of guessing by name."""
    n = max(len(df), 1)
    prof = []
    for col in df.columns:
        s = df[col]
        nonnull = s.dropna()
        distinct = int(nonnull.nunique())
        missing = int(s.isna().sum())
        dt = _to_datetime(s)
        if dt is not None and dt.notna().sum() > 0:
            role = "date"
        elif pd.api.types.is_numeric_dtype(s):
            strong_id = re.search(r"(\bid\b|\bhsn\b|\bpin\b|\bzip\b|\bsku\b|\bref\b|account|\bcode\b)", str(col).lower())
            role = "id" if (strong_id or (_ID_HINT.search(str(col)) and distinct > 0.9 * len(nonnull))) else "number"
        else:
            if _ID_HINT.search(str(col)) and distinct > 0.7 * len(nonnull):
                role = "id"
            elif 1 < distinct <= max(50, int(n * 0.5)):
                role = "category"
            else:
                role = "text"
        prof.append({"column": col, "role": role, "distinct": distinct, "missing": missing})
    return prof


def build_filter_schema(df: pd.DataFrame, profile: list[dict]) -> list[dict]:
    """Turn the profile into UI-ready filter controls — one per useful column."""
    schema = []
    for p in profile:
        col, role = p["column"], p["role"]
        try:
            if role == "category" or (role == "id" and p["distinct"] <= 100):
                vals = [str(v) for v in df[col].dropna().unique().tolist()][:200]
                schema.append({"column": col, "type": "category", "values": sorted(vals)})
            elif role == "number":
                s = pd.to_numeric(df[col], errors="coerce").dropna()
                if len(s):
                    schema.append({"column": col, "type": "number",
                                   "min": float(s.min()), "max": float(s.max())})
            elif role == "date":
                d = _to_datetime(df[col]).dropna()
                if len(d):
                    schema.append({"column": col, "type": "date",
                                   "min": str(d.min().date()), "max": str(d.max().date())})
        except Exception:
            continue
    return schema


def apply_filters(df: pd.DataFrame, filters: dict) -> pd.DataFrame:
    """Apply user-chosen filters (any column) to the cleaned dataframe."""
    if not filters:
        return df
    out = df
    for col, crit in filters.items():
        if col not in out.columns or not isinstance(crit, dict):
            continue
        try:
            if "values" in crit and crit["values"]:
                wanted = {str(v) for v in crit["values"]}
                out = out[out[col].astype(str).isin(wanted)]
            else:
                if crit.get("min") is not None or crit.get("max") is not None:
                    if _to_datetime(out[col]) is not None and not pd.api.types.is_numeric_dtype(out[col]):
                        s = _to_datetime(out[col])
                        if crit.get("min"): out = out[s >= pd.to_datetime(crit["min"])]
                        if crit.get("max"): out = out[s <= pd.to_datetime(crit["max"])]
                    else:
                        s = pd.to_numeric(out[col], errors="coerce")
                        if crit.get("min") is not None: out = out[s >= float(crit["min"])]
                        if crit.get("max") is not None: out = out[s <= float(crit["max"])]
        except Exception:
            continue
    return out


def _groupable(df: pd.DataFrame, profile: list[dict]) -> list[str]:
    """Columns worth grouping a measure by — text/category/id with real grouping
    (more than 1 value, not all-unique, not mostly blank). Lowest-cardinality first
    so the cleanest charts surface, but high-card names still allowed for top-N."""
    n = max(len(df), 1)
    cols = [p["column"] for p in profile
            if p["role"] in ("category", "text", "id")
            and 1 < p["distinct"] < len(df)
            and p["missing"] < 0.6 * n]
    dist = {p["column"]: p["distinct"] for p in profile}
    return sorted(cols, key=lambda c: dist[c])


def analyze_auto(df: pd.DataFrame, profile: list[dict],
                 group_by: str = "", measure: str = "") -> dict:
    """Data-driven analysis that works for ANY sheet. Builds KPIs from numeric
    columns, breakdowns of the primary measure by each category, and a time trend
    if there's a date column. If group_by/measure are given, builds that view too."""
    measures   = [p["column"] for p in profile if p["role"] == "number"]
    dimensions = [p["column"] for p in profile if p["role"] == "category"]
    date_cols  = [p["column"] for p in profile if p["role"] == "date"]
    dist = {p["column"]: p["distinct"] for p in profile}
    miss = {p["column"]: p["missing"] for p in profile}

    def _measure_rank(name: str) -> int:
        n = str(name).lower()
        # 2 = never sum (per-unit, rates, durations, scores — adding them is meaningless)
        if re.search(r"(per\b|/|\brate\b|ratio|percent|%|unit\s*price|\bprice\b|avg|average|mean|tenure|\bage\b|score|rating|index|\bnpa\b|\bemi\b)", n): return 2
        # 0 = real additive money/quantity totals (English + currency symbols + EU terms)
        if re.search(r"(total|amount|amt|value|revenue|sales|net|gross|disbursed|outstanding|balance|cost|spend|budget|\bfee|income|profit|paid|collected|\bdue\b|salary|payroll|turnover|charge|freight|€|£|\$|₹|umsatz|betrag|chiffre|montant|ventas|importe)", n): return 0
        return 1
    # Primary measure: prefer real money totals, never a per-unit price
    pool = [m for m in measures if _is_money(m)] or measures
    primary = sorted(pool, key=_measure_rank)[0] if pool else None

    # ── KPIs: only auto-total columns that are CLEARLY additive (money/value).
    # Accuracy over richness — better to show fewer KPIs than wrong sums. ──
    kpi_measures = [m for m in measures if _measure_rank(m) == 0][:6]
    kpis: dict = {"rows": len(df)}
    for m in kpi_measures:
        try:
            total = float(pd.to_numeric(df[m], errors="coerce").sum())
            kpis[f"{m} (total)"] = round(total, 2)
            if _is_money(m):
                kpis[f"{m} (total)_fmt"] = _fmt(total)
        except Exception:
            continue

    result: dict = {"kpis": kpis}

    def _breakdown(dim, meas):
        cols = [dim] + ([meas] if meas else [])
        g = df[cols].copy()
        g = g[g[dim].notna()]
        lab = g[dim].astype(str).str.strip()
        keep = ~lab.str.lower().isin(["nan", "none", ""])
        g, lab = g[keep], lab[keep]
        vals = pd.to_numeric(g[meas], errors="coerce") if meas else pd.Series(1, index=g.index)
        grouped = vals.groupby(lab).sum().sort_values(ascending=False)
        grand = float(grouped.sum()) or 1.0          # % is honest — vs the true total
        return [{"name": str(k), "value": round(float(v), 2), "pct": round(float(v) / grand * 100, 1)}
                for k, v in grouped.head(10).items()]

    # ── Auto breakdowns: primary measure by the most useful dimensions (top-N each) ──
    if primary:
        for dim in _groupable(df, profile)[:4]:
            try:
                rows = _breakdown(dim, primary)
                if rows:
                    result[f"{primary} by {dim}"] = rows
            except Exception:
                continue

    # ── On-demand view: measure by chosen group_by ──
    if group_by and group_by in df.columns:
        meas = measure if (measure and measure in df.columns) else primary
        try:
            result[f"{(meas or 'Count')} by {group_by}"] = _breakdown(group_by, meas)
        except Exception:
            pass

    # ── Time trend ──
    if primary and date_cols:
        try:
            dcol = date_cols[0]
            t = df.copy()
            t["_d"] = _to_datetime(t[dcol])
            t["_m"] = pd.to_numeric(t[primary], errors="coerce")
            t = t.dropna(subset=["_d"])
            if len(t):
                ser = t.groupby(t["_d"].dt.to_period("M"))["_m"].sum().sort_index().tail(24)
                result["monthly_trend"] = [{"month": str(k), "value": round(float(v), 2)} for k, v in ser.items()]
        except Exception:
            pass

    # ── Honest auto-insights (no fabrication, ranked by real impact) ──
    alerts = []
    if primary:
        try:
            bys = [k for k in result if k.startswith(f"{primary} by ")]
            if bys:
                top = result[bys[0]][0]
                alerts.append({"type": "Opportunity",
                               "text": f"{top['name']} is the largest in {bys[0].split(' by ')[-1]} "
                                       f"({top['pct']}% of {primary})."})
        except Exception:
            pass
    worst = max(profile, key=lambda p: p["missing"], default=None)
    if worst and worst["missing"] > 0:
        pct = round(worst["missing"] / max(len(df), 1) * 100)
        if pct >= 10:
            alerts.append({"type": "Warning",
                           "text": f"Column '{worst['column']}' is {pct}% blank — results for it may be partial."})
    if not alerts:
        alerts.append({"type": "Info", "text": f"Analysed {len(df)} rows across {len(df.columns)} columns."})
    result["alerts"] = alerts

    result["columns_detected"] = {p["role"]: p["column"] for p in profile
                                  if p["role"] in ("date", "number", "category")}
    return result


# ════════════════════════════════════════════════════════════════════════════
# DOMAIN SMART PACKS — industry-specific KPIs, breakdowns & alerts layered on top
# of the universal engine. Each pack only runs the parts whose columns it finds,
# so it adds real value when the data fits and silently skips when it doesn't
# (the universal base always ran first, so nothing ever breaks).
# ════════════════════════════════════════════════════════════════════════════

# Synonym / translation table — when a pack asks for a concept, _col ALSO tries
# these equivalents (abbreviations + EU/Latin languages) so engines capture as
# many real-world column names as possible. Kept to true equivalents to avoid
# false matches; the pack's explicit keyword always has priority.
_SYN = {
    "revenue":    ["umsatz", "turnover", "ventas", "ingresos", "receita", "fatturato", "chiffre", "net sales", "gross sales"],
    "sales":      ["verkauf", "ventas", "vendas", "vendite"],
    "amount":     ["betrag", "montant", "importe", "importo", "montante"],
    "total":      ["gesamt", "totale", "suma"],
    "quantity":   ["menge", "cantidad", "quantité", "quantidade", "stück", "qté", "no. of units"],
    "units":      ["einheiten", "unidades", "unità"],
    "price":      ["preis", "precio", "prix", "prezzo", "preço"],
    "cost":       ["kosten", "costo", "coût", "custo", "cogs"],
    "profit":     ["gewinn", "beneficio", "bénéfice", "profitto", "lucro"],
    "budget":     ["haushalt", "presupuesto"],
    "target":     ["ziel", "objetivo", "objectif", "obiettivo", "meta", "quota", "goal"],
    "discount":   ["rabatt", "descuento", "remise", "sconto", "desconto", "markdown"],
    "date":       ["datum", "fecha", "data", "dato", "tarih"],
    "customer":   ["kunde", "kunden", "cliente", "clientes", "client", "buyer"],
    "product":    ["produkt", "producto", "artikel", "articolo", "produto", "prodotto"],
    "category":   ["kategorie", "categoria", "catégorie", "categoría"],
    "region":     ["gebiet", "région", "regione", "região", "territory", "zone"],
    "city":       ["stadt", "ciudad", "ville", "città", "cidade"],
    "country":    ["land", "pays", "país", "paese"],
    "salesperson":["verkäufer", "vendeur", "vendedor", "venditore"],
    "supplier":   ["lieferant", "proveedor", "fournisseur", "fornitore", "vendor"],
    "vendor":     ["lieferant", "proveedor", "fournisseur", "fornitore", "supplier", "payee"],
    "warehouse":  ["lager", "almacén", "entrepôt", "magazzino", "godown", "depot"],
    "department": ["abteilung", "département", "departamento", "reparto", "dept", "division"],
    "employee":   ["mitarbeiter", "empleado", "employé", "dipendente", "personnel"],
    "salary":     ["gehalt", "salario", "salaire", "stipendio", "wage", "ctc", "compensation"],
    "stock":      ["bestand", "existencias"],
    "channel":    ["kanal", "canal", "canale", "marketplace"],
    "rating":     ["bewertung", "valoración", "valutazione", "csat"],
}

def _col(profile, *keywords, role=None):
    # expand each requested keyword with its known equivalents (priority preserved)
    expanded = []
    for kw in keywords:
        expanded.append(kw)
        expanded.extend(_SYN.get(kw, []))
    for kw in expanded:
        for p in profile:
            if kw in str(p["column"]).lower() and (role is None or p["role"] == role):
                return p["column"]
    return None

def _numv(df, c):
    return pd.to_numeric(df[c], errors="coerce")

def _pct(a, b):
    return round(a / b * 100, 1) if b else 0.0

def _sumrows(df, label_col, val, n=10, asc=False):
    lab = df[label_col].astype(str).str.strip()
    keep = ~lab.str.lower().isin(["nan", "none", ""])
    s = val[keep].groupby(lab[keep]).sum().sort_values(ascending=asc)
    grand = float(s.sum()) or 1.0
    return [{"name": str(k), "value": round(float(v), 2), "pct": round(float(v) / grand * 100, 1)}
            for k, v in s.head(n).items()]

def _meanrows(df, label_col, val, n=10):
    lab = df[label_col].astype(str).str.strip()
    keep = ~lab.str.lower().isin(["nan", "none", ""])
    s = val[keep].groupby(lab[keep]).mean().sort_values(ascending=False)
    return [{"name": str(k), "value": round(float(v), 1)} for k, v in s.head(n).items()]

def _countrows(df, label_col, n=10):
    """Count of rows per category — uses a non-money 'count' key so the UI doesn't $-format it."""
    lab = df[label_col].astype(str).str.strip()
    keep = ~lab.str.lower().isin(["nan", "none", ""])
    vc = lab[keep].value_counts().head(n)
    grand = int(vc.sum()) or 1
    return [{"name": str(k), "count": int(v), "pct": round(v / grand * 100, 1)} for k, v in vc.items()]

def _money_kpi(K, label, value):
    K[label] = round(float(value), 2)
    K[label + "_fmt"] = _fmt(float(value))


def pack_sales(df, profile):
    """Best-in-class Sales engine — the salesman's real questions:
    targets, growth, forecast, segments, margin, customers, products, anomalies."""
    K, S, A = {}, {}, []
    rev    = _col(profile, "revenue", "sales", "amount", "umsatz", "ventas", "total", role="number")
    tgt    = _col(profile, "target", "goal", "quota")
    prod   = _col(profile, "product", "item", "sku", "article")
    person = _col(profile, "salesperson", "sales person", "rep", "executive", "agent", "seller")
    region = _col(profile, "region", "territory", "zone", "state", "country", "area", "ville", "land")
    channel= _col(profile, "channel", "source", "platform", "medium")
    cust   = _col(profile, "customer", "client", "account", "buyer", "company")
    cost   = _col(profile, "total cost", "cogs", "cost of goods", "purchase cost")
    qty    = _col(profile, "quantity", "qty", "units", "menge", role="number")
    date   = _col(profile, "date", "order date", "invoice date", "day", "month", role="date")
    if not rev:
        return K, S, A

    r = _numv(df, rev); tot = float(r.sum()); n = int(r.notna().sum()) or 1
    _money_kpi(K, "Total Revenue", tot)
    K["Orders"] = n
    _money_kpi(K, "Avg Order Value", tot / n)
    if qty: K["Units Sold"] = int(_numv(df, qty).sum())

    # ── Target attainment ──
    if tgt:
        t = float(_numv(df, tgt).sum())
        if t:
            ach = _pct(tot, t); K["Target Achievement %"] = ach
            if ach < 80:    A.append({"type": "Critical",    "text": f"Revenue at {ach}% of target — below 80%."})
            elif ach >= 100: A.append({"type": "Opportunity", "text": f"Target beaten — {ach}% achieved."})

    # ── Growth + forecast (needs a date column) ──
    if date:
        try:
            t2 = df[[date, rev]].copy()
            t2["_d"] = _to_datetime(t2[date]); t2["_r"] = _numv(t2, rev)
            t2 = t2.dropna(subset=["_d"])
            monthly = t2.groupby(t2["_d"].dt.to_period("M"))["_r"].sum().sort_index()
            if len(monthly) >= 2:
                mom = _pct(float(monthly.iloc[-1] - monthly.iloc[-2]), float(monthly.iloc[-2]))
                K["MoM Growth %"] = mom
                if mom < 0: A.append({"type": "Warning", "text": f"Revenue down {abs(mom)}% vs last month."})
                S["Monthly revenue"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in monthly.tail(12).items()]
            yearly = t2.groupby(t2["_d"].dt.year)["_r"].sum().sort_index()
            if len(yearly) >= 2:
                K["YoY Growth %"] = _pct(float(yearly.iloc[-1] - yearly.iloc[-2]), float(yearly.iloc[-2]))
            if len(monthly) >= 4:
                from sklearn.linear_model import LinearRegression
                y = monthly.values.astype(float); x = np.arange(len(y)).reshape(-1, 1)
                pred = LinearRegression().fit(x, y).predict(np.arange(len(y), len(y) + 3).reshape(-1, 1))
                pred = [max(0.0, float(p)) for p in pred]
                S["Revenue forecast (next 3 mo)"] = [{"month": f"+{i+1}mo", "value": round(pred[i], 2)} for i in range(3)]
                _money_kpi(K, "Forecast (next 3 mo)", sum(pred))
        except Exception:
            pass

    # ── Margin (needs a total-cost column) ──
    if cost:
        c = float(_numv(df, cost).sum())
        if c and tot:
            _money_kpi(K, "Gross Profit", tot - c)
            margin = _pct(tot - c, tot); K["Gross Margin %"] = margin
            if margin < 15: A.append({"type": "Warning", "text": f"Gross margin is {margin}% — thin."})

    # ── Segments ──
    if region:  S["Revenue by region"]  = _sumrows(df, region, r, 10)
    if channel: S["Revenue by channel"] = _sumrows(df, channel, r, 8)

    # ── Customers (top, new vs repeat, concentration risk) ──
    if cust:
        vc = df[cust].astype(str).value_counts()
        K["Customers"] = int(vc.shape[0]); K["Repeat Customers"] = int((vc > 1).sum())
        top_cust = _sumrows(df, cust, r, 8)
        if top_cust:
            S["Top customers"] = top_cust
            if top_cust[0]["pct"] >= 40:
                A.append({"type": "Warning", "text": f"{top_cust[0]['name']} is {top_cust[0]['pct']}% of revenue — concentration risk."})

    # ── Products: best + slowest ──
    if prod:
        S["Top products"] = _sumrows(df, prod, r, 10)
        S["Slowest products"] = _sumrows(df, prod, r, 5, asc=True)

    # ── People ──
    if person:
        S["Top performers"] = _sumrows(df, person, r, 8)
        S["Needs attention (lowest)"] = _sumrows(df, person, r, 5, asc=True)

    # ── Pipeline / funnel (CRM-style data: stage + won/lost) ──
    status = _col(profile, "stage", "deal stage", "pipeline stage", "opportunity stage", "status", "disposition")
    if status:
        sv = df[status].astype(str).str.lower()
        won  = int(sv.str.contains(r"won|success|complete|closed.?won|paid|deal", na=False).sum())
        lost = int(sv.str.contains(r"lost|cancel|reject|dropped|closed.?lost|fail", na=False).sum())
        if won + lost > 0:
            wr = _pct(won, won + lost); K["Win Rate %"] = wr
            if wr < 30: A.append({"type": "Warning", "text": f"Win rate is {wr}% — most deals are lost."})
        S["Deals by stage"] = _countrows(df, status, 8)

    # ── Discount leakage ──
    discount = _col(profile, "discount", "disc", "markdown", "rebate", role="number")
    if discount:
        dv = _numv(df, discount).dropna()
        if len(dv):
            if dv.max() <= 100 and dv.mean() < 100:   # looks like a percentage
                K["Avg Discount %"] = round(float(dv.mean()), 1)
                if dv.mean() > 20: A.append({"type": "Warning", "text": f"Avg discount {round(float(dv.mean()),1)}% — eating margin."})
            else:                                      # money amount
                _money_kpi(K, "Total Discount", float(dv.sum()))

    # ── Returns ──
    returns = _col(profile, "return status", "returned", "refund", "is return", "return")
    if returns:
        low = df[returns].astype(str).str.lower()
        rc = int(low.str.contains(r"yes|true|return|refund", na=False).sum())
        if 0 < rc < len(df):
            rr2 = _pct(rc, len(df)); K["Return Rate %"] = rr2
            if rr2 > 10: A.append({"type": "Warning", "text": f"Return rate {rr2}% — above 10%."})

    # ── RFM customer segmentation + win-back list (needs customer + date) ──
    if cust and date:
        try:
            t = df[[cust, date, rev]].copy()
            t["_d"] = _to_datetime(t[date]); t["_r"] = pd.to_numeric(t[rev], errors="coerce").fillna(0)
            t = t.dropna(subset=["_d"])
            if len(t) >= 4:
                now = t["_d"].max()
                g = t.groupby(t[cust].astype(str)).agg(recency=("_d", lambda x: (now - x.max()).days),
                                                       freq=("_d", "size"), monetary=("_r", "sum"))
                rmed = g["recency"].median(); fmed = max(g["freq"].median(), 2); mmed = g["monetary"].median()
                def _seg(row):
                    recent = row["recency"] <= rmed; freq = row["freq"] >= fmed; big = row["monetary"] >= mmed
                    if row["freq"] <= 1: return "New / one-time"
                    if recent and freq and big: return "Champions"
                    if recent and (freq or big): return "Loyal"
                    if (not recent) and (freq or big): return "At-risk"
                    if not recent: return "Lost"
                    return "Regular"
                g["seg"] = g.apply(_seg, axis=1)
                vc = g["seg"].value_counts()
                S["Customer segments (RFM)"] = [{"name": k, "count": int(v), "pct": round(v/len(g)*100, 1)} for k, v in vc.items()]
                wb = g[g["seg"].isin(["At-risk", "Lost"])].sort_values("monetary", ascending=False).head(8)
                if len(wb):
                    S["Win-back list (lapsed, by past spend)"] = [{"name": idx, "value": round(float(rw["monetary"]), 2)} for idx, rw in wb.iterrows()]
                    A.append({"type": "Warning", "text": f"{int((g['seg'].isin(['At-risk','Lost'])).sum())} customers have gone quiet — win-back opportunity."})
        except Exception:
            pass

    # ── Seasonality: revenue by calendar month ──
    if date:
        try:
            t = df[[date, rev]].copy(); t["_d"] = _to_datetime(t[date]); t["_r"] = pd.to_numeric(t[rev], errors="coerce").fillna(0); t = t.dropna(subset=["_d"])
            mn = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
            bym = t.groupby(t["_d"].dt.month)["_r"].sum()
            if len(bym) >= 3:
                S["Revenue by month (seasonality)"] = [{"month": mn[int(m)-1], "value": round(float(v), 2)} for m, v in bym.items()]
                A.append({"type": "Info", "text": f"Peak month: {mn[int(bym.idxmax())-1]}."})
        except Exception:
            pass

    # ── Growth drivers: what moved revenue last month (needs date + a dimension) ──
    drv = region or channel or prod
    if date and drv:
        try:
            t = df[[date, drv, rev]].copy(); t["_d"] = _to_datetime(t[date]); t["_r"] = pd.to_numeric(t[rev], errors="coerce").fillna(0); t = t.dropna(subset=["_d"])
            t["_m"] = t["_d"].dt.to_period("M")
            months = sorted(t["_m"].unique())
            if len(months) >= 2:
                last, prev = months[-1], months[-2]
                cur = t[t["_m"] == last].groupby(t[t["_m"] == last][drv].astype(str))["_r"].sum()
                pre = t[t["_m"] == prev].groupby(t[t["_m"] == prev][drv].astype(str))["_r"].sum()
                keys = set(cur.index) | set(pre.index)
                deltas = {k: float(cur.get(k, 0) - pre.get(k, 0)) for k in keys}
                if deltas:
                    up = max(deltas.items(), key=lambda kv: kv[1]); dn = min(deltas.items(), key=lambda kv: kv[1])
                    if up[1] > 0: A.append({"type": "Opportunity", "text": f"Biggest growth: {up[0]} ({_fmt(up[1])} vs last month)."})
                    if dn[1] < 0: A.append({"type": "Warning", "text": f"Biggest drop: {dn[0]} ({_fmt(dn[1])} vs last month)."})
        except Exception:
            pass

    # ── Target pacing (gap to target) ──
    if tgt:
        try:
            tt = float(_numv(df, tgt).sum())
            if tt > 0:
                gap = tt - tot
                if gap > 0: A.append({"type": "Warning", "text": f"{_fmt(gap)} to go to hit target ({_pct(tot, tt)}% there)."})
        except Exception:
            pass

    # ── Cross-sell: products frequently bought together ──
    order = _col(profile, "order id", "invoice", "bill no", "order no", "transaction id", "order")
    bk = order or cust
    if bk and prod:
        try:
            from itertools import combinations
            from collections import Counter
            baskets = df.groupby(df[bk].astype(str))[prod].apply(lambda s: sorted({str(x) for x in s if str(x).strip().lower() not in ("nan","none","")}))
            pairs = Counter()
            for items in baskets:
                if 2 <= len(items) <= 20:
                    for a, b in combinations(items, 2): pairs[(a, b)] += 1
            top = pairs.most_common(8)
            if top:
                S["Frequently bought together"] = [{"name": f"{a} + {b}", "count": int(c)} for (a, b), c in top]
        except Exception:
            pass

    # ── Anomalies: unusually large orders ──
    try:
        rr = r.dropna()
        if len(rr) > 12:
            thr = float(rr.mean() + 3 * rr.std())
            big = int((rr > thr).sum())
            if big: A.append({"type": "Warning", "text": f"{big} unusually large order(s) (> {_fmt(thr)}) — worth checking."})
    except Exception:
        pass

    return K, S, A


def pack_inventory(df, profile):
    """Best-in-class Inventory engine — stock value, reorder, days-of-cover,
    revenue-at-risk, overstock, turnover, margin/GMROI, movers, supplier spend,
    ABC concentration, dead stock and expiry."""
    K, S, A = {}, {}, []
    stock   = _col(profile, "current stock", "balance", "stock", "on hand", "units in stock", "qty", "quantity", role="number")
    value   = _col(profile, "stock value", "inventory value", "total cost", role="number")
    ucost   = _col(profile, "unit cost", "cost per", "cost/piece", "purchase cost", role="number")
    sell    = _col(profile, "selling price", "sale price", "retail price", "mrp", "list price", role="number")
    reorder = _col(profile, "reorder", "min level", "min stock", "reorder point", "safety stock", role="number")
    maxs    = _col(profile, "max stock", "maximum stock", "max level", role="number")
    demand  = _col(profile, "monthly demand", "demand", "avg demand", "forecast demand", "monthly sales", role="number")
    sold30  = _col(profile, "sold last 30", "units sold last 30", "last 30 days", "30 day", role="number")
    sold90  = _col(profile, "sold last 90", "units sold last 90", "last 90 days", "90 day", role="number")
    lastsold= _col(profile, "last sold", "days since", "last movement", role="number")
    name    = _col(profile, "item name", "product", "item", "description", "name", "sku")
    cat     = _col(profile, "category", "type", "group")
    wh      = _col(profile, "warehouse", "location", "store", "godown", "branch")
    supplier= _col(profile, "supplier", "vendor")
    brand   = _col(profile, "brand", "make", "manufacturer")
    expiry  = _col(profile, "expiry", "expiration", "exp date", "best before", role="date")

    K["Total SKUs"] = len(df)
    s  = _numv(df, stock) if stock else None
    uc = _numv(df, ucost) if ucost else None
    sp = _numv(df, sell)  if sell  else None

    # stock value: prefer a value column, else stock × unit cost
    val = None
    if value is not None:                       val = _numv(df, value)
    elif uc is not None and s is not None:      val = s * uc
    if val is not None:
        _money_kpi(K, "Stock Value", float(val.sum()))

    if s is not None:
        K["Units in Stock"] = int(round(float(s.fillna(0).sum())))
        oos = int((s <= 0).sum()); low = int(((s > 0) & (s <= 5)).sum())
        K["Out of Stock"] = oos
        if reorder is not None:
            rl = _numv(df, reorder)
            need = int(((s.notna()) & (rl.notna()) & (s < rl)).sum())
            if need:
                K["Reorder Needed"] = need
                A.append({"type": "Warning", "text": f"{need} item(s) below reorder level — restock."})
        if oos: A.append({"type": "Critical", "text": f"{oos} item(s) out of stock."})
        if low: K["Low Stock (≤5)"] = low

    # ── Revenue at risk — out-of-stock items × demand × price (the cost of stockouts) ──
    if s is not None and demand and sp is not None:
        dm = _numv(df, demand); oosm = (s <= 0)
        atrisk = float((dm[oosm].fillna(0) * sp[oosm].fillna(0)).sum())
        if atrisk > 0:
            _money_kpi(K, "Revenue at Risk", atrisk)
            A.append({"type": "Critical", "text": f"{_fmt(atrisk)}/mo of sales at risk from out-of-stock items."})

    # ── Days of cover — how long until each item runs out at current demand ──
    if s is not None and demand:
        daily = (_numv(df, demand) / 30.0)
        cover = s / daily.replace(0, np.nan)
        valid = cover.replace([np.inf, -np.inf], np.nan).dropna()
        if len(valid):
            K["Avg Days of Cover"] = int(round(float(valid.mean())))
            soon = int((valid <= 14).sum())
            if soon: A.append({"type": "Warning", "text": f"{soon} item(s) will run out within 2 weeks at current demand."})
            if name is not None:
                try:
                    cd = pd.DataFrame({"name": df[name].astype(str), "days": cover})
                    cd = cd[(cd["days"].notna()) & np.isfinite(cd["days"]) & (cd["days"] > 0)].sort_values("days").head(8)
                    if len(cd): S["Running out soonest (days of cover)"] = [{"name": r["name"], "days": int(r["days"])} for _, r in cd.iterrows()]
                except Exception:
                    pass

    # ── Margin & potential profit (needs selling price + unit cost) ──
    if sp is not None and uc is not None:
        mp = (sp - uc) / sp.replace(0, np.nan)
        if mp.notna().any():
            K["Avg Margin %"] = round(float(mp.dropna().mean()) * 100, 1)
        if s is not None:
            pot = float(((sp - uc) * s).clip(lower=0).sum())
            if pot > 0: _money_kpi(K, "Potential Profit", pot)

    # ── Overstock / excess capital (stock above max level) ──
    if s is not None and maxs:
        mx = _numv(df, maxs)
        over = (s.notna()) & (mx.notna()) & (s > mx)
        nover = int(over.sum())
        if nover:
            K["Overstock Items"] = nover
            if uc is not None:
                excess = float(((s - mx)[over] * uc[over]).clip(lower=0).sum())
                if excess > 0:
                    _money_kpi(K, "Excess Capital", excess)
                    A.append({"type": "Warning", "text": f"{nover} overstocked item(s) — {_fmt(excess)} tied up in excess inventory."})
            else:
                A.append({"type": "Warning", "text": f"{nover} item(s) above max stock level."})

    # ── Turnover + fast/slow movers (needs a units-sold column) ──
    movecol = sold90 or sold30
    if movecol and name is not None:
        mv = _numv(df, movecol).fillna(0)
        if s is not None:
            tot_units = float(s.fillna(0).sum()) or 1.0
            K["Stock Turnover"] = round(float(mv.sum()) / tot_units, 2)
        try:
            g = mv.groupby(df[name].astype(str)).sum()
            g = g[~g.index.str.lower().isin(["nan", "none", ""])].sort_values(ascending=False)
            tot = float(g.sum()) or 1.0
            S["Fastest movers (units sold)"] = [{"name": str(k), "units": int(v), "pct": round(v/tot*100, 1)} for k, v in g.head(8).items()]
            slow = g[g > 0].sort_values().head(5) if (g > 0).any() else g.sort_values().head(5)
            S["Slowest movers (units sold)"] = [{"name": str(k), "units": int(v), "pct": round(v/tot*100, 1)} for k, v in slow.items()]
        except Exception:
            pass

    # ── Dead stock — no recent sales (prefer 90-day units, else days-since-sold) ──
    dead = None
    if sold90:    dead = int((_numv(df, sold90).fillna(0) <= 0).sum())
    elif lastsold: dead = int((_numv(df, lastsold) > 90).sum())
    if dead:
        K["Dead Stock"] = dead
        A.append({"type": "Warning", "text": f"{dead} item(s) with no recent sales — dead stock tying up cash."})

    # ── Segment breakdowns by value ──
    if val is not None:
        if name is not None:     S["Top items by value"] = _sumrows(df, name, val, 10)
        if cat is not None:      S["Stock value by category"] = _sumrows(df, cat, val, 10)
        if wh is not None:       S["Stock value by warehouse"] = _sumrows(df, wh, val, 8)
        if supplier is not None: S["Stock value by supplier"] = _sumrows(df, supplier, val, 8)
        elif brand is not None:  S["Stock value by brand"] = _sumrows(df, brand, val, 8)

        # ── ABC analysis (value concentration) ──
        if name is not None:
            try:
                g = val.groupby(df[name].astype(str)).sum().sort_values(ascending=False)
                total = float(g.sum()) or 1.0
                cum = g.cumsum() / total
                a = int((cum <= 0.8).sum()); b = int(((cum > 0.8) & (cum <= 0.95)).sum()); c = len(g) - a - b
                va = float(g[cum <= 0.8].sum()); vb = float(g[(cum > 0.8) & (cum <= 0.95)].sum()); vc = total - va - vb
                S["Value by ABC class"] = [
                    {"name": f"A — {a} items", "value": round(va, 2), "pct": round(va/total*100, 1)},
                    {"name": f"B — {b} items", "value": round(vb, 2), "pct": round(vb/total*100, 1)},
                    {"name": f"C — {c} items", "value": round(vc, 2), "pct": round(vc/total*100, 1)},
                ]
                if a: A.append({"type": "Opportunity", "text": f"{a} 'A' items ({round(a/len(g)*100)}% of SKUs) hold {round(va/total*100)}% of stock value — focus here."})
            except Exception:
                pass

    # ── Expiry ──
    if expiry is not None:
        d = _to_datetime(df[expiry])
        if d is not None and d.notna().any():
            delta = (d - pd.Timestamp.now()).dt.days
            expired = int(((d.notna()) & (delta < 0)).sum())
            soon30  = int(((d.notna()) & (delta >= 0) & (delta <= 30)).sum())
            soon90  = int(((d.notna()) & (delta >= 0) & (delta <= 90)).sum())
            if soon30: K["Expiring ≤30d"] = soon30
            if soon90: K["Expiring ≤90d"] = soon90
            if expired:   A.append({"type": "Critical", "text": f"{expired} item(s) already expired."})
            elif soon30:  A.append({"type": "Warning",  "text": f"{soon30} item(s) expire within 30 days."})
            if name is not None:
                try:
                    ed = pd.DataFrame({"name": df[name].astype(str), "days": delta})
                    ed = ed[(ed["days"].notna()) & (ed["days"] >= 0)].sort_values("days").head(8)
                    if len(ed): S["Expiring soonest (days left)"] = [{"name": r["name"], "days": int(r["days"])} for _, r in ed.iterrows()]
                except Exception:
                    pass

    return K, S, A


def pack_finance(df, profile):
    """CFO-grade finance engine: P&L (revenue / expense / net / margin), budget
    variance, monthly cash-flow trend, expense drivers, department & vendor spend.
    Works on any ledger / transactions / P&L sheet."""
    K, S, A = {}, {}, []
    amt    = _col(profile, "amount", "value", "total", "amt", "umsatz", "betrag", role="number")
    budget = _col(profile, "budget", "planned", "forecast", "target", role="number")
    typ    = _col(profile, "type", "transaction type", "txn type", "head", "account", "category", "flow")
    dept   = _col(profile, "department", "dept", "cost center", "cost centre", "team", "division", "unit")
    subcat = _col(profile, "sub-category", "subcategory", "sub category", "line item", "category", "head", "account")
    vendor = _col(profile, "vendor", "supplier", "payee", "merchant", "counterparty", "paid to")
    date   = _col(profile, "date", "month", "period", "posting date", "txn date", role="date")
    if not amt:
        return K, S, A
    a = _numv(df, amt)
    K["Transactions"] = int(a.notna().sum())

    # ── classify income vs expense (by a type column, else by sign) ──
    inc_mask = exp_mask = None
    if typ:
        low = df[typ].astype(str).str.lower()
        inc_mask = low.str.contains(r"revenue|income|credit|sales|inflow|receipt|earning", na=False)
        exp_mask = low.str.contains(r"expense|cost|debit|cogs|purchase|spend|payment|outflow|payroll|opex|capex|bill", na=False)
    if (inc_mask is None or not (inc_mask.any() or exp_mask.any())) and (a < 0).any():
        inc_mask, exp_mask = (a > 0), (a < 0)

    income  = float(a[inc_mask].sum())       if inc_mask is not None else 0.0
    expense = float(a[exp_mask].abs().sum())  if exp_mask is not None else 0.0

    if income or expense:
        if income:  _money_kpi(K, "Total Revenue", income)
        if expense: _money_kpi(K, "Total Expenses", expense)
        net = income - expense
        _money_kpi(K, "Net Profit", net)
        if income:
            K["Profit Margin %"] = _pct(net, income)
            K["Expense Ratio %"] = _pct(expense, income)
        S["Income vs expense"] = [
            {"name": "Revenue",  "value": round(income, 2),  "pct": _pct(income, income + expense)},
            {"name": "Expenses", "value": round(expense, 2), "pct": _pct(expense, income + expense)}]
        if net < 0:
            A.append({"type": "Critical", "text": f"Operating at a loss of {_fmt(abs(net))} — expenses exceed revenue."})
        elif income and _pct(net, income) < 10:
            A.append({"type": "Warning", "text": f"Profit margin is {_pct(net, income)}% — thin; watch costs."})
        elif income and _pct(net, income) >= 25:
            A.append({"type": "Opportunity", "text": f"Healthy {_pct(net, income)}% profit margin."})
    else:
        _money_kpi(K, "Total Amount", float(a.sum()))

    # ── budget variance (actual vs planned, like-for-like over all lines) ──
    if budget:
        b = float(_numv(df, budget).sum())
        actual = float(a.sum())
        if b:
            var = _pct(actual - b, b); K["Budget Variance %"] = var
            if dept:
                bd = _numv(df, budget)
                rows = []
                for d, g in df.groupby(df[dept].astype(str)):
                    if str(d).lower() in ("nan", "none", ""):
                        continue
                    diff = float(a[g.index].sum() - bd[g.index].sum())
                    rows.append({"name": str(d), "value": round(diff, 2), "pct": 0.0})
                rows.sort(key=lambda r: r["value"], reverse=True)
                if rows:
                    S["Budget variance by department"] = rows
                    worst = rows[0]
                    if worst["value"] > 0:
                        A.append({"type": "Warning", "text": f"{worst['name']} is the biggest budget overrun ({_fmt(worst['value'])} over plan)."})

    # ── expense drivers & revenue mix ──
    if exp_mask is not None and exp_mask.any():
        ea = a[exp_mask].abs(); ed = df[exp_mask]
        if subcat: S["Top expenses (by category)"] = _sumrows(ed, subcat, ea, 8)
        if dept:   S["Expenses by department"]     = _sumrows(ed, dept, ea, 10)
        if vendor:
            tv = _sumrows(ed, vendor, ea, 8)
            if tv:
                S["Top vendors by spend"] = tv
                if tv[0]["pct"] >= 40:
                    A.append({"type": "Warning", "text": f"{tv[0]['name']} is {tv[0]['pct']}% of all spend — vendor concentration risk."})
        top_exp = S.get("Top expenses (by category)") or []
        if top_exp:
            A.append({"type": "Info", "text": f"Biggest cost: {top_exp[0]['name']} ({_fmt(top_exp[0]['value'])}, {top_exp[0]['pct']}% of spend)."})
    if inc_mask is not None and inc_mask.any():
        ia = a[inc_mask]; idf = df[inc_mask]
        if subcat: S["Revenue by category"]   = _sumrows(idf, subcat, ia, 8)
        if dept:   S["Revenue by department"]  = _sumrows(idf, dept, ia, 10)

    # ── monthly cash-flow trend (net = income − expense per month) ──
    if date and (income or expense):
        try:
            t = df[[date]].copy()
            t["_d"] = _to_datetime(df[date]); t["_a"] = a
            t["_inc"] = a.where(inc_mask, 0.0); t["_exp"] = a.where(exp_mask, 0.0).abs()
            t = t.dropna(subset=["_d"])
            m = t.groupby(t["_d"].dt.to_period("M"))
            net_m = (m["_inc"].sum() - m["_exp"].sum()).sort_index()
            if len(net_m) >= 2:
                S["Net cash flow by month"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in net_m.tail(12).items()]
                rev_m = m["_inc"].sum().sort_index()
                if float(rev_m.sum()) > 0:
                    S["Revenue by month"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in rev_m.tail(12).items()]
                exp_m = m["_exp"].sum().sort_index()
                if len(exp_m) >= 2 and float(exp_m.iloc[-2]) > 0:
                    mom = _pct(float(exp_m.iloc[-1] - exp_m.iloc[-2]), float(exp_m.iloc[-2]))
                    if mom > 20:
                        A.append({"type": "Warning", "text": f"Expenses jumped {mom}% vs last month — review spend."})
        except Exception:
            pass
    return K, S, A


def pack_hr(df, profile):
    """HR / people-analytics — headcount, payroll, attrition, tenure, diversity
    & pay-gap, performance, overtime/burnout, leave."""
    K, S, A = {}, {}, []
    sal     = _col(profile, "salary", "ctc", "pay", "compensation", "gross pay", "wage", role="number")
    dept    = _col(profile, "department", "dept", "team", "function")
    desig   = _col(profile, "designation", "title", "role", "position", "job title", "grade")
    gender  = _col(profile, "gender", "sex")
    age     = _col(profile, "age", role="number")
    exp     = _col(profile, "experience", "tenure", "years of service", role="number")
    loc     = _col(profile, "location", "city", "office", "site", "branch")
    perf    = _col(profile, "performance", "appraisal", "rating", "score", role="number")
    overtime= _col(profile, "overtime", "ot hours", "extra hours", role="number")
    leave   = _col(profile, "leave", "absence", "pto", "time off", role="number")
    status  = _col(profile, "status", "employment status", "active", "attrition")
    K["Headcount"] = len(df)

    if status:
        low = df[status].astype(str).str.lower()
        leftm = low.str.contains(r"left|resign|inactive|exit|attrit|terminat|separat", na=False)
        left = int(leftm.sum())
        if left:
            K["Active Headcount"] = len(df) - left
            ar = _pct(left, len(df)); K["Attrition %"] = ar
            if ar > 15: A.append({"type": "Critical", "text": f"Attrition is {ar}% — above the healthy 10–12% range."})
            if dept:
                abd = _countrows(df[leftm], dept, 6)
                if abd:
                    S["Resignations by department"] = abd
                    if abd[0]["count"] >= 3: A.append({"type": "Warning", "text": f"{abd[0]['name']} has the most exits ({abd[0]['count']}) — investigate."})

    if sal:
        s = _numv(df, sal)
        _money_kpi(K, "Total Payroll", float(s.sum())); _money_kpi(K, "Avg Salary", float(s.mean()))
        if dept:
            try:
                g = s.groupby(df[dept].astype(str)).mean().sort_values(ascending=False)
                g = g[~g.index.str.lower().isin(["nan", "none", ""])]
                if len(g): S["Avg salary by department"] = [{"name": str(k), "value": round(float(v), 0)} for k, v in g.head(10).items()]
            except Exception: pass
        if gender:
            try:
                gs = s.groupby(df[gender].astype(str)).mean(); gs = gs[~gs.index.str.lower().isin(["nan", "none", ""])]
                if len(gs) >= 2:
                    gap = _pct(float(gs.max()) - float(gs.min()), float(gs.max()))
                    if gap >= 10: A.append({"type": "Warning", "text": f"Gender pay gap of {gap}% — {gs.idxmax()} paid more on average."})
            except Exception: pass

    if exp: K["Avg Tenure (yrs)"] = round(float(_numv(df, exp).mean()), 1)
    if age: K["Avg Age"] = round(float(_numv(df, age).mean()), 1)
    if perf:
        pv = _numv(df, perf).dropna()
        if len(pv):
            K["Avg Performance"] = round(float(pv.mean()), 1)
            scale = 5 if pv.max() <= 5 else (10 if pv.max() <= 10 else 100)
            lowp = int((pv < scale * 0.5).sum())
            if lowp: A.append({"type": "Warning", "text": f"{lowp} employee(s) below half the performance scale."})
            if dept:
                try:
                    g = _numv(df, perf).groupby(df[dept].astype(str)).mean().sort_values()
                    g = g[~g.index.str.lower().isin(["nan", "none", ""])]
                    if len(g): S["Avg performance by department"] = [{"name": str(k), "score": round(float(v), 2)} for k, v in g.head(10).items()]
                except Exception: pass

    if gender:
        gc = _countrows(df, gender, 6)
        if gc:
            S["Gender split"] = gc
            if gc[0]["pct"] >= 70: A.append({"type": "Info", "text": f"Workforce is {gc[0]['pct']}% {gc[0]['name']} — limited gender balance."})
    if overtime:
        ot = _numv(df, overtime).dropna()
        if len(ot):
            K["Avg Overtime (hrs)"] = round(float(ot.mean()), 1)
            burn = int((ot > max(float(ot.mean()) * 2, 20)).sum())
            if burn: A.append({"type": "Warning", "text": f"{burn} employee(s) logging very high overtime — burnout risk."})
    if leave:
        lv = _numv(df, leave).dropna()
        if len(lv): K["Avg Leave Days"] = round(float(lv.mean()), 1)

    if dept:  S["Headcount by department"] = _countrows(df, dept, 12)
    if loc:   S["Headcount by location"] = _countrows(df, loc, 10)
    if desig: S["Headcount by designation"] = _countrows(df, desig, 10)
    return K, S, A


def pack_retail(df, profile):
    """Best-in-class Retail / e-commerce engine — revenue, channels, returns,
    discount leakage, new-vs-returning, geography, delivery & ratings."""
    K, S, A = {}, {}, []
    rev      = _col(profile, "revenue", "sales", "amount", "total", "net amount", role="number")
    qty      = _col(profile, "quantity", "qty", "units", role="number")
    price    = _col(profile, "unit price", "price", "rate", "mrp", role="number")
    disc     = _col(profile, "discount", "disc", "markdown", role="number")
    channel  = _col(profile, "channel", "source", "platform", "marketplace")
    cat      = _col(profile, "category", "department", "segment")
    product  = _col(profile, "product", "item", "sku", "article")
    city     = _col(profile, "city", "location", "region", "state", "store", "zone")
    pay      = _col(profile, "payment method", "payment", "pay mode", "tender")
    custtype = _col(profile, "customer type", "customer segment", "new", "returning", "buyer type")
    returns  = _col(profile, "returned", "return", "refund")
    rating   = _col(profile, "rating", "review", "csat", "stars", role="number")
    delivery = _col(profile, "delivery days", "delivery time", "days to deliver", "shipping days", "lead time", role="number")
    date     = _col(profile, "date", "order date", "invoice date", role="date")

    if not rev:
        return K, S, A
    r = _numv(df, rev); tot = float(r.sum()); n = int(r.notna().sum()) or 1
    _money_kpi(K, "Total Revenue", tot)
    K["Orders"] = n
    _money_kpi(K, "Avg Order Value", tot / n)
    if qty: K["Units Sold"] = int(_numv(df, qty).fillna(0).sum())

    # ── Monthly sales trend (shown first) ──
    if date:
        try:
            t = df[[date]].copy(); t["_d"] = _to_datetime(df[date]); t["_r"] = r
            t = t.dropna(subset=["_d"])
            m = t.groupby(t["_d"].dt.to_period("M"))["_r"].sum().sort_index()
            if len(m) >= 2:
                S["Monthly revenue"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in m.tail(12).items()]
        except Exception:
            pass

    # ── Revenue breakdowns ──
    if channel:
        S["Revenue by channel"] = _sumrows(df, channel, r, 8)
        tc = S["Revenue by channel"]
        if tc and tc[0]["pct"] >= 50:
            A.append({"type": "Warning", "text": f"{tc[0]['name']} is {tc[0]['pct']}% of revenue — channel concentration risk."})
    if cat:     S["Revenue by category"] = _sumrows(df, cat, r, 10)
    if product: S["Top products"] = _sumrows(df, product, r, 10)
    if city:    S["Revenue by city"] = _sumrows(df, city, r, 10)
    if pay:     S["Revenue by payment method"] = _sumrows(df, pay, r, 6)

    # ── Discounts & leakage ──
    if disc:
        dv = _numv(df, disc).dropna()
        if len(dv) and dv.max() <= 100:
            K["Avg Discount %"] = round(float(dv.mean()), 1)
            if price and qty:
                gross = float((_numv(df, price) * _numv(df, qty)).sum())
                if gross > tot:
                    _money_kpi(K, "Discount Given", gross - tot)
            if float(dv.mean()) > 25:
                A.append({"type": "Warning", "text": f"Avg discount is {round(float(dv.mean()),1)}% — discount leakage eating margin."})

    # ── Returns ──
    if returns:
        low = df[returns].astype(str).str.lower()
        retm = low.str.contains(r"\byes\b|\btrue\b|return|refund|^1$", na=False)
        rc = int(retm.sum())
        if rc:
            rr = _pct(rc, len(df)); K["Return Rate %"] = rr
            lost = float(r[retm].sum())
            if lost > 0: _money_kpi(K, "Returned Revenue", lost)
            if rr > 10: A.append({"type": "Warning", "text": f"Return rate is {rr}% — above the 10% benchmark ({_fmt(lost)} returned)."})
            if cat:
                rcat = _countrows(df[retm], cat, 6)
                if rcat: S["Returns by category"] = rcat

    # ── New vs returning customers ──
    if custtype:
        S["Revenue by customer type"] = _sumrows(df, custtype, r, 6)
        low = df[custtype].astype(str).str.lower()
        ret_cust = int(low.str.contains("return|repeat|existing|loyal", na=False).sum())
        if ret_cust:
            rp = _pct(ret_cust, len(df)); K["Repeat Customer %"] = rp
            if rp < 30: A.append({"type": "Opportunity", "text": f"Only {rp}% returning customers — retention upside."})

    # ── Ratings ──
    if rating:
        rt = _numv(df, rating).dropna()
        if len(rt):
            K["Avg Rating"] = round(float(rt.mean()), 1)
            if float(rt.mean()) < 3.5:
                A.append({"type": "Warning", "text": f"Avg rating is {round(float(rt.mean()),1)}/5 — customer satisfaction is low."})
            if cat:
                try:
                    g = _numv(df, rating).groupby(df[cat].astype(str)).mean().sort_values()
                    g = g[~g.index.str.lower().isin(["nan", "none", ""])]
                    if len(g): S["Lowest-rated categories"] = [{"name": str(k), "rating": round(float(v), 2)} for k, v in g.head(6).items()]
                except Exception:
                    pass

    # ── Delivery performance ──
    if delivery:
        dd = _numv(df, delivery).dropna()
        if len(dd):
            K["Avg Delivery Days"] = round(float(dd.mean()), 1)
            slow = int((dd > 7).sum())
            if slow: A.append({"type": "Warning", "text": f"{slow} order(s) took over a week to deliver."})
    return K, S, A


def pack_logistics(df, profile):
    """Logistics / fleet — delivery cost & variance, on-time %, delays, damage
    rate, cost-per-km, route/vehicle/cargo breakdowns and ratings."""
    K, S, A = {}, {}, []
    cost    = _col(profile, "delivery cost", "freight", "shipping cost", "cost", "charge", "amount", role="number")
    budget  = _col(profile, "budget", "planned cost", role="number")
    route   = _col(profile, "route", "lane", "destination")
    vehicle = _col(profile, "vehicle", "truck", "fleet")
    driver  = _col(profile, "driver", "carrier")
    cargo   = _col(profile, "cargo", "goods type", "commodity")
    ontime  = _col(profile, "on time", "ontime", "on-time")
    damaged = _col(profile, "damaged", "damage", "broken")
    planned_d = _col(profile, "planned delivery", "planned days", "expected days", role="number")
    actual_d  = _col(profile, "actual delivery", "actual days", "delivery days", role="number")
    rating  = _col(profile, "rating", "csat", role="number")
    dist    = _col(profile, "distance", "km", "miles", role="number")
    weight  = _col(profile, "weight", "load", role="number")
    date    = _col(profile, "date", "ship date", "dispatch date", role="date")
    K["Shipments"] = len(df)

    if cost:
        c = _numv(df, cost)
        _money_kpi(K, "Total Cost", float(c.sum())); _money_kpi(K, "Avg Cost/Shipment", float(c.mean()))
        if date:
            try:
                t = df[[date]].copy(); t["_d"] = _to_datetime(df[date]); t["_c"] = c; t = t.dropna(subset=["_d"])
                m = t.groupby(t["_d"].dt.to_period("M"))["_c"].sum().sort_index()
                if len(m) >= 2: S["Monthly delivery cost"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in m.tail(12).items()]
            except Exception: pass
        if route:   S["Cost by route"] = _sumrows(df, route, c, 10)
        if vehicle: S["Cost by vehicle"] = _sumrows(df, vehicle, c, 8)
        if cargo:   S["Cost by cargo type"] = _sumrows(df, cargo, c, 8)
        if dist:
            tk = float(_numv(df, dist).sum())
            if tk: K["Cost per KM"] = round(float(c.sum()) / tk, 2)
        if budget:
            b = float(_numv(df, budget).sum())
            if b:
                var = _pct(float(c.sum()) - b, b); K["Cost Variance %"] = var
                if var > 5: A.append({"type": "Warning", "text": f"Delivery cost is {var}% over budget."})

    if ontime:
        low = df[ontime].astype(str).str.lower()
        ot = int(low.str.contains(r"\byes\b|true|on.?time|^1$", na=False).sum())
        rate = _pct(ot, len(df)); K["On-Time %"] = rate
        if rate < 85: A.append({"type": "Warning", "text": f"On-time delivery only {rate}% — below the 85% benchmark."})
    if planned_d and actual_d:
        delay = (_numv(df, actual_d) - _numv(df, planned_d)).dropna()
        if len(delay): K["Avg Delay (days)"] = round(float(delay.mean()), 1)
    elif actual_d:
        K["Avg Delivery Days"] = round(float(_numv(df, actual_d).mean()), 1)
    if damaged:
        low = df[damaged].astype(str).str.lower()
        dmgm = low.str.contains(r"\byes\b|true|damage|^1$", na=False); dmg = int(dmgm.sum())
        if dmg:
            dr = _pct(dmg, len(df)); K["Damage Rate %"] = dr
            if dr > 3: A.append({"type": "Warning", "text": f"{dmg} damaged shipment(s) ({dr}%) — handling issue."})
            if cargo:
                dc = _countrows(df[dmgm], cargo, 6)
                if dc: S["Damages by cargo type"] = dc
    if rating:
        rt = _numv(df, rating).dropna()
        if len(rt):
            K["Avg Rating"] = round(float(rt.mean()), 1)
            if float(rt.mean()) < 3.5: A.append({"type": "Warning", "text": f"Customer rating {round(float(rt.mean()),1)}/5 — service quality slipping."})
    if weight: K["Avg Weight (KG)"] = round(float(_numv(df, weight).mean()), 1)
    if driver: S["Shipments by driver"] = _countrows(df, driver, 8)
    return K, S, A


def pack_restaurant(df, profile):
    """Restaurant / F&B — revenue, food-cost & margin, wastage, table occupancy,
    channel mix (dine-in vs delivery), menu performance, shift & staff, ratings."""
    K, S, A = {}, {}, []
    rev    = _col(profile, "revenue", "sales", "amount", "total", role="number")
    qty    = _col(profile, "quantity", "qty", "units sold", "covers", role="number")
    sell   = _col(profile, "selling price", "sale price", "price", role="number")
    costp  = _col(profile, "cost price", "food cost", "cost", "cogs", role="number")
    item   = _col(profile, "menu item", "item", "dish", "product")
    cat    = _col(profile, "category", "section")
    channel= _col(profile, "channel", "platform", "source")
    shift  = _col(profile, "shift", "meal", "daypart")
    staff  = _col(profile, "staff", "server", "waiter", "employee")
    occ    = _col(profile, "tables occupied", "occupied", role="number")
    tott   = _col(profile, "total tables", "tables available", role="number")
    waste  = _col(profile, "wastage", "waste", "spoilage", role="number")
    rating = _col(profile, "rating", "review", role="number")
    date   = _col(profile, "date", role="date")
    if not rev:
        return K, S, A
    r = _numv(df, rev); tot = float(r.sum()); n = int(r.notna().sum()) or 1
    _money_kpi(K, "Total Revenue", tot); K["Orders"] = n; _money_kpi(K, "Avg Order Value", tot / n)
    if qty: K["Units Sold"] = int(_numv(df, qty).fillna(0).sum())

    # ── Food cost & margin ──
    if costp:
        ctot = float(_numv(df, costp).sum())
        if sell and qty:
            ctot = float((_numv(df, costp) * _numv(df, qty)).sum())
        if tot and ctot:
            K["Food Cost %"] = _pct(ctot, tot)
            _money_kpi(K, "Gross Profit", tot - ctot); K["Gross Margin %"] = _pct(tot - ctot, tot)
            if _pct(ctot, tot) > 35: A.append({"type": "Warning", "text": f"Food cost is {_pct(ctot,tot)}% of revenue — above the 30–35% target."})

    if date:
        try:
            t = df[[date]].copy(); t["_d"] = _to_datetime(df[date]); t["_r"] = r; t = t.dropna(subset=["_d"])
            m = t.groupby(t["_d"].dt.to_period("M"))["_r"].sum().sort_index()
            if len(m) >= 2: S["Monthly revenue"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in m.tail(12).items()]
        except Exception: pass
    if cat:     S["Revenue by category"] = _sumrows(df, cat, r, 8)
    if item:    S["Top menu items"] = _sumrows(df, item, r, 10)
    if channel:
        S["Revenue by channel"] = _sumrows(df, channel, r, 8)
        tc = S["Revenue by channel"]
        if tc and tc[0]["pct"] >= 50: A.append({"type": "Info", "text": f"{tc[0]['name']} drives {tc[0]['pct']}% of revenue."})
    if shift:   S["Revenue by shift"] = _sumrows(df, shift, r, 6)
    if staff:   S["Top staff by sales"] = _sumrows(df, staff, r, 8)

    # ── Table occupancy ──
    if occ and tott:
        o = float(_numv(df, occ).sum()); tt = float(_numv(df, tott).sum())
        if tt:
            K["Table Occupancy %"] = _pct(o, tt)
            if _pct(o, tt) < 50: A.append({"type": "Warning", "text": f"Table occupancy only {_pct(o,tt)}% — capacity under-used."})

    # ── Wastage ──
    if waste:
        w = _numv(df, waste).fillna(0); K["Wastage Units"] = int(w.sum())
        if item and float(w.sum()) > 0:
            try:
                g = w.groupby(df[item].astype(str)).sum().sort_values(ascending=False)
                g = g[(g > 0) & ~g.index.str.lower().isin(["nan", "none", ""])]
                if len(g): S["Most wasted items"] = [{"name": str(k), "units": int(v)} for k, v in g.head(8).items()]
            except Exception: pass

    if rating:
        rt = _numv(df, rating).dropna()
        if len(rt):
            K["Avg Rating"] = round(float(rt.mean()), 1)
            if float(rt.mean()) < 3.5: A.append({"type": "Warning", "text": f"Avg rating {round(float(rt.mean()),1)}/5 — guest satisfaction is low."})
    return K, S, A


def pack_healthcare(df, profile):
    """Healthcare — revenue & budget, readmission, bed occupancy, payer mix,
    department/doctor performance, patient type, ratings."""
    K, S, A = {}, {}, []
    rev     = _col(profile, "revenue", "bill", "amount", "charge", "total", role="number")
    budget  = _col(profile, "budget", "planned", role="number")
    dept    = _col(profile, "department", "ward", "specialty", "unit")
    doctor  = _col(profile, "doctor", "physician", "consultant")
    diag    = _col(profile, "diagnosis", "condition", "procedure")
    ptype   = _col(profile, "patient type", "admission type", "visit type")
    payer   = _col(profile, "payment type", "payer", "insurance", "payment method")
    city    = _col(profile, "city", "location", "region")
    readmit = _col(profile, "readmit", "readmission")
    beddays = _col(profile, "bed days", "bed-days", "occupied beds", role="number")
    bedsav  = _col(profile, "beds available", "total beds", "capacity", role="number")
    rating  = _col(profile, "rating", "satisfaction", "csat", role="number")
    date    = _col(profile, "date", "admission date", "visit date", role="date")
    K["Patients"] = len(df)

    if rev:
        r = _numv(df, rev); tot = float(r.sum())
        _money_kpi(K, "Total Revenue", tot); _money_kpi(K, "Avg Revenue/Patient", tot / (len(df) or 1))
        if date:
            try:
                t = df[[date]].copy(); t["_d"] = _to_datetime(df[date]); t["_r"] = r; t = t.dropna(subset=["_d"])
                m = t.groupby(t["_d"].dt.to_period("M"))["_r"].sum().sort_index()
                if len(m) >= 2: S["Monthly revenue"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in m.tail(12).items()]
            except Exception: pass
        if dept:   S["Revenue by department"] = _sumrows(df, dept, r, 10)
        if doctor: S["Revenue by doctor"] = _sumrows(df, doctor, r, 10)
        if payer:  S["Revenue by payer"] = _sumrows(df, payer, r, 6)
        if city:   S["Revenue by city"] = _sumrows(df, city, r, 10)
        if budget:
            b = float(_numv(df, budget).sum())
            if b:
                var = _pct(tot - b, b); K["Budget Variance %"] = var
                if var < -10: A.append({"type": "Warning", "text": f"Revenue is {abs(var)}% under budget."})

    if readmit:
        low = df[readmit].astype(str).str.lower()
        rcm = low.str.contains(r"\byes\b|true|^1$|readmit", na=False); rc = int(rcm.sum())
        if rc:
            rr = _pct(rc, len(df)); K["Readmission %"] = rr
            if rr > 10: A.append({"type": "Critical", "text": f"Readmission rate {rr}% — above the 10% quality threshold."})
            if dept:
                rd = _countrows(df[rcm], dept, 6)
                if rd: S["Readmissions by department"] = rd
    if beddays and bedsav:
        bd = float(_numv(df, beddays).sum()); ba = float(_numv(df, bedsav).sum())
        if ba: K["Bed Occupancy %"] = _pct(bd, ba)
    if ptype:  S["Patients by type"] = _countrows(df, ptype, 6)
    if diag:   S["Top diagnoses"] = _countrows(df, diag, 10)
    if rating:
        rt = _numv(df, rating).dropna()
        if len(rt):
            K["Avg Patient Rating"] = round(float(rt.mean()), 1)
            if float(rt.mean()) < 3.5: A.append({"type": "Warning", "text": f"Patient rating {round(float(rt.mean()),1)}/5 — experience needs attention."})
    return K, S, A


def pack_manufacturing(df, profile):
    """Manufacturing — output vs plan (efficiency/OEE), defect rate, downtime &
    causes, production cost & cost/unit, energy, machine/product/shift breakdowns."""
    K, S, A = {}, {}, []
    actual  = _col(profile, "actual output", "actual", "produced", "output", role="number")
    planned = _col(profile, "planned output", "planned", "target", "plan", role="number")
    defect  = _col(profile, "defective", "defect", "reject", "scrap", role="number")
    downtime= _col(profile, "downtime", "stoppage", role="number")
    dtcause = _col(profile, "downtime cause", "cause", "reason")
    dtcost  = _col(profile, "downtime cost", role="number")
    energy  = _col(profile, "energy", "kwh", "power", "consumption", role="number")
    labour  = _col(profile, "labour cost", "labor cost", role="number")
    material= _col(profile, "material cost", "raw material", role="number")
    machine = _col(profile, "machine", "line", "equipment")
    product = _col(profile, "product", "item", "sku")
    shift   = _col(profile, "shift", "crew")
    date    = _col(profile, "date", role="date")

    if actual:
        a = _numv(df, actual); ao = float(a.sum()); K["Total Output"] = round(ao, 0)
        if planned:
            p = float(_numv(df, planned).sum())
            if p:
                eff = _pct(ao, p); K["Efficiency %"] = eff
                if eff < 85: A.append({"type": "Warning", "text": f"Output efficiency {eff}% — below 85% of plan."})
        if date:
            try:
                t = df[[date]].copy(); t["_d"] = _to_datetime(df[date]); t["_a"] = a; t = t.dropna(subset=["_d"])
                m = t.groupby(t["_d"].dt.to_period("M"))["_a"].sum().sort_index()
                if len(m) >= 2: S["Monthly output"] = [{"month": str(p2), "value": round(float(v), 2)} for p2, v in m.tail(12).items()]
            except Exception: pass
        if machine: S["Output by machine"] = [{"name": x["name"], "units": int(x["value"]), "pct": x["pct"]} for x in _sumrows(df, machine, a, 10)]
        if product: S["Output by product"] = [{"name": x["name"], "units": int(x["value"]), "pct": x["pct"]} for x in _sumrows(df, product, a, 10)]
        if defect:
            d = _numv(df, defect); dt = float(d.sum()); K["Total Defects"] = round(dt, 0)
            if ao:
                K["Defect Rate %"] = _pct(dt, ao + dt)
                if _pct(dt, ao + dt) > 5: A.append({"type": "Warning", "text": f"Defect rate {_pct(dt, ao+dt)}% — above 5%."})
            if machine:
                try:
                    g = d.groupby(df[machine].astype(str)).sum().sort_values(ascending=False)
                    g = g[(g > 0) & ~g.index.str.lower().isin(["nan", "none", ""])]
                    if len(g): S["Defects by machine"] = [{"name": str(k), "units": int(v)} for k, v in g.head(8).items()]
                except Exception: pass

    if downtime:
        dh = _numv(df, downtime); K["Downtime (hrs)"] = round(float(dh.sum()), 1)
        if dtcause:
            try:
                g = dh.groupby(df[dtcause].astype(str)).sum().sort_values(ascending=False)
                g = g[(g > 0) & ~g.index.str.lower().isin(["nan", "none", ""])]
                if len(g):
                    S["Downtime by cause"] = [{"name": str(k), "hours": round(float(v), 1)} for k, v in g.head(8).items()]
                    A.append({"type": "Info", "text": f"Top downtime cause: {g.index[0]} ({round(float(g.iloc[0]),1)} hrs)."})
            except Exception: pass
    if dtcost: _money_kpi(K, "Downtime Cost", float(_numv(df, dtcost).sum()))
    if labour or material:
        pc = 0.0
        if labour: pc += float(_numv(df, labour).sum())
        if material: pc += float(_numv(df, material).sum())
        if pc:
            _money_kpi(K, "Production Cost", pc)
            if actual and float(_numv(df, actual).sum()): K["Cost per Unit"] = round(pc / float(_numv(df, actual).sum()), 2)
    if energy: K["Energy (KWH)"] = round(float(_numv(df, energy).sum()), 0)
    return K, S, A


def pack_marketing(df, profile):
    """Marketing — spend, revenue, ROAS, CAC, funnel (impressions→clicks→leads→
    conversions), CTR & conversion rate, channel/campaign/geo performance."""
    K, S, A = {}, {}, []
    spend   = _col(profile, "spend", "cost", "budget", "ad spend", role="number")
    rev     = _col(profile, "revenue generated", "revenue", "sales", "conversion value", role="number")
    impr    = _col(profile, "impressions", "views", "reach", role="number")
    clicks  = _col(profile, "clicks", "click", role="number")
    leads   = _col(profile, "leads", "lead", "signup", role="number")
    conv    = _col(profile, "conversions", "conversion", "purchases", "orders", role="number")
    channel = _col(profile, "channel", "source", "platform", "medium")
    campaign= _col(profile, "campaign", "ad", "initiative")
    city    = _col(profile, "city", "region", "location", "geo")
    date    = _col(profile, "date", role="date")

    if spend:
        sp = _numv(df, spend); spv = float(sp.sum())
        _money_kpi(K, "Total Spend", spv)
        if rev:
            rv = float(_numv(df, rev).sum()); _money_kpi(K, "Revenue Generated", rv)
            if spv:
                K["ROAS"] = round(rv / spv, 2)
                if rv / spv < 1: A.append({"type": "Critical", "text": f"ROAS is {round(rv/spv,2)}x — spending more than you earn."})
                elif rv / spv >= 4: A.append({"type": "Opportunity", "text": f"Strong {round(rv/spv,2)}x ROAS — scale what's working."})
        if date:
            try:
                t = df[[date]].copy(); t["_d"] = _to_datetime(df[date]); t["_s"] = sp; t = t.dropna(subset=["_d"])
                m = t.groupby(t["_d"].dt.to_period("M"))["_s"].sum().sort_index()
                if len(m) >= 2: S["Monthly spend"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in m.tail(12).items()]
            except Exception: pass
        if channel: S["Spend by channel"] = _sumrows(df, channel, sp, 8)

    if conv: K["Conversions"] = int(_numv(df, conv).fillna(0).sum())
    if leads: K["Total Leads"] = int(_numv(df, leads).fillna(0).sum())
    if spend and conv:
        cv = float(_numv(df, conv).fillna(0).sum())
        if cv: _money_kpi(K, "CAC", float(_numv(df, spend).sum()) / cv)
    if clicks and impr:
        ci = float(_numv(df, impr).sum())
        if ci: K["CTR %"] = _pct(float(_numv(df, clicks).sum()), ci)
    if conv and clicks:
        cl = float(_numv(df, clicks).sum())
        if cl: K["Conversion Rate %"] = _pct(float(_numv(df, conv).sum()), cl)

    if rev and channel:
        S["Revenue by channel"] = _sumrows(df, channel, _numv(df, rev), 8)
        # ROAS by channel
        if spend:
            try:
                gs = _numv(df, spend).groupby(df[channel].astype(str)).sum()
                gr = _numv(df, rev).groupby(df[channel].astype(str)).sum()
                roas = (gr / gs.replace(0, np.nan)).dropna().sort_values(ascending=False)
                roas = roas[~roas.index.str.lower().isin(["nan", "none", ""])]
                if len(roas):
                    S["ROAS by channel"] = [{"name": str(k), "roas": round(float(v), 2)} for k, v in roas.head(8).items()]
                    if float(roas.iloc[-1]) < 1: A.append({"type": "Warning", "text": f"{roas.index[-1]} is unprofitable ({round(float(roas.iloc[-1]),2)}x ROAS) — cut or fix it."})
            except Exception: pass
    if campaign: S["Spend by campaign"] = _sumrows(df, campaign, _numv(df, spend), 8) if spend else None
    if conv and channel: S["Conversions by channel"] = [{"name": x["name"], "units": int(x["value"]), "pct": x["pct"]} for x in _sumrows(df, channel, _numv(df, conv), 8)]
    if city and rev: S["Revenue by city"] = _sumrows(df, city, _numv(df, rev), 10)
    S = {k: v for k, v in S.items() if v}
    return K, S, A


def pack_education(df, profile):
    """Education — scores & pass rate, attendance, fee collection & dues, dropout,
    performance by class/subject/teacher, gender split."""
    K, S, A = {}, {}, []
    score   = _col(profile, "avg score", "score", "marks", "percentage", "result", "grade", role="number")
    maxsc   = _col(profile, "max score", "maximum score", "out of", role="number")
    present = _col(profile, "students present", "present", "attended", role="number")
    totstu  = _col(profile, "total students", "enrolled", "strength", "class size", role="number")
    passcol = _col(profile, "pass")
    fee_ch  = _col(profile, "fee charged", "fees charged", "fee due", "billed", role="number")
    fee_pd  = _col(profile, "fee paid", "fees paid", "collected", "received", role="number")
    cls     = _col(profile, "class", "section", "standard", "batch")
    subj    = _col(profile, "subject", "course")
    teacher = _col(profile, "teacher", "faculty", "instructor", "professor")
    gender  = _col(profile, "gender", "sex")
    sstatus = _col(profile, "student status", "status", "enrollment status")

    if score:
        s = _numv(df, score)
        K["Avg Score"] = round(float(s.mean()), 1)
        if maxsc:
            mx = float(_numv(df, maxsc).mean())
            if mx: K["Score %"] = _pct(float(s.mean()), mx)
        if cls:     S["Avg score by class"] = [{"name": x["name"], "score": round(x["value"], 1)} for x in _meanrows(df, cls, s, 12)]
        if subj:    S["Avg score by subject"] = [{"name": x["name"], "score": round(x["value"], 1)} for x in _meanrows(df, subj, s, 12)]
        if teacher: S["Avg score by teacher"] = [{"name": x["name"], "score": round(x["value"], 1)} for x in _meanrows(df, teacher, s, 10)]
    # ── Pass rate (from a pass flag, else score ≥ 40) ──
    if passcol:
        low = df[passcol].astype(str).str.lower(); passed = int(low.str.contains(r"\byes\b|true|pass|^1$", na=False).sum())
        pr = _pct(passed, len(df)); K["Pass Rate %"] = pr
        if pr < 60: A.append({"type": "Warning", "text": f"Pass rate {pr}% — over a third of students failing."})
    elif score:
        s = _numv(df, score); pr = _pct(int((s >= 40).sum()), int(s.notna().sum()) or 1); K["Pass Rate %"] = pr
        if pr < 60: A.append({"type": "Warning", "text": f"Pass rate {pr}% — many below the pass mark."})

    # ── Attendance ──
    if present and totstu:
        p = float(_numv(df, present).sum()); t = float(_numv(df, totstu).sum())
        if t:
            K["Attendance %"] = _pct(p, t)
            if _pct(p, t) < 75: A.append({"type": "Warning", "text": f"Attendance {_pct(p,t)}% — below 75%."})
            if cls:
                try:
                    pr_ = _numv(df, present).groupby(df[cls].astype(str)).sum()
                    tr_ = _numv(df, totstu).groupby(df[cls].astype(str)).sum()
                    at = (pr_ / tr_.replace(0, np.nan) * 100).dropna().sort_values()
                    at = at[~at.index.str.lower().isin(["nan", "none", ""])]
                    if len(at): S["Attendance by class"] = [{"name": str(k), "rate": round(float(v), 1)} for k, v in at.head(12).items()]
                except Exception: pass

    # ── Fees ──
    if fee_pd:
        paid = float(_numv(df, fee_pd).sum()); _money_kpi(K, "Fees Collected", paid)
        if fee_ch:
            ch = float(_numv(df, fee_ch).sum())
            if ch:
                K["Fee Collection %"] = _pct(paid, ch)
                if ch > paid: _money_kpi(K, "Outstanding Fees", ch - paid)
                if _pct(paid, ch) < 80: A.append({"type": "Warning", "text": f"Only {_pct(paid,ch)}% of fees collected — {_fmt(ch-paid)} outstanding."})
        if cls: S["Fees collected by class"] = _sumrows(df, cls, _numv(df, fee_pd), 12)

    # ── Dropout ──
    if sstatus:
        low = df[sstatus].astype(str).str.lower(); dropm = low.str.contains(r"drop|left|inactive|discontinu", na=False); dropped = int(dropm.sum())
        if dropped:
            dr = _pct(dropped, len(df)); K["Dropout %"] = dr
            if dr > 10: A.append({"type": "Critical", "text": f"Dropout rate {dr}% — high student churn."})
            if cls:
                dc = _countrows(df[dropm], cls, 8)
                if dc: S["Dropouts by class"] = dc
    if gender: S["Gender split"] = _countrows(df, gender, 6)
    return K, S, A


def pack_banking(df, profile):
    """Banking / lending — disbursement, outstanding, NPA, collection efficiency,
    overdue, interest & ticket size, branch/agent/product/geo & NPA concentration."""
    K, S, A = {}, {}, []
    disb    = _col(profile, "loan amount", "disbursed", "principal", "sanctioned", "amount", role="number")
    out     = _col(profile, "outstanding", "balance", "due amount", role="number")
    status  = _col(profile, "status", "npa", "default", "loan status")
    branch  = _col(profile, "branch", "region")
    agent   = _col(profile, "agent", "officer", "rm", "relationship manager")
    ltype   = _col(profile, "loan type", "product", "scheme", "loan product")
    state   = _col(profile, "state", "province")
    rate    = _col(profile, "interest rate", "interest", "roi", role="number")
    tenure  = _col(profile, "tenure", "term", "duration", role="number")
    overdue = _col(profile, "days overdue", "overdue", "dpd", role="number")
    coll    = _col(profile, "collection", "recovery", "repaid", role="number")
    date    = _col(profile, "date", "disbursal date", role="date")

    if disb:
        d = _numv(df, disb)
        _money_kpi(K, "Total Disbursed", float(d.sum())); K["Loans"] = int(d.notna().sum())
        _money_kpi(K, "Avg Ticket Size", float(d.mean()))
        if date:
            try:
                t = df[[date]].copy(); t["_d"] = _to_datetime(df[date]); t["_a"] = d; t = t.dropna(subset=["_d"])
                m = t.groupby(t["_d"].dt.to_period("M"))["_a"].sum().sort_index()
                if len(m) >= 2: S["Monthly disbursement"] = [{"month": str(p), "value": round(float(v), 2)} for p, v in m.tail(12).items()]
            except Exception: pass
        if branch: S["Disbursed by branch"] = _sumrows(df, branch, d, 10)
        if ltype:  S["Disbursed by loan type"] = _sumrows(df, ltype, d, 10)
        if state:  S["Disbursed by state"] = _sumrows(df, state, d, 10)
        if agent:  S["Top agents by disbursal"] = _sumrows(df, agent, d, 8)
    if out: _money_kpi(K, "Outstanding", float(_numv(df, out).sum()))
    if rate: K["Avg Interest Rate %"] = round(float(_numv(df, rate).mean()), 2)
    if tenure: K["Avg Tenure (mo)"] = round(float(_numv(df, tenure).mean()), 1)
    if coll:
        cv = _numv(df, coll).dropna()
        if len(cv):
            K["Avg Collection %"] = round(float(cv.mean()), 1)
            if float(cv.mean()) < 80: A.append({"type": "Warning", "text": f"Collection efficiency {round(float(cv.mean()),1)}% — recovery lagging."})
    if overdue:
        od = _numv(df, overdue).fillna(0); odn = int((od > 0).sum())
        if odn:
            K["Overdue Loans"] = odn
            if _pct(odn, len(df)) > 15: A.append({"type": "Warning", "text": f"{odn} loans overdue ({_pct(odn,len(df))}%) — collection risk."})

    if status:
        low = df[status].astype(str).str.lower()
        npam = low.str.contains(r"npa|default|overdue|bad|non.?perform", na=False); npa = int(npam.sum())
        if npa:
            nr = _pct(npa, len(df)); K["NPA %"] = nr
            if out: _money_kpi(K, "NPA Amount", float(_numv(df, out)[npam].sum()))
            if nr > 5: A.append({"type": "Critical", "text": f"NPA rate {nr}% — above the 5% red line."})
            if branch:
                nb = _countrows(df[npam], branch, 6)
                if nb: S["NPA by branch"] = nb
            if ltype:
                nl = _countrows(df[npam], ltype, 6)
                if nl:
                    S["NPA by loan type"] = nl
                    A.append({"type": "Warning", "text": f"{nl[0]['name']} has the most NPAs ({nl[0]['count']} loans)."})
    return K, S, A


def pack_realestate(df, profile):
    K, S, A = {}, {}, []
    price = _col(profile, "price", "asking", "value", "amount", role="number")
    area = _col(profile, "area", "sqft", "size", "carpet", role="number")
    city = _col(profile, "city", "location", "area", "locality")
    status = _col(profile, "status", "availability")
    agent = _col(profile, "agent", "broker")
    if price:
        p = _numv(df, price); _money_kpi(K, "Avg Asking Price", float(p.mean()))
        K["Listings"] = int(p.notna().sum())
        if city: S["Avg price by city"] = _meanrows(df, city, p, 10)
        if agent: S["Listings by agent"] = _sumrows(df, agent, pd.Series(1, index=df.index), 10)
        if area:
            ar = _numv(df, area)
            pps = (p / ar).replace([np.inf, -np.inf], np.nan)
            if pps.notna().any(): _money_kpi(K, "Avg Price / sqft", float(pps.mean()))
    if status:
        low = df[status].astype(str).str.lower()
        sold = int(low.str.contains("sold", na=False).sum())
        avail = int(low.str.contains("avail|open", na=False).sum())
        if sold or avail:
            K["Sold"] = sold; K["Available"] = avail
    return K, S, A


def pack_hospitality(df, profile):
    K, S, A = {}, {}, []
    rev = _col(profile, "revenue", "amount", "rate", "total", role="number")
    room = _col(profile, "room", "type", "category")
    source = _col(profile, "source", "channel", "booking")
    if rev:
        r = _numv(df, rev); _money_kpi(K, "Total Revenue", float(r.sum())); K["Bookings"] = int(r.notna().sum())
        if room: S["Revenue by room type"] = _sumrows(df, room, r, 8)
        if source: S["Revenue by source"] = _sumrows(df, source, r, 8)
    return K, S, A


MODULE_PACKS = [
    ("sales", pack_sales), ("inventory", pack_inventory), ("hr", pack_hr), ("payroll", pack_hr),
    ("finance", pack_finance), ("accounting", pack_finance), ("retail", pack_retail), ("e-commerce", pack_retail),
    ("logistics", pack_logistics), ("restaurant", pack_restaurant), ("healthcare", pack_healthcare),
    ("manufacturing", pack_manufacturing), ("marketing", pack_marketing), ("education", pack_education),
    ("banking", pack_banking), ("real estate", pack_realestate), ("hospitality", pack_hospitality),
]


def apply_module_pack(module: str, df: pd.DataFrame, profile: list[dict], result: dict) -> str:
    """Layer the matching industry pack on top of the universal result. Returns the
    pack name used (or "")."""
    m = (module or "").lower()
    pack = next((fn for key, fn in MODULE_PACKS if key in m), None)
    if not pack:
        return ""
    try:
        K, S, A = pack(df, profile)
    except Exception:
        return ""
    if K:  # domain KPIs are curated & meaningful → drop the generic auto-summed totals
        base = result.get("kpis", {})
        kept = {k: v for k, v in base.items()
                if k == "rows" or not (k.endswith("(total)") or k.endswith("(total)_fmt"))}
        result["kpis"] = {**K, **{k: v for k, v in kept.items() if k not in K}}
    # domain breakdowns inserted right after kpis, before the generic "X by Y"
    if S:
        rebuilt = {"kpis": result.get("kpis", {})}
        for nm, rows in S.items():
            if rows:
                rebuilt[nm] = rows
        # keep meta (alerts, columns_detected) but DROP the universal generic
        # breakdowns/trend — the pack's curated sections replace them (no duplicates)
        for k, v in result.items():
            if k == "kpis":
                continue
            if isinstance(v, list) and k != "alerts":
                continue
            rebuilt[k] = v
        result.clear(); result.update(rebuilt)
    if A:
        # drop the low-value generic "X is the largest in Y" alerts — the pack's
        # curated sections already convey this, so they'd just be noise.
        base_alerts = [al for al in result.get("alerts", []) if "is the largest in" not in al.get("text", "")]
        result["alerts"] = A + base_alerts
    return pack.__name__.replace("pack_", "")


# ════════════════════════════════════════════════════════════════════════════
# DATA HEALTH SCORE — a "credit score" for the sheet (0–100) + an audit of what's
# wrong and what we handled. This is what the speedometer gauge displays.
# ════════════════════════════════════════════════════════════════════════════

def compute_health(df: pd.DataFrame, profile: list[dict], fixes: list[str]) -> dict:
    issues = []
    n_rows = len(df)
    n_cols = len(df.columns)
    cells = max(n_rows * n_cols, 1)
    score = 100.0

    # 1. Blank cells — biggest quality signal
    total_missing = int(df.isna().sum().sum())
    missing_ratio = total_missing / cells
    if total_missing:
        score -= min(missing_ratio * 100 * 0.8, 35)
        sev = "high" if missing_ratio > 0.20 else "medium" if missing_ratio > 0.05 else "low"
        issues.append({"label": f"{total_missing:,} blank cells ({round(missing_ratio*100)}% of data)",
                       "severity": sev,
                       "detail": "Kept blank, never invented — totals stay honest."})

    # 2. Mostly-empty columns
    low_cols = [p["column"] for p in profile if n_rows and p["missing"] > 0.4 * n_rows]
    if low_cols:
        score -= min(len(low_cols) * 4, 16)
        issues.append({"label": f"{len(low_cols)} mostly-empty column(s)",
                       "severity": "medium", "detail": ", ".join(map(str, low_cols[:4]))})

    # 3. Tiny dataset
    if n_rows < 5:
        score -= 10
        issues.append({"label": "Very few rows", "severity": "medium", "detail": f"Only {n_rows} rows of data."})

    # 4. No numbers to analyse
    if not any(p["role"] == "number" for p in profile):
        score -= 10
        issues.append({"label": "No numeric columns detected", "severity": "low",
                       "detail": "Limited metrics possible — mostly text data."})

    # 5. Things we already handled during cleaning (shown for transparency)
    handled = [{"label": fx, "severity": "fixed", "detail": ""} for fx in fixes]

    score = int(max(0, min(100, round(score))))
    rating = ("Excellent" if score >= 85 else "Good" if score >= 70
              else "Fair" if score >= 50 else "Poor")
    return {
        "score": score,
        "rating": rating,
        "issues": issues + handled,
        "problems_found": len(issues),
        "auto_handled": len(handled),
        "summary": f"{len(issues)} issue(s) need attention · {len(handled)} auto-handled",
    }


def build_metadata(df: pd.DataFrame, profile: list[dict]) -> list[dict]:
    """A column-by-column data dictionary for the Metadata tool: type, uniqueness,
    completeness, and a sample value / range for each column."""
    n = max(len(df), 1)
    meta = []
    for p in profile:
        col, role = p["column"], p["role"]
        sample = ""
        try:
            if role == "number":
                sv = pd.to_numeric(df[col], errors="coerce").dropna()
                if len(sv):
                    sample = f"{round(float(sv.min()), 2):g} – {round(float(sv.max()), 2):g}"
            elif role == "date":
                d = _to_datetime(df[col])
                if d is not None and d.notna().any():
                    sample = f"{d.min().date()} → {d.max().date()}"
            else:
                nn = df[col].dropna()
                if len(nn):
                    sample = str(nn.iloc[0])[:40]
        except Exception:
            pass
        meta.append({
            "column": str(col),
            "type": role,
            "distinct": p["distinct"],
            "missing": p["missing"],
            "fill_pct": round((n - p["missing"]) / n * 100),
            "sample": sample,
        })
    return meta


# ════════════════════════════════════════════════════════════════════════════
# FOOTPRINT / PRIVACY SCANNER — find & remove everything hidden in a file before
# it's shared: author, company, hidden sheets, comments, external links, macros,
# and a scan for personal data (emails/phones/IDs). The differentiator.
# ════════════════════════════════════════════════════════════════════════════

# Value patterns — EU/US first (GDPR), India kept as secondary
_PII = {
    "Email":       re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+"),
    "Phone":       re.compile(r"(?<!\d)(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{2,4}\)?[\s.-]?){2,4}\d{2,3}(?!\d)"),
    "IBAN":        re.compile(r"\b[A-Z]{2}\d{2}[A-Z0-9]{10,30}\b"),
    "Credit card": re.compile(r"(?<!\d)(?:\d[ -]?){15,16}\d(?!\d)"),
    "US SSN":      re.compile(r"\b\d{3}-\d{2}-\d{4}\b"),
    "UK NINO":     re.compile(r"\b[A-CEGHJ-PR-TW-Z]{2}\d{6}[A-D]\b"),
    "IP address":  re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b"),
    "PAN (India)": re.compile(r"\b[A-Z]{5}[0-9]{4}[A-Z]\b"),
    "Aadhaar":     re.compile(r"(?<!\d)\d{4}\s?\d{4}\s?\d{4}(?!\d)"),
}


# Column-name patterns — catches PII even when values are numeric/clean
_PII_NAME = {
    "Email":         re.compile(r"e-?mail", re.I),
    "Phone":         re.compile(r"phone|mobile|\btel\b|telefon|t[ée]l[ée]phone|whatsapp", re.I),
    "Name":          re.compile(r"(first|last|full|sur|customer|client|employee|contact|patient|user)[\s_]*name|\bsurname\b|\bvorname\b|\bnachname\b", re.I),
    "Address":       re.compile(r"\baddress\b|\badresse\b|street|post[\s_]?code|\bzip\b|\bcity\b", re.I),
    "Date of birth": re.compile(r"\b(dob|date of birth|birth ?date|geburtsdatum)\b", re.I),
    "IBAN":          re.compile(r"\biban\b|account number|bank account", re.I),
    "VAT":           re.compile(r"\bvat\b|ust-?id|\btva\b", re.I),
    "US SSN":        re.compile(r"\bssn\b|social security", re.I),
    "UK NINO":       re.compile(r"national insurance|\bnino\b|ni number", re.I),
    "Passport":      re.compile(r"passport", re.I),
    "PAN (India)":   re.compile(r"\bpan\b", re.I),
    "Aadhaar":       re.compile(r"aadh?aar|uidai", re.I),
}


def scan_pii(df: pd.DataFrame) -> dict:
    """Detect personal data by (a) column NAME (catches numeric phone columns) and
    (b) value PATTERN on text columns — requiring a real share of matches so
    financial/batch numbers don't false-trigger."""
    found: dict = {}
    try:
        # (a) name-based
        for label, nrx in _PII_NAME.items():
            for c in df.columns:
                if nrx.search(str(c)):
                    found.setdefault(label, []).append(str(c))
        # (b) pattern-based on text columns only. Skip ambiguous patterns
        # (phone/IP match IBANs, VATs, dates) — those rely on column-name detection.
        obj = df.select_dtypes(include="object")
        for label, rx in _PII.items():
            if label in ("Phone", "IP address"):
                continue
            for c in obj.columns:
                if str(c) in found.get(label, []):
                    continue
                vals = obj[c].dropna().astype(str).head(400)
                if len(vals) < 3:
                    continue
                hits = sum(1 for v in vals if rx.search(v))
                if hits >= 3 and hits / len(vals) >= 0.3:
                    found.setdefault(label, []).append(str(c))
        found = {k: list(dict.fromkeys(v))[:5] for k, v in found.items() if v}
    except Exception:
        pass
    return found


def scan_footprint(contents: bytes, name: str, df: pd.DataFrame | None = None) -> dict:
    rep = {"applicable": name.endswith((".xlsx", ".xlsm")), "leaked": {}, "hidden_sheets": [],
           "comments": 0, "hidden_cols": 0, "hidden_rows": 0, "external_links": [],
           "defined_names": 0, "macros": name.endswith(".xlsm"), "pii": {}, "risks": [], "score": 100}
    if df is not None:
        rep["pii"] = scan_pii(df)

    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=False, keep_links=True)
            p = wb.properties
            for a in ["creator", "lastModifiedBy", "title", "subject", "keywords", "description", "category", "manager", "company"]:
                v = getattr(p, a, None)
                if v: rep["leaked"][a] = str(v)
            if getattr(p, "created", None):  rep["leaked"]["created"]  = str(p.created)
            if getattr(p, "modified", None): rep["leaked"]["modified"] = str(p.modified)
            rep["hidden_sheets"] = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
            cells = 0
            for ws in wb.worksheets:
                try:
                    rep["hidden_cols"] += sum(1 for d in ws.column_dimensions.values() if d.hidden)
                    rep["hidden_rows"] += sum(1 for d in ws.row_dimensions.values() if d.hidden)
                    for row in ws.iter_rows():
                        for cell in row:
                            cells += 1
                            if cell.comment: rep["comments"] += 1
                        if cells > 60000: break
                    if cells > 60000: break
                except Exception:
                    continue
            try: rep["defined_names"] = len(list(wb.defined_names))
            except Exception: pass
            try: rep["external_links"] = [str(l.file_link.target) for l in (wb._external_links or [])][:10]
            except Exception: pass
        except Exception:
            rep["applicable"] = False

    # Privacy score (100 = safe) + plain-language risks
    s, R = 100, rep["risks"]
    ids = [rep["leaked"][k] for k in ("creator", "lastModifiedBy", "company", "manager") if k in rep["leaked"]]
    if ids:
        s -= 15; R.append({"label": f"Reveals identity: {', '.join(ids)[:70]}", "severity": "high",
                           "detail": "Author / editor / company name is embedded in the file."})
    if rep["hidden_sheets"]:
        s -= 20; R.append({"label": f"{len(rep['hidden_sheets'])} hidden sheet(s): {', '.join(rep['hidden_sheets'][:4])}",
                           "severity": "high", "detail": "Hidden sheets often hold internal data (costs, margins)."})
    if rep["hidden_cols"] or rep["hidden_rows"]:
        s -= 10; R.append({"label": f"{rep['hidden_cols']} hidden column(s), {rep['hidden_rows']} hidden row(s)",
                           "severity": "medium", "detail": "Hidden cells still travel inside the file."})
    if rep["comments"]:
        s -= 10; R.append({"label": f"{rep['comments']} cell comment(s) / note(s)", "severity": "medium",
                           "detail": "Internal notes may be private."})
    if rep["external_links"]:
        s -= 10; R.append({"label": f"{len(rep['external_links'])} external link(s)", "severity": "medium",
                           "detail": "Links can leak file paths or other workbooks."})
    if rep["macros"]:
        s -= 10; R.append({"label": "Contains macros (VBA code)", "severity": "medium", "detail": "Macros can carry executable code."})
    if rep["pii"]:
        s -= 15; R.append({"label": f"Possible personal data: {', '.join(rep['pii'].keys())}", "severity": "high",
                           "detail": "Emails / phones / ID numbers detected in the cells."})
    rep["score"] = max(0, min(100, s))
    if not R:
        R.append({"label": "No hidden footprint found — safe to share", "severity": "ok", "detail": ""})

    # ── GDPR "safe to share" verdict (plain English) ──
    pii = rep.get("pii") or {}
    direct = [k for k in pii if k in ("Name", "Email", "Phone", "Address", "Date of birth",
                                      "IBAN", "Aadhaar", "PAN (India)", "US SSN", "UK NINO", "Credit card", "VAT")]
    if direct:
        rep["gdpr"] = {
            "personal_data": True, "level": "high",
            "verdict": "This file contains personal data — under GDPR, anonymise or get consent before sharing it.",
            "categories": list(pii.keys()),
            "advice": "Remove or mask the personal columns below before you send this file externally.",
        }
    elif pii:
        rep["gdpr"] = {
            "personal_data": True, "level": "medium",
            "verdict": "This file may contain personal data — review before sharing.",
            "categories": list(pii.keys()),
            "advice": "Check the flagged columns and remove anything that identifies a person.",
        }
    else:
        rep["gdpr"] = {
            "personal_data": False, "level": "ok",
            "verdict": "No obvious personal data detected — likely safe to share.",
            "categories": [], "advice": "Still review hidden sheets/author info below before sending.",
        }
    return rep


# ════════════════════════════════════════════════════════════════════════════
# IMAGE / PHOTO METADATA — photos secretly carry GPS location, camera, owner,
# date. Same "find → show → clean" model as spreadsheets, just an EXIF reader.
# ════════════════════════════════════════════════════════════════════════════
_IMG_EXT = (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".bmp", ".heic", ".heif")

def _register_heif():
    try:
        import pillow_heif; pillow_heif.register_heif_opener()
    except Exception:
        pass

def _gps_dec(v, ref):
    try:
        d = float(v[0]); m = float(v[1]); s = float(v[2])
        dec = d + m / 60 + s / 3600
        r = (ref.decode() if isinstance(ref, bytes) else str(ref)).strip()
        return -dec if r in ("S", "W") else dec
    except Exception:
        return None

def scan_image(contents: bytes, name: str) -> dict:
    """Read what a photo secretly reveals — GPS location, camera, owner, date.
    Uses Pillow's universal EXIF reader so JPEG, HEIC (iPhone), PNG, TIFF all work.
    Each item is tagged with a category so the client can pick what to remove."""
    rep = {"applicable": True, "kind": "image", "leaked": {}, "maps": "", "risks": [],
           "pii": {}, "score": 100, "readable": True}
    try:
        from PIL import Image, ExifTags
        _register_heif()
        img = Image.open(io.BytesIO(contents))
        exif = img.getexif()
        def txt(v):
            if isinstance(v, bytes): v = v.decode(errors="ignore")
            return str(v).strip("\x00 ").strip()
        make = exif.get(271); model = exif.get(272); sw = exif.get(305)
        art = exif.get(315); cpy = exif.get(33432); dt0 = exif.get(306)
        try: sub = exif.get_ifd(ExifTags.IFD.Exif)
        except Exception: sub = {}
        dto = sub.get(36867) if sub else None
        if make or model: rep["leaked"]["Camera"] = (txt(make or "") + " " + txt(model or "")).strip()
        if sw:  rep["leaked"]["Software"] = txt(sw)
        if art: rep["leaked"]["Author / owner"] = txt(art)
        if cpy: rep["leaked"]["Copyright"] = txt(cpy)
        dv = dto or dt0
        if dv: rep["leaked"]["Date taken"] = txt(dv)
        try: gps = exif.get_ifd(ExifTags.IFD.GPSInfo)
        except Exception: gps = {}
        if gps and gps.get(2) and gps.get(4):
            lat = _gps_dec(gps.get(2), gps.get(1, "N")); lon = _gps_dec(gps.get(4), gps.get(3, "E"))
            if lat is not None and lon is not None:
                rep["leaked"]["GPS location"] = f"{lat:.5f}, {lon:.5f}"
                rep["maps"] = f"https://maps.google.com/?q={lat},{lon}"

        # ── Advanced details (beyond what a phone shows) ──
        det = {}
        rat = lambda v: round(float(v), 2)
        try:
            if sub.get(42036): det["Lens"] = txt(sub.get(42036))
            bsn = sub.get(42033) or exif.get(42033)
            if bsn: det["Camera serial no."] = txt(bsn)
            if sub.get(42037): det["Lens serial no."] = txt(sub.get(42037))
            if sub.get(33437): det["Aperture"] = f"f/{rat(sub.get(33437))}"
            iso = sub.get(34855)
            if iso: det["ISO"] = int(iso[0]) if isinstance(iso, (tuple, list)) else int(iso)
            exp = sub.get(33434)
            if exp:
                ev = float(exp); det["Shutter"] = (f"1/{round(1/ev)}s" if 0 < ev < 1 else f"{rat(exp)}s")
            if sub.get(37386): det["Focal length"] = f"{rat(sub.get(37386))}mm"
            fl = sub.get(37385)
            if fl is not None: det["Flash"] = "On" if (int(fl) & 1) else "Off"
            if gps:
                alt = gps.get(6)
                if alt is not None:
                    det["Altitude"] = f"{round(float(alt))} m" + (" below sea level" if gps.get(5) in (1, b"\x01") else "")
                d = gps.get(17)
                if d is not None:
                    dv = float(d); dirs = ["N", "NE", "E", "SE", "S", "SW", "W", "NW"]
                    det["Facing"] = f"{dirs[int((dv + 22.5) % 360 // 45)]} ({round(dv)}°)"
        except Exception:
            pass
        rep["details"] = det
        # embedded thumbnail — often the ORIGINAL before crop/edit (advanced leak)
        try:
            import piexif, base64
            pex = piexif.load(contents)
            if pex.get("thumbnail"):
                rep["thumbnail"] = "data:image/jpeg;base64," + base64.b64encode(pex["thumbnail"]).decode()
        except Exception:
            pass
    except Exception:
        rep["readable"] = False

    s, R = 100, rep["risks"]
    if "GPS location" in rep["leaked"]:
        s -= 35; R.append({"label": f"📍 GPS location embedded: {rep['leaked']['GPS location']}", "severity": "high",
                           "cat": "gps", "detail": "This reveals exactly where the photo was taken — a real safety risk."})
    if "Author / owner" in rep["leaked"] or "Copyright" in rep["leaked"]:
        s -= 15; R.append({"label": f"Owner identity: {rep['leaked'].get('Author / owner') or rep['leaked'].get('Copyright')}", "severity": "high",
                           "cat": "owner", "detail": "Your name / copyright is stored inside the photo."})
    if "Camera" in rep["leaked"] or "Software" in rep["leaked"]:
        s -= 10; R.append({"label": f"Device info: {rep['leaked'].get('Camera','')} {rep['leaked'].get('Software','')}".strip(), "severity": "medium",
                           "cat": "camera", "detail": "The exact device and software are recorded."})
    if "Date taken" in rep["leaked"]:
        s -= 5; R.append({"label": f"Date/time taken: {rep['leaked']['Date taken']}", "severity": "low",
                          "cat": "date", "detail": "When the photo was captured."})
    _det = rep.get("details") or {}
    if _det.get("Camera serial no.") or _det.get("Lens serial no."):
        s -= 15; R.append({"label": f"Device serial number: {_det.get('Camera serial no.') or _det.get('Lens serial no.')}", "severity": "high",
                           "cat": "camera", "detail": "A serial number uniquely fingerprints your exact device across EVERY photo you take."})
    if rep.get("thumbnail"):
        s -= 15; R.append({"label": "Hidden thumbnail of the original embedded", "severity": "high",
                           "cat": "thumbnail", "detail": "A cropped/edited photo can still secretly contain the full original image as a thumbnail."})
    rep["score"] = max(0, min(100, s))
    rep["categories_present"] = sorted({r.get("cat") for r in R if r.get("cat") and r.get("cat") != "thumbnail"})
    if not R:
        msg = ("No hidden metadata found — this photo is clean (apps like WhatsApp/Instagram strip it on send)."
               if rep["readable"] else "Couldn't read this image format — try a JPG/PNG/HEIC original.")
        R.append({"label": msg, "severity": "ok", "detail": ""})
    has_loc = "GPS location" in rep["leaked"]; has_id = "Author / owner" in rep["leaked"] or "Copyright" in rep["leaked"]
    rep["gdpr"] = ({"personal_data": True, "level": "high",
                    "verdict": "This photo reveals personal data — strip it before posting or sharing.",
                    "categories": [k for k in ("GPS location", "Author / owner", "Date taken") if k in rep["leaked"]],
                    "advice": "Tick what to remove below, then clean & download — the photo looks identical."}
                   if (has_loc or has_id) else
                   {"personal_data": False, "level": "ok" if not rep["leaked"] else "medium",
                    "verdict": "No location or identity found." if not rep["leaked"] else "Minor device metadata only — low risk.",
                    "categories": list(rep["leaked"].keys()),
                    "advice": "You can still strip the remaining metadata below."})

    # ── "What this reveals about you" narrative ──
    parts = []
    if rep["leaked"].get("Camera"): parts.append(f"on a {rep['leaked']['Camera']}")
    if rep["leaked"].get("Date taken"): parts.append(f"on {rep['leaked']['Date taken']}")
    if rep["leaked"].get("GPS location"): parts.append(f"at {rep['leaked']['GPS location']}")
    if _det.get("Facing"): parts.append(f"facing {_det['Facing']}")
    if _det.get("Altitude"): parts.append(f"{_det['Altitude']} altitude")
    if parts:
        rep["story"] = "Anyone with this photo can see it was taken " + ", ".join(parts) + "."
    return rep

def strip_image(contents: bytes, name: str = "", categories=None) -> tuple[bytes, str]:
    """Return a visually-identical copy with the SELECTED metadata removed.
    categories: list subset of {gps,camera,owner,date}; None/empty = remove all."""
    from PIL import Image
    nm = (name or "").lower()
    if nm.endswith((".heic", ".heif")):
        _register_heif()
    img = Image.open(io.BytesIO(contents))
    fmt = (img.format or "JPEG").upper()
    is_jpeg = nm.endswith((".jpg", ".jpeg")) or fmt in ("JPEG", "JPG", "MPO")
    ALL = {"gps", "camera", "owner", "date"}
    cats = ALL if not categories else (set(categories) & ALL or ALL)

    # JPEG + a partial selection → edit EXIF in place (keeps the rest, no re-encode)
    if is_jpeg and cats != ALL:
        try:
            import piexif
            ex = piexif.load(contents)
            if "gps" in cats: ex["GPS"] = {}
            if "camera" in cats:
                for t in (piexif.ImageIFD.Make, piexif.ImageIFD.Model, piexif.ImageIFD.Software):
                    ex["0th"].pop(t, None)
            if "owner" in cats:
                for t in (piexif.ImageIFD.Artist, piexif.ImageIFD.Copyright):
                    ex["0th"].pop(t, None)
            if "date" in cats:
                ex["Exif"].pop(piexif.ExifIFD.DateTimeOriginal, None)
                ex["0th"].pop(piexif.ImageIFD.DateTime, None)
            # always drop the embedded thumbnail (it can be the uncropped original)
            ex["1st"] = {}; ex["thumbnail"] = None
            # also drop serial numbers when removing camera/device info
            if "camera" in cats:
                for t in (0xA431, 0xA435):  # BodySerialNumber, LensSerialNumber
                    ex["Exif"].pop(t, None)
            nb = piexif.dump(ex)
            out = io.BytesIO(); piexif.insert(nb, contents, out); out.seek(0)
            return out.read(), "JPEG"
        except Exception:
            pass

    # Full strip — re-encode with no metadata at all
    save_fmt = "JPEG" if (is_jpeg or nm.endswith((".heic", ".heif"))) else fmt
    clean = Image.new(img.mode, img.size)
    clean.putdata(list(img.getdata()))
    if save_fmt == "JPEG" and clean.mode in ("RGBA", "P", "LA"):
        clean = clean.convert("RGB")
    out = io.BytesIO(); clean.save(out, format=save_fmt); out.seek(0)
    return out.read(), save_fmt


# ════════════════════════════════════════════════════════════════════════════
# PDF METADATA — author, producer/software, dates + hidden content (JavaScript,
# attachments, forms, annotations) and the big one: REDACTION FAILURES (text
# "blacked out" but still readable underneath).
# ════════════════════════════════════════════════════════════════════════════
_PDF_EXT = (".pdf",)

def scan_pdf(contents: bytes, name: str) -> dict:
    rep = {"applicable": True, "kind": "pdf", "leaked": {}, "risks": [], "pii": {},
           "score": 100, "details": {}, "readable": True, "_flags": {}}
    try:
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(contents))
        meta = r.metadata or {}
        def m(k):
            v = meta.get(k); return str(v).strip() if v else ""
        if m("/Author"):   rep["leaked"]["Author"] = m("/Author")
        if m("/Creator"):  rep["leaked"]["Created with"] = m("/Creator")
        if m("/Producer"): rep["leaked"]["Producer (software)"] = m("/Producer")
        if m("/Title"):    rep["leaked"]["Title"] = m("/Title")
        if m("/Subject"):  rep["leaked"]["Subject"] = m("/Subject")
        if m("/Keywords"): rep["leaked"]["Keywords"] = m("/Keywords")
        if meta.get("/CreationDate"): rep["leaked"]["Created"] = str(meta.get("/CreationDate"))
        if meta.get("/ModDate"):      rep["leaked"]["Modified"] = str(meta.get("/ModDate"))

        det = rep["details"]; det["Pages"] = len(r.pages)
        has_js = False
        try:
            root = r.trailer["/Root"]
            names = root.get("/Names")
            if names and "/JavaScript" in names: has_js = True
            if "/OpenAction" in root or "/AA" in root: has_js = True
        except Exception:
            pass
        attach = 0
        try:
            ef = (r.trailer["/Root"].get("/Names") or {}).get("/EmbeddedFiles")
            if ef and ef.get("/Names"): attach = len(ef["/Names"]) // 2
        except Exception:
            pass
        forms = 0
        try:
            acro = r.trailer["/Root"].get("/AcroForm")
            if acro and acro.get("/Fields"): forms = len(acro["/Fields"])
        except Exception:
            pass
        annots = redact = blackbox = 0
        for p in r.pages:
            try:
                for a in (p.get("/Annots") or []):
                    o = a.get_object(); annots += 1
                    st = str(o.get("/Subtype"))
                    if st == "/Redact": redact += 1
                    elif st in ("/Square", "/Redaction"):
                        ic = o.get("/IC")
                        if ic and all(float(x) < 0.25 for x in ic): blackbox += 1
            except Exception:
                continue
        text = ""
        try:
            text = "".join((pg.extract_text() or "") for pg in r.pages[:5])
        except Exception:
            pass
        det["Has JavaScript"] = "Yes" if has_js else "No"
        if attach: det["Attachments"] = attach
        if forms:  det["Form fields"] = forms
        if annots: det["Annotations / comments"] = annots
        det["Text copyable"] = "Yes" if len(text.strip()) > 20 else "No (scanned/secured)"
        rep["_flags"] = {"js": has_js, "attach": attach, "forms": forms, "annots": annots,
                         "redact": redact, "blackbox": blackbox, "has_text": len(text.strip()) > 20}
    except Exception:
        rep["readable"] = False

    s, R = 100, rep["risks"]; f = rep["_flags"]
    ids = [rep["leaked"][k] for k in ("Author", "Created with", "Producer (software)") if k in rep["leaked"]]
    if ids:
        s -= 15; R.append({"label": f"Reveals identity / software: {', '.join(ids)[:80]}", "severity": "high", "cat": "metadata",
                           "detail": "Author and the software used are embedded — can leak names, usernames, internal tools."})
    if f.get("redact"):
        s -= 40; R.append({"label": f"{f['redact']} UNAPPLIED redaction(s) — blacked-out text is still readable", "severity": "high", "cat": "redaction",
                           "detail": "Redactions were marked but never applied — the hidden text can be copied out."})
    if f.get("blackbox") and f.get("has_text"):
        s -= 30; R.append({"label": "Black boxes over text, but the text is still copyable underneath", "severity": "high", "cat": "redaction",
                           "detail": "Drawing a black box does NOT remove the text — it's still in the file."})
    if f.get("js"):
        s -= 20; R.append({"label": "Contains JavaScript / auto-actions", "severity": "high", "cat": "js",
                           "detail": "Embedded scripts can run when the PDF is opened."})
    if f.get("attach"):
        s -= 10; R.append({"label": f"{f['attach']} embedded file attachment(s)", "severity": "medium", "cat": "attachments",
                           "detail": "Files hidden inside the PDF travel with it."})
    if f.get("annots"):
        s -= 8; R.append({"label": f"{f['annots']} annotation(s) / comment(s)", "severity": "medium", "cat": "annotations",
                          "detail": "Notes and markups may contain private remarks."})
    if f.get("forms"):
        s -= 5; R.append({"label": f"{f['forms']} form field(s) with possible saved data", "severity": "low", "cat": "metadata",
                          "detail": "Form fields can hold previously entered data."})
    rep["score"] = max(0, min(100, s))
    rep["categories_present"] = sorted({r2.get("cat") for r2 in R if r2.get("cat")})
    if not R:
        R.append({"label": "No hidden data found — this PDF looks safe to share" if rep["readable"]
                  else "Couldn't read this PDF (it may be encrypted).", "severity": "ok", "detail": ""})
    has_redact = bool(f.get("redact") or (f.get("blackbox") and f.get("has_text")))
    rep["gdpr"] = ({"personal_data": True, "level": "high",
                    "verdict": "This PDF may expose data that was meant to be hidden — review before sharing.",
                    "categories": [k for k in ("Author",) if k in rep["leaked"]] + (["Failed redaction"] if has_redact else []),
                    "advice": "Clean it below to strip metadata, scripts and attachments. NOTE: failed-redaction text can't be auto-fixed — re-do the redaction in your PDF editor."}
                   if (ids or has_redact or f.get("js")) else
                   {"personal_data": bool(ids), "level": "ok" if not (ids or rep["leaked"]) else "medium",
                    "verdict": "No major exposure found." if not rep["leaked"] else "Only document info found — low risk.",
                    "categories": list(rep["leaked"].keys()), "advice": "You can still strip the metadata below."})
    return rep

def clean_pdf(contents: bytes) -> bytes:
    """Strip metadata + drop document-level JavaScript/attachments/forms (a fresh
    writer doesn't carry them) + remove page annotations. NOTE: cannot fix amateur
    black-box redaction — that text lives in the page content itself."""
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(io.BytesIO(contents)); w = PdfWriter()
    for p in r.pages:
        try:
            if "/Annots" in p: del p["/Annots"]
        except Exception:
            pass
        w.add_page(p)
    try:
        w.add_metadata({"/Author": "", "/Creator": "", "/Producer": "", "/Title": "", "/Subject": "", "/Keywords": ""})
    except Exception:
        pass
    out = io.BytesIO(); w.write(out); out.seek(0)
    return out.read()


# ════════════════════════════════════════════════════════════════════════════
# WORD / POWERPOINT METADATA — .docx/.pptx are ZIP+XML, so we read the hidden
# document properties (author, company, edit history, template/username path)
# and content (comments, tracked changes, speaker notes, hidden slides) with no
# extra libraries. Same find → report → clean model.
# ════════════════════════════════════════════════════════════════════════════
_DOC_EXT = (".docx", ".pptx")
_MEDIA = {".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
          ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}

def scan_document(contents: bytes, name: str) -> dict:
    import zipfile, xml.etree.ElementTree as ET
    rep = {"applicable": True, "kind": "document", "leaked": {}, "risks": [], "pii": {},
           "score": 100, "details": {}, "readable": True, "_flags": {}}
    loc = lambda t: t.split("}")[-1]
    is_word = name.endswith(".docx"); is_ppt = name.endswith(".pptx")
    try:
        z = zipfile.ZipFile(io.BytesIO(contents)); names = set(z.namelist())
        core, app = {}, {}
        if "docProps/core.xml" in names:
            for el in ET.fromstring(z.read("docProps/core.xml")):
                core[loc(el.tag)] = (el.text or "").strip()
        if "docProps/app.xml" in names:
            for el in ET.fromstring(z.read("docProps/app.xml")):
                app[loc(el.tag)] = (el.text or "").strip()
        L = rep["leaked"]
        if core.get("creator"):        L["Author"] = core["creator"]
        if core.get("lastModifiedBy"): L["Last modified by"] = core["lastModifiedBy"]
        if app.get("Company"):         L["Company"] = app["Company"]
        if app.get("Manager"):         L["Manager"] = app["Manager"]
        if core.get("title"):          L["Title"] = core["title"]
        if core.get("subject"):        L["Subject"] = core["subject"]
        if core.get("keywords"):       L["Keywords"] = core["keywords"]
        if core.get("created"):        L["Created"] = core["created"]
        if core.get("modified"):       L["Modified"] = core["modified"]
        det = rep["details"]
        if app.get("Application"): det["Created with"] = (app["Application"] + " " + app.get("AppVersion", "")).strip()
        if core.get("revision"):  det["Revisions"] = core["revision"]
        if app.get("TotalTime") and app["TotalTime"] != "0": det["Total edit time"] = app["TotalTime"] + " min"
        if app.get("Template") and app["Template"].lower() not in ("normal.dotm", "normal", ""): det["Template"] = app["Template"]
        for k in ("Words", "Slides", "Pages", "Paragraphs"):
            if app.get(k): det[k] = app[k]
        f = rep["_flags"]
        if is_word:
            f["comments"] = 1 if "word/comments.xml" in names else 0
            try:
                doc = z.read("word/document.xml").decode("utf-8", "ignore")
                f["tracked"] = ("<w:ins" in doc or "<w:del" in doc)
                f["hidden_text"] = ("<w:vanish" in doc)
            except Exception:
                pass
        if is_ppt:
            f["comments"] = sum(1 for n in names if n.startswith("ppt/comments/") and n.endswith(".xml"))
            f["notes"] = sum(1 for n in names if n.startswith("ppt/notesSlides/notesSlide"))
            try:
                hidden = 0
                for n in names:
                    if n.startswith("ppt/slides/slide") and n.endswith(".xml"):
                        if 'show="0"' in z.read(n).decode("utf-8", "ignore")[:400]:
                            hidden += 1
                f["hidden_slides"] = hidden
            except Exception:
                pass
    except Exception:
        rep["readable"] = False

    s, R, f = 100, rep["risks"], rep["_flags"]
    ids = [rep["leaked"][k] for k in ("Author", "Last modified by", "Company", "Manager") if k in rep["leaked"]]
    if ids:
        s -= 15; R.append({"label": f"Reveals identity: {', '.join(ids)[:80]}", "severity": "high", "cat": "metadata",
                           "detail": "Author / editor / company name is embedded in the file."})
    tpl = rep["details"].get("Template", "")
    if "\\users\\" in tpl.lower() or ":\\" in tpl.lower():
        s -= 10; R.append({"label": f"File path leaks a username: {tpl}", "severity": "high", "cat": "metadata",
                           "detail": "The template path reveals a person's name / computer."})
    if f.get("tracked"):
        s -= 20; R.append({"label": "Tracked changes present — deleted text is recoverable", "severity": "high", "cat": "comments",
                           "detail": "Accept/reject all changes in your editor before sharing."})
    if f.get("comments"):
        s -= 12; R.append({"label": f"{f['comments']} comment thread(s) / note(s)", "severity": "medium", "cat": "comments",
                           "detail": "Internal review comments travel with the file."})
    if f.get("notes"):
        s -= 10; R.append({"label": f"{f['notes']} slide(s) with speaker notes", "severity": "medium", "cat": "comments",
                           "detail": "Speaker notes are often private remarks."})
    if f.get("hidden_slides"):
        s -= 12; R.append({"label": f"{f['hidden_slides']} hidden slide(s)", "severity": "high", "cat": "metadata",
                           "detail": "Hidden slides still travel inside the file."})
    if f.get("hidden_text"):
        s -= 8; R.append({"label": "Hidden text detected", "severity": "medium", "cat": "metadata",
                          "detail": "Text formatted as hidden is still in the file."})
    rep["score"] = max(0, min(100, s))
    rep["categories_present"] = sorted({r2.get("cat") for r2 in R if r2.get("cat")})
    if not R:
        R.append({"label": "No hidden data found — safe to share" if rep["readable"]
                  else "Couldn't read this document.", "severity": "ok", "detail": ""})
    rep["gdpr"] = ({"personal_data": True, "level": "high",
                    "verdict": "This document carries hidden author/edit data — clean it before sharing.",
                    "categories": [k for k in ("Author", "Company") if k in rep["leaked"]],
                    "advice": "Download the cleaned copy to strip author/company/edit history. Tracked changes & comments are flagged — clear those in your editor too."}
                   if ids or f.get("tracked") or f.get("comments") else
                   {"personal_data": bool(ids), "level": "ok" if not rep["leaked"] else "medium",
                    "verdict": "No major exposure found." if not rep["leaked"] else "Only document info found — low risk.",
                    "categories": list(rep["leaked"].keys()), "advice": "You can still strip the metadata below."})
    return rep

def clean_document(contents: bytes, name: str) -> bytes:
    """Strip document properties (author/company/edit history) reliably by
    replacing core.xml/app.xml with empty ones. (Comments/tracked-changes are
    flagged for the user to clear in their editor — removing them safely needs
    rel/content-type surgery we don't risk corrupting the file with.)"""
    import zipfile
    CORE = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
            ' xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/"'
            ' xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            '<dc:creator></dc:creator><cp:lastModifiedBy></cp:lastModifiedBy></cp:coreProperties>')
    z_in = zipfile.ZipFile(io.BytesIO(contents)); out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in z_in.infolist():
            n = item.filename
            if n == "docProps/core.xml":
                zo.writestr(n, CORE)
            elif n == "docProps/app.xml":
                data = z_in.read(n).decode("utf-8", "ignore")
                # blank the identity-bearing fields, keep the file structure valid
                for tag in ("Company", "Manager", "Template"):
                    data = re.sub(rf"<{tag}>.*?</{tag}>", f"<{tag}></{tag}>", data, flags=re.S)
                zo.writestr(n, data)
            else:
                zo.writestr(item, z_in.read(n))
    out.seek(0); return out.read()


# ════════════════════════════════════════════════════════════════════════════
# VIDEO METADATA — phone videos (MP4/MOV) embed GPS location, device & date,
# just like photos. Read via ffmpeg, strip with ffmpeg -map_metadata -1 -c copy
# (no re-encode → fast, no quality loss).
# ════════════════════════════════════════════════════════════════════════════
_VIDEO_EXT = (".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm")
_VIDEO_MAX = 250 * 1024 * 1024   # 250 MB cap (keeps the in-memory model safe)

def _ffmpeg_exe():
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()

def scan_video(contents: bytes, name: str) -> dict:
    import tempfile, subprocess, os
    rep = {"applicable": True, "kind": "video", "leaked": {}, "maps": "", "risks": [],
           "pii": {}, "score": 100, "details": {}, "readable": True}
    tmp = None
    try:
        ext = "." + name.rsplit(".", 1)[-1]
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as f:
            f.write(contents); tmp = f.name
        info = subprocess.run([_ffmpeg_exe(), "-i", tmp], capture_output=True, text=True, timeout=90).stderr
        def grab(key):
            m = re.search(rf"^\s*{key}\s*:\s*(.+?)\s*$", info, re.I | re.M)
            return m.group(1).strip() if m else ""
        make = grab(r"com\.apple\.quicktime\.make") or grab(r"make")
        model = grab(r"com\.apple\.quicktime\.model") or grab(r"model")
        sw = grab(r"com\.apple\.quicktime\.software") or grab(r"encoder")
        ct = grab(r"com\.apple\.quicktime\.creationdate") or grab(r"creation_time")
        loc = grab(r"com\.apple\.quicktime\.location\.iso6709") or grab(r"location")
        if make or model: rep["leaked"]["Camera"] = (make + " " + model).strip()
        if sw: rep["leaked"]["Software"] = sw
        if ct: rep["leaked"]["Date taken"] = ct
        if loc:
            mll = re.match(r"\s*([+-]\d+\.?\d*)([+-]\d+\.?\d*)", loc)
            if mll:
                lat = float(mll.group(1)); lon = float(mll.group(2))
                rep["leaked"]["GPS location"] = f"{lat:.5f}, {lon:.5f}"
                rep["maps"] = f"https://maps.google.com/?q={lat},{lon}"
        det = rep["details"]
        d = re.search(r"Duration:\s*([0-9:.]+)", info)
        r = re.search(r", (\d{3,5}x\d{3,5})", info)
        if d: det["Duration"] = d.group(1)
        if r: det["Resolution"] = r.group(1)
        det["Size"] = f"{len(contents)/1024/1024:.1f} MB"
    except Exception:
        rep["readable"] = False
    finally:
        if tmp and os.path.exists(tmp):
            try: os.remove(tmp)
            except Exception: pass

    s, R = 100, rep["risks"]
    if "GPS location" in rep["leaked"]:
        s -= 35; R.append({"label": f"📍 GPS location embedded: {rep['leaked']['GPS location']}", "severity": "high",
                           "cat": "gps", "detail": "This video reveals exactly where it was filmed — a real safety risk."})
    if "Camera" in rep["leaked"] or "Software" in rep["leaked"]:
        s -= 10; R.append({"label": f"Device info: {rep['leaked'].get('Camera','')} {rep['leaked'].get('Software','')}".strip(), "severity": "medium",
                           "cat": "camera", "detail": "The exact device and software are recorded."})
    if "Date taken" in rep["leaked"]:
        s -= 5; R.append({"label": f"Date/time filmed: {rep['leaked']['Date taken']}", "severity": "low",
                          "cat": "date", "detail": "When the video was captured."})
    rep["score"] = max(0, min(100, s))
    rep["categories_present"] = sorted({r2.get("cat") for r2 in R if r2.get("cat")})
    if not R:
        R.append({"label": "No hidden metadata found — this video looks clean" if rep["readable"]
                  else "Couldn't read this video format.", "severity": "ok", "detail": ""})
    has_loc = "GPS location" in rep["leaked"]
    rep["gdpr"] = ({"personal_data": True, "level": "high",
                    "verdict": "This video reveals personal data — strip it before posting or sharing.",
                    "categories": [k for k in ("GPS location", "Date taken") if k in rep["leaked"]],
                    "advice": "Download the cleaned video below — it's identical but carries no hidden data."}
                   if has_loc else
                   {"personal_data": False, "level": "ok" if not rep["leaked"] else "medium",
                    "verdict": "No location found." if not rep["leaked"] else "Minor device metadata only — low risk.",
                    "categories": list(rep["leaked"].keys()), "advice": "You can still strip the remaining metadata below."})
    return rep

def strip_video(contents: bytes, name: str) -> tuple[bytes, str]:
    import tempfile, subprocess, os
    ext = "." + name.rsplit(".", 1)[-1]
    tin = tout = None
    try:
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as f:
            f.write(contents); tin = f.name
        tout = tin.rsplit(".", 1)[0] + "_clean" + ext
        subprocess.run([_ffmpeg_exe(), "-y", "-i", tin, "-map_metadata", "-1",
                        "-map", "0", "-c", "copy", tout], capture_output=True, timeout=240)
        with open(tout, "rb") as f:
            return f.read(), ext.lstrip(".")
    finally:
        for p in (tin, tout):
            if p and os.path.exists(p):
                try: os.remove(p)
                except Exception: pass


# ════════════════════════════════════════════════════════════════════════════
# EMAIL METADATA — .eml/.msg headers are the ONE place IP addresses really live.
# Received headers reveal the sender/server IPs (≈ location); Bcc exposes hidden
# recipients; X-Mailer reveals the software. Scan + strip to a clean .eml.
# ════════════════════════════════════════════════════════════════════════════
_EMAIL_EXT = (".eml", ".msg")

def _email_extract(contents: bytes, name: str):
    """Return (fields dict, received headers, x-originating headers, client)."""
    From = To = Cc = Bcc = Subject = Date = client = ""
    received, xorig = [], []
    if name.endswith(".eml"):
        import email
        from email import policy
        m = email.message_from_bytes(contents, policy=policy.default)
        g = lambda k: (str(m[k]) if m[k] else "")
        From, To, Cc, Bcc, Subject, Date = g("From"), g("To"), g("Cc"), g("Bcc"), g("Subject"), g("Date")
        client = g("X-Mailer") or g("User-Agent")
        received = [str(x) for x in (m.get_all("Received") or [])]
        xorig = [str(x) for x in (m.get_all("X-Originating-IP") or [])]
        for k in ("X-Sender-IP", "X-Source-IP", "X-Real-IP"):
            if m[k]: xorig.append(str(m[k]))
    else:  # .msg
        import extract_msg, tempfile, os
        tmp = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".msg", delete=False) as f:
                f.write(contents); tmp = f.name
            em = extract_msg.Message(tmp)
            From = str(em.sender or ""); To = str(em.to or ""); Cc = str(em.cc or "")
            Bcc = str(em.bcc or ""); Subject = str(em.subject or ""); Date = str(em.date or "")
            hdr = em.header
            if hdr:
                received = [str(x) for x in (hdr.get_all("Received") or [])]
                xorig = [str(x) for x in (hdr.get_all("X-Originating-IP") or [])]
                client = str(hdr.get("X-Mailer") or hdr.get("User-Agent") or "")
            em.close()
        finally:
            if tmp and os.path.exists(tmp):
                try: os.remove(tmp)
                except Exception: pass
    return {"From": From, "To": To, "Cc": Cc, "Bcc": Bcc, "Subject": Subject, "Date": Date}, received, xorig, client

def scan_email(contents: bytes, name: str) -> dict:
    rep = {"applicable": True, "kind": "email", "leaked": {}, "risks": [], "pii": {},
           "score": 100, "details": {}, "readable": True}
    f = {}; received = []; xorig = []; client = ""; ips = set()
    try:
        f, received, xorig, client = _email_extract(contents, name)
        iprx = re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b")
        for h in received + xorig:
            for ip in iprx.findall(h):
                if not ip.startswith(("127.", "0.")):
                    ips.add(ip)
    except Exception:
        rep["readable"] = False
    L = rep["leaked"]
    if f.get("From"):    L["From"] = f["From"]
    if f.get("To"):      L["To (recipients)"] = f["To"]
    if f.get("Cc"):      L["Cc"] = f["Cc"]
    if f.get("Bcc"):     L["Bcc (hidden recipients)"] = f["Bcc"]
    if f.get("Subject"): L["Subject"] = f["Subject"]
    if f.get("Date"):    L["Date"] = f["Date"]
    if ips:              L["IP addresses"] = ", ".join(sorted(ips)[:6])
    det = rep["details"]
    if client:   det["Email client"] = client
    if received: det["Mail servers (hops)"] = len(received)

    s, R = 100, rep["risks"]
    if ips:
        s -= 30; R.append({"label": f"IP address(es) exposed: {', '.join(sorted(ips)[:4])}", "severity": "high", "cat": "ip",
                           "detail": "Email headers reveal the sender / mail-server IPs — roughly where it came from."})
    if f.get("Bcc"):
        s -= 25; R.append({"label": f"Bcc (hidden recipients) exposed: {f['Bcc'][:60]}", "severity": "high", "cat": "recipients",
                           "detail": "Blind-copied recipients are visible in this saved email."})
    if f.get("To") or f.get("Cc"):
        s -= 10; R.append({"label": "Recipient email addresses present", "severity": "medium", "cat": "recipients",
                           "detail": "To / Cc addresses are personal data."})
    if client:
        s -= 5; R.append({"label": f"Email client / software: {client[:50]}", "severity": "low", "cat": "client",
                          "detail": "Reveals the software used to send."})
    rep["score"] = max(0, min(100, s))
    rep["categories_present"] = sorted({r2.get("cat") for r2 in R if r2.get("cat")})
    if not R:
        R.append({"label": "No sensitive headers found" if rep["readable"] else "Couldn't read this email file.", "severity": "ok", "detail": ""})
    rep["gdpr"] = ({"personal_data": True, "level": "high",
                    "verdict": "This email exposes IPs / recipients — clean the headers before sharing it.",
                    "categories": [k for k in ("IP addresses", "Bcc (hidden recipients)", "From") if k in L],
                    "advice": "Download the cleaned .eml — Received / IP / Message-ID / X-* headers and Bcc are removed."}
                   if (ips or f.get("Bcc") or f.get("From")) else
                   {"personal_data": False, "level": "ok", "verdict": "No major exposure found.", "categories": [], "advice": ""})
    return rep

def clean_email(contents: bytes, name: str) -> tuple[bytes, str]:
    import email
    from email import policy
    if name.endswith(".msg"):
        import extract_msg, tempfile, os
        from email.message import EmailMessage
        tmp = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".msg", delete=False) as fobj:
                fobj.write(contents); tmp = fobj.name
            em = extract_msg.Message(tmp)
            nm = EmailMessage()
            if em.sender:  nm["From"] = str(em.sender)
            if em.to:      nm["To"] = str(em.to)
            if em.cc:      nm["Cc"] = str(em.cc)
            if em.subject: nm["Subject"] = str(em.subject)
            nm.set_content(str(em.body or ""))
            em.close()
            return nm.as_bytes(), "eml"
        finally:
            if tmp and os.path.exists(tmp):
                try: os.remove(tmp)
                except Exception: pass
    m = email.message_from_bytes(contents, policy=policy.default)
    drop = {"received", "x-originating-ip", "x-sender-ip", "x-source-ip", "x-real-ip", "message-id",
            "x-mailer", "user-agent", "bcc", "return-path", "x-originating-email", "dkim-signature",
            "authentication-results", "received-spf"}
    names = {k for k in m.keys() if k.lower() in drop or k.lower().startswith("x-")}
    for h in names:
        while h in m:
            del m[h]
    return m.as_bytes(), "eml"


# ── API Routes ────────────────────────────────────────────────────────────────
@app.get("/")
def root():
    return {"status": "ok", "product": "Velytics API", "version": "1.0.0"}


@app.get("/health")
def health():
    return {"status": "healthy"}


def _open_excel(contents: bytes) -> pd.ExcelFile:
    """Open an Excel workbook, falling back to the calamine engine for files that
    openpyxl chokes on (LibreOffice exports, slightly-malformed XML, etc.)."""
    try:
        return pd.ExcelFile(io.BytesIO(contents))
    except Exception:
        return pd.ExcelFile(io.BytesIO(contents), engine="calamine")


def _pick_sheet(xls: pd.ExcelFile, requested: str = "") -> str:
    """Choose which sheet to analyse. Honours an explicit request, else picks the
    first sheet that actually holds a table (≥2 non-blank rows)."""
    names = xls.sheet_names
    if requested and requested in names:
        return requested
    for sh in names:
        try:
            probe = xls.parse(sheet_name=sh, header=None, nrows=30)
        except Exception:
            continue
        if probe.dropna(how="all").shape[0] >= 2:
            return sh
    return names[0]


def _read_and_clean(contents: bytes, name: str, sheet: str = ""):
    """Shared by /analyze and /download: read any format raw → smart_clean."""
    sheets: list[str] = []
    sheet_used: str = ""
    if name.endswith(".csv"):
        df_raw = pd.read_csv(io.BytesIO(contents), header=None, low_memory=False)
    elif name.endswith(".json"):
        parsed = pd.read_json(io.BytesIO(contents))
        df_raw = pd.DataFrame([list(parsed.columns)] + parsed.values.tolist())
    else:
        xls = _open_excel(contents)
        sheets = list(xls.sheet_names)
        sheet_used = _pick_sheet(xls, sheet)
        df_raw = xls.parse(sheet_name=sheet_used, header=None)
    currency = detect_currency(df_raw)
    df, fixes = smart_clean(df_raw)
    if len(sheets) > 1:
        fixes.insert(0, f"Workbook has {len(sheets)} sheets — analysing '{sheet_used}'")
    return df, fixes, sheets, sheet_used, currency


# Auto-detect the kind of data from its columns, so the user never has to pick.
_TYPE_SIGNATURES = [
    ("Inventory",    ["stock", "sku", "reorder", "expiry", "balance", "warehouse", "on hand", "inventory", "quantity"]),
    ("Sales",        ["revenue", "sales", "target", "salesperson", "sales person", "order value", "quota"]),
    ("Finance",      ["budget", "expense", "income", "ledger", "invoice", "profit", "gst", "debit", "credit", "vendor", "spend", "payable", "receivable", "cost center", "cash flow", "p&l"]),
    ("HR & Payroll", ["salary", "payroll", "employee", "attrition", "headcount", "designation", "department"]),
    ("Retail",       ["order", "return", "cart", "channel", "ecommerce", "e-commerce"]),
    ("Healthcare",   ["patient", "diagnosis", "doctor", "admission", "ward", "readmission"]),
    ("Education",    ["student", "attendance", "marks", "score", "subject", "exam", "grade"]),
    ("Logistics",    ["shipment", "route", "driver", "delivery", "freight", "vehicle", "dispatch"]),
    ("Restaurant",   ["menu", "dish", "food cost", "waste", "shift", "cuisine"]),
    ("Manufacturing",["machine", "downtime", "defect", "output", "production", "yield"]),
    ("Marketing",    ["campaign", "spend", "ctr", "roas", "impressions", "clicks", "leads", "conversion"]),
    ("Banking",      ["loan", "npa", "disbursed", "interest", "emi", "outstanding"]),
    ("Real Estate",  ["property", "listing", "sqft", "bedroom", "rent", "tenant", "carpet"]),
]
def detect_type(profile: list[dict]) -> str:
    names = " ".join(str(p["column"]).lower() for p in profile)
    best, best_score = "Data", 0
    for label, kws in _TYPE_SIGNATURES:
        score = sum(1 for kw in kws if kw in names)
        if score > best_score:
            best, best_score = label, score
    return best if best_score >= 2 else "Data"


def _first_breakdown(result: dict):
    """First non-trend breakdown section (has name + pct), e.g. 'Revenue by region'."""
    for k, v in result.items():
        if k in ("kpis", "alerts", "columns"):
            continue
        if isinstance(v, list) and v and isinstance(v[0], dict):
            if "month" in v[0]:
                continue
            if "pct" in v[0] and "name" in v[0]:
                return k, v
    return None


def build_summary(result: dict, detected_type: str) -> list[dict]:
    """Plain-English narrative built deterministically from the computed numbers.
    Returns a list of {text, tone} where tone is good | bad | neutral."""
    K = result.get("kpis") or {}
    pts: list[dict] = []
    fmt = lambda label: K.get(label + "_fmt") or _fmt(float(K[label])) if label in K else ""

    # 1) Headline — the primary money KPI (first money kpi, else first kpi)
    money_keys = [k[:-4] for k in K if k.endswith("_fmt")]
    head = money_keys[0] if money_keys else next((k for k in K if not k.endswith("_fmt")), None)
    if head:
        extra = f" across {int(K['Orders']):,} orders" if isinstance(K.get("Orders"), (int, float)) else ""
        val = fmt(head) if head in money_keys else f"{K[head]:,}" if isinstance(K[head], (int, float)) else K[head]
        pts.append({"text": f"{head} is {val}{extra}.", "tone": "neutral"})

    # 2) Growth — momentum vs last period
    for gk in ("MoM Growth %", "YoY Growth %"):
        if gk in K:
            v = K[gk]; period = "last month" if "MoM" in gk else "last year"
            if v > 0:   pts.append({"text": f"Up {v}% vs {period} — momentum is positive.", "tone": "good"})
            elif v < 0: pts.append({"text": f"Down {abs(v)}% vs {period} — momentum is slowing.", "tone": "bad"})
            break

    # 3) Target attainment
    if "Target Achievement %" in K:
        a = K["Target Achievement %"]
        if a >= 100:  pts.append({"text": f"Target beaten — {a}% of goal achieved.", "tone": "good"})
        elif a >= 80: pts.append({"text": f"At {a}% of target — close, push to finish.", "tone": "neutral"})
        else:         pts.append({"text": f"Only {a}% of target reached — below the 80% line.", "tone": "bad"})

    # 4) Top driver — who/what leads the biggest breakdown
    bd = _first_breakdown(result)
    if bd:
        name, rows = bd; top = rows[0]
        valtxt = _fmt(top["value"]) if "value" in top else f"{top.get('count', '')}"
        pts.append({"text": f"{top['name']} leads {name.lower()} at {valtxt} ({top['pct']}% of total).", "tone": "neutral"})

    # 5) Margin health
    if "Gross Margin %" in K:
        m = K["Gross Margin %"]
        pts.append({"text": f"Gross margin is {m}%." + ("" if m >= 15 else " That's thin — watch costs."),
                    "tone": "good" if m >= 15 else "bad"})

    # 6) The single most important risk (first Critical/Warning alert not already covered)
    seen_nums = set()
    for p in pts:
        seen_nums.update(re.findall(r"\d+\.?\d*", p["text"]))
    for a in (result.get("alerts") or []):
        if a.get("type") in ("Critical", "Warning"):
            nums = re.findall(r"\d+\.?\d*", a["text"])
            if nums and all(n in seen_nums for n in nums):
                continue   # this alert just repeats a point we already made
            pts.append({"text": a["text"], "tone": "bad"}); break

    # 7) Forward look
    for fk in ("Forecast (next 3 mo)", "Forecast (next 3 Mo)"):
        if fk in K:
            pts.append({"text": f"Projected next 3 months: {fmt(fk)}.", "tone": "neutral"}); break

    return pts[:6]


def _apply_edits_and_cleaning(df, fixes, edits_json, clean_options_json):
    """Apply manual corrections + the client's blank-handling choice. Shared by
    /analyze and /report so the exported report matches the dashboard exactly."""
    applied_edits = 0
    try:
        edit_list = json.loads(edits_json) if edits_json else []
    except Exception:
        edit_list = []
    if isinstance(edit_list, list):
        for e in edit_list:
            try:
                r = int(e["row"]); c = str(e["col"]); v = e.get("value")
                if c in df.columns and 0 <= r < len(df):
                    ci = df.columns.get_loc(c)
                    if pd.api.types.is_numeric_dtype(df[c]):
                        v = pd.to_numeric(str(v).replace(",", "").strip(), errors="coerce")
                    elif v is None:
                        v = ""
                    df.iat[r, ci] = v
                    applied_edits += 1
            except Exception:
                pass
    if applied_edits:
        fixes.insert(0, f"{applied_edits} manual correction{'s' if applied_edits != 1 else ''} applied before analysis")

    raw_blanks = int(df.isna().sum().sum())
    try:
        copts = json.loads(clean_options_json) if clean_options_json else {}
        if not isinstance(copts, dict):
            copts = {}
    except Exception:
        copts = {}
    blanks_mode = copts.get("blanks", "leave")
    if raw_blanks and blanks_mode != "leave":
        if blanks_mode == "drop_rows":
            before = len(df)
            df = df.dropna().reset_index(drop=True)
            dropped = before - len(df)
            if dropped:
                fixes.insert(0, f"Removed {dropped} row(s) containing blanks (your choice)")
        else:
            filled = 0
            for c in df.select_dtypes(include=[np.number]).columns:
                na = int(df[c].isna().sum())
                if not na:
                    continue
                if blanks_mode == "zero":
                    fillv = 0.0
                elif blanks_mode == "mean":
                    fillv = float(df[c].mean()) if df[c].notna().any() else 0.0
                elif blanks_mode == "median":
                    fillv = float(df[c].median()) if df[c].notna().any() else 0.0
                else:
                    fillv = None
                if fillv is not None:
                    df[c] = df[c].fillna(fillv)
                    filled += na
            if filled:
                label = {"zero": "0", "mean": "the column average", "median": "the column median"}.get(blanks_mode, blanks_mode)
                fixes.insert(0, f"Filled {filled} blank number(s) with {label} (your choice)")
    return df, fixes, applied_edits, raw_blanks, blanks_mode


@app.post("/analyze")
async def analyze(
    file: UploadFile = File(...),
    module: str = Form(default=""),       # optional override; normally auto-detected
    sheet: str = Form(default=""),
    filters: str = Form(default=""),     # JSON: {col: {values:[...]} | {min,max}}
    group_by: str = Form(default=""),    # optional: column to break a measure down by
    measure: str = Form(default=""),     # optional: numeric column to measure
    edits: str = Form(default=""),       # JSON: [{"row":int,"col":str,"value":any}] manual corrections
    clean_options: str = Form(default=""),  # JSON: {"blanks": "leave"|"zero"|"mean"|"median"|"drop_rows"}
):
    # Validate file type
    name = file.filename.lower() if file.filename else ""
    if not any(name.endswith(ext) for ext in [".xlsx", ".csv", ".json", *_IMG_EXT, *_PDF_EXT, *_DOC_EXT, *_VIDEO_EXT, *_EMAIL_EXT]):
        raise HTTPException(400, "Unsupported file. Please upload a spreadsheet, image, PDF, Word/PowerPoint, video, or email.")

    contents = await file.read()

    # ── Image / PDF / Word / PowerPoint / Video / Email path: scan footprint only ──
    if name.endswith(_IMG_EXT) or name.endswith(_PDF_EXT) or name.endswith(_DOC_EXT) or name.endswith(_VIDEO_EXT) or name.endswith(_EMAIL_EXT):
        if name.endswith(_VIDEO_EXT) and len(contents) > _VIDEO_MAX:
            fp = {"applicable": True, "kind": "video", "leaked": {}, "risks": [{"label": f"Video is too large ({len(contents)/1024/1024:.0f} MB) — max 250 MB. Try a shorter clip.", "severity": "medium", "detail": ""}],
                  "score": 100, "details": {}, "categories_present": [], "gdpr": {"personal_data": False, "level": "ok", "verdict": "File too large to scan here.", "categories": [], "advice": ""}}
            ft, dt = "video", "Video"
        elif name.endswith(_PDF_EXT):   fp, ft, dt = scan_pdf(contents, name), "pdf", "PDF"
        elif name.endswith(_DOC_EXT):   fp, ft, dt = scan_document(contents, name), "document", ("Word" if name.endswith(".docx") else "PowerPoint")
        elif name.endswith(_VIDEO_EXT): fp, ft, dt = scan_video(contents, name), "video", "Video"
        elif name.endswith(_EMAIL_EXT): fp, ft, dt = scan_email(contents, name), "email", "Email"
        else:                            fp, ft, dt = scan_image(contents, name), "image", "Image"
        return {
            "file": file.filename, "file_type": ft, "currency": "$",
            "footprint": fp, "result": {}, "metadata": [], "health": None, "summary": [],
            "fixes": [], "filter_schema": [], "dimensions": [], "measures": [],
            "applied_filters": {}, "applied_edits": 0, "raw_blanks": 0, "blanks_mode": "leave",
            "preview": [], "rows": 0, "total_rows": 0, "columns": 0, "sheets": [], "sheet_used": "",
            "detected_type": dt,
        }

    try:
        df, fixes, sheets, sheet_used, currency = _read_and_clean(contents, name, sheet)
        global _CURRENCY
        _CURRENCY = currency
    except Exception as e:
        raise HTTPException(400, f"Could not read file: {str(e)}")

    # ── Manual corrections + cleaning-with-approval (deterministic indices) ──
    df, fixes, applied_edits, raw_blanks, blanks_mode = _apply_edits_and_cleaning(df, fixes, edits, clean_options)

    # ── Universal adaptive analysis — SAME engine for every module ──
    try:
        flt = json.loads(filters) if filters else {}
        if not isinstance(flt, dict):
            flt = {}
    except Exception:
        flt = {}

    try:
        profile = profile_columns(df)
        filter_schema = build_filter_schema(df, profile)
        dimensions = _groupable(df, profile)
        measures   = [p["column"] for p in profile if p["role"] == "number"]
        detected_type = (module or detect_type(profile))   # auto-detect, unless overridden
        dff = apply_filters(df, flt)
        result = analyze_auto(dff, profile, group_by=group_by, measure=measure)
        apply_module_pack(detected_type, dff, profile, result)   # layer industry features on top
        summary = build_summary(result, detected_type)     # plain-English narrative
        health = compute_health(df, profile, fixes)        # data health score
        metadata = build_metadata(df, profile)             # column data dictionary
        footprint = scan_footprint(contents, name, df)     # privacy / hidden-footprint scan
    except Exception as e:
        raise HTTPException(500, f"Analysis failed: {str(e)}")

    # Cleaned preview for the "Clean & Organize" view (JSON-safe via pandas).
    preview = json.loads(df.head(50).to_json(orient="records", date_format="iso"))

    return {
        "module": module,
        "detected_type": detected_type,
        "currency": currency,
        "file": file.filename,
        "rows": int(len(dff)),
        "total_rows": int(len(df)),
        "columns": int(len(df.columns)),
        "sheets": sheets,
        "sheet_used": sheet_used,
        "fixes": fixes,
        "summary": summary,
        "health": health,
        "metadata": metadata,
        "footprint": footprint,
        "filter_schema": filter_schema,
        "dimensions": dimensions,
        "measures": measures,
        "applied_filters": flt,
        "applied_edits": applied_edits,
        "raw_blanks": raw_blanks,
        "blanks_mode": blanks_mode,
        "preview": preview,
        "result": result,
    }


@app.post("/download")
async def download(
    file: UploadFile = File(...),
    sheet: str = Form(default=""),
    filters: str = Form(default=""),
    fmt: str = Form(default="xlsx"),     # "xlsx" or "csv"
):
    """Return the CLEANED dataset as a downloadable file (the 'Clean & Organize' tool)."""
    name = file.filename.lower() if file.filename else ""
    if not any(name.endswith(ext) for ext in [".xlsx", ".csv", ".json"]):
        raise HTTPException(400, "Unsupported file.")
    contents = await file.read()
    try:
        df, _, _, _, _ = _read_and_clean(contents, name, sheet)
    except Exception as e:
        raise HTTPException(400, f"Could not read file: {str(e)}")
    try:
        flt = json.loads(filters) if filters else {}
        if isinstance(flt, dict):
            df = apply_filters(df, flt)
    except Exception:
        pass

    base = (file.filename or "data").rsplit(".", 1)[0]
    buf = io.BytesIO()
    if fmt == "csv":
        df.to_csv(buf, index=False)
        media = "text/csv"; fname = f"{base}_cleaned.csv"
    else:
        with pd.ExcelWriter(buf, engine="openpyxl") as xw:
            df.to_excel(xw, index=False, sheet_name="Cleaned")
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        fname = f"{base}_cleaned.xlsx"
    buf.seek(0)
    return StreamingResponse(buf, media_type=media,
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ── Full Excel report: summary + data + a sheet & native chart per breakdown ──
def _safe_sheet_name(name: str, used: set) -> str:
    s = "".join(c for c in str(name) if c not in '[]:*?/\\').strip()[:31] or "Sheet"
    base, i = s, 1
    while s.lower() in used:
        suffix = f"_{i}"; s = base[:31 - len(suffix)] + suffix; i += 1
    used.add(s.lower())
    return s

def _section_value_key(row: dict):
    for k in ("value", "revenue", "amount", "total", "count", "qty", "units", "days"):
        if k in row and isinstance(row[k], (int, float)):
            return k
    for k, v in row.items():
        if k != "pct" and isinstance(v, (int, float)):
            return k
    return None

@app.post("/report")
async def report(
    file: UploadFile = File(...),
    module: str = Form(default=""),
    sheet: str = Form(default=""),
    filters: str = Form(default=""),
    edits: str = Form(default=""),
    clean_options: str = Form(default=""),
):
    """Full analytical report as a multi-sheet .xlsx with native Excel charts —
    mirrors exactly what the dashboard shows (same edits, cleaning, filters)."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, DoughnutChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
    from openpyxl.chart.label import DataLabelList

    name = file.filename.lower() if file.filename else ""
    if not any(name.endswith(ext) for ext in [".xlsx", ".csv", ".json"]):
        raise HTTPException(400, "Unsupported file.")
    contents = await file.read()
    try:
        df, fixes, sheets, sheet_used, currency = _read_and_clean(contents, name, sheet)
        global _CURRENCY
        _CURRENCY = currency
    except Exception as e:
        raise HTTPException(400, f"Could not read file: {str(e)}")

    df, fixes, _ae, _rb, _bm = _apply_edits_and_cleaning(df, fixes, edits, clean_options)
    try:
        flt = json.loads(filters) if filters else {}
        if not isinstance(flt, dict):
            flt = {}
    except Exception:
        flt = {}
    try:
        profile = profile_columns(df)
        detected_type = (module or detect_type(profile))
        dff = apply_filters(df, flt)
        result = analyze_auto(dff, profile)
        apply_module_pack(detected_type, dff, profile, result)
        summary = build_summary(result, detected_type)
    except Exception as e:
        raise HTTPException(500, f"Report build failed: {str(e)}")

    # ── Theme (matches the app's indigo dashboard) ──
    INDIGO, INK, GREY, CARD, CARDBDR = "6366F1", "0F172A", "64748B", "F5F7FF", "C7D2FE"
    WHITE_BOLD = Font(bold=True, color="FFFFFF", size=11)
    HDRFILL = PatternFill("solid", fgColor=INDIGO)
    CARDFILL = PatternFill("solid", fgColor=CARD)
    box = Border(*([Side(style="thin", color=CARDBDR)] * 4))
    CUR_FMT = f'"{currency}"#,##0'
    MONEY_KEYS = {"value", "revenue", "amount", "total", "sales", "profit", "cost", "spend", "payroll", "disbursed"}
    TONE = {"good": "059669", "bad": "DC2626", "neutral": INK}

    def _chart(is_trend, wss, n, title, money, w, h):
        """A native chart that SHOWS its data — category labels on the axis + the
        value printed on every bar/point — so the numbers are readable in-chart."""
        c = LineChart() if is_trend else BarChart()
        if not is_trend:
            c.type = "bar"; c.gapWidth = 60
        c.title = title; c.legend = None; c.width = w; c.height = h
        c.add_data(Reference(wss, min_col=2, min_row=1, max_row=n + 1), titles_from_data=True)
        c.set_categories(Reference(wss, min_col=1, min_row=2, max_row=n + 1))
        # print the values on the chart
        c.dataLabels = DataLabelList()
        c.dataLabels.showVal = True
        c.dataLabels.numFmt = CUR_FMT if money else "#,##0"
        c.dataLabels.showSerName = c.dataLabels.showCatName = c.dataLabels.showLegendKey = False
        # make sure both axes (incl. the category labels) are visible
        c.x_axis.delete = False; c.y_axis.delete = False
        c.x_axis.majorGridlines = None
        if not is_trend:
            c.y_axis.majorGridlines = None
        try:
            c.series[0].graphicalProperties.solidFill = INDIGO
        except Exception:
            pass
        return c

    wb = Workbook()
    used: set = set()
    K = result.get("kpis") or {}
    kkeys = [k for k in K if not k.endswith("_fmt")]
    sections = [(k, v) for k, v in result.items()
                if isinstance(v, list) and v and isinstance(v[0], dict) and k not in ("alerts", "columns")]

    # ── Build a styled sheet + native chart for each section; remember refs ──
    refs = []
    for title, rows in sections:
        try:
            is_trend = "month" in rows[0]
            lk = "month" if is_trend else ("name" if "name" in rows[0] else list(rows[0].keys())[0])
            vk = _section_value_key(rows[0])
            money = vk in MONEY_KEYS
            wss = wb.create_sheet(_safe_sheet_name(title, used))
            for cc, txt in ((1, lk.title()), (2, "Value")):
                c = wss.cell(1, cc, txt); c.font = WHITE_BOLD; c.fill = HDRFILL
                c.alignment = Alignment(horizontal="left", vertical="center")
            wss.row_dimensions[1].height = 18
            for i, row in enumerate(rows, start=2):
                wss.cell(i, 1, str(row.get(lk, "")))
                if vk is not None:
                    cell = wss.cell(i, 2, row.get(vk))
                    cell.number_format = CUR_FMT if money else "#,##0"
            n = len(rows)
            wss.column_dimensions["A"].width = 30; wss.column_dimensions["B"].width = 18
            wss.freeze_panes = "A2"
            # in-cell data bars on the value column (like the references' "Top drivers")
            if vk is not None and n >= 2 and not is_trend:
                wss.conditional_formatting.add(
                    f"B2:B{n+1}",
                    DataBarRule(start_type="num", start_value=0, end_type="max",
                                color="6366F1", showValue=True))
            if vk is not None and n >= 1:
                ch = _chart(is_trend, wss, n, title, money, w=17, h=max(8.5, n * 0.55))
                wss.add_chart(ch, "D2")
                refs.append({"wss": wss, "n": n, "is_trend": is_trend, "money": money, "title": title})
        except Exception:
            continue

    # ── Dashboard (Summary) sheet — KPI cards + at-a-glance + charts ──
    ws = wb.create_sheet(_safe_sheet_name("Dashboard", used), 0)   # first sheet
    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 15.5
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    t = ws.cell(1, 1, f"{detected_type} Report"); t.font = Font(bold=True, size=18, color="FFFFFF")
    t.fill = HDRFILL; t.alignment = Alignment(horizontal="left", vertical="center"); ws.row_dimensions[1].height = 32
    ws.merge_cells("A2:H2")
    ws.cell(2, 1, f"Generated by Velytics   ·   {len(dff):,} rows   ·   currency {currency}").font = Font(color=GREY, size=10)

    # KPI cards (4 per row, each 2 cols × 2 rows)
    base_row = 4
    for i, k in enumerate(kkeys):
        br, bc = i // 4, i % 4
        lr = base_row + br * 3; vr = lr + 1
        c1 = 1 + bc * 2; c2 = c1 + 1
        ws.merge_cells(start_row=lr, start_column=c1, end_row=lr, end_column=c2)
        ws.merge_cells(start_row=vr, start_column=c1, end_row=vr, end_column=c2)
        lab = ws.cell(lr, c1, k.upper()); lab.font = Font(bold=True, color=GREY, size=8)
        lab.alignment = Alignment(horizontal="left", vertical="center")
        raw = K.get(k)
        is_growth = any(g in k.lower() for g in ("growth", "mom", "yoy", "change")) and isinstance(raw, (int, float))
        val = K.get(k + "_fmt"); val = (raw if val is None else val)
        vcolor = INK
        if is_growth:
            val = f'{"▲" if raw >= 0 else "▼"} {abs(raw)}%'; vcolor = "059669" if raw >= 0 else "DC2626"
        vc = ws.cell(vr, c1, val); vc.font = Font(bold=True, color=vcolor, size=14)
        vc.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[lr].height = 14; ws.row_dimensions[vr].height = 22
        for rr in (lr, vr):
            for cc in (c1, c2):
                cell = ws.cell(rr, cc); cell.fill = CARDFILL; cell.border = box
    cards_end = base_row + ((len(kkeys) + 3) // 4) * 3

    # At a glance
    r = cards_end + 1
    ws.cell(r, 1, "AT A GLANCE").font = Font(bold=True, color=INDIGO, size=11); r += 1
    for p in (summary or []):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(r, 1, "•  " + p.get("text", ""))
        c.font = Font(color=TONE.get(p.get("tone"), INK), size=10.5); r += 1
    r += 1

    # Two charts on the dashboard: a trend line + a composition donut
    trend_ref = next((x for x in refs if x["is_trend"]), None)
    bd_ref = next((x for x in refs if not x["is_trend"]), None)
    if trend_ref:
        ws.add_chart(_chart(True, trend_ref["wss"], trend_ref["n"], trend_ref["title"], trend_ref["money"], w=15, h=8), f"A{r}")
    if bd_ref:
        dn = DoughnutChart(); dn.title = bd_ref["title"]; dn.width = 15; dn.height = 8; dn.holeSize = 55
        dn.add_data(Reference(bd_ref["wss"], min_col=2, min_row=1, max_row=bd_ref["n"] + 1), titles_from_data=True)
        dn.set_categories(Reference(bd_ref["wss"], min_col=1, min_row=2, max_row=bd_ref["n"] + 1))
        dn.dataLabels = DataLabelList(); dn.dataLabels.showPercent = True
        ws.add_chart(dn, f"E{r}")

    # ── Cleaned data sheet (last) ──
    wsd = wb.create_sheet(_safe_sheet_name("Data", used))
    wsd.append([str(c) for c in dff.columns])
    for c in range(1, len(dff.columns) + 1):
        hc = wsd.cell(1, c); hc.font = WHITE_BOLD; hc.fill = HDRFILL
    wsd.freeze_panes = "A2"
    for _, row in dff.iterrows():
        wsd.append([None if pd.isna(v) else (v if isinstance(v, (int, float, str)) else str(v)) for v in row.tolist()])

    base = (file.filename or "data").rsplit(".", 1)[0]
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{base}_report.xlsx"'})


def _mask_pii(s: str) -> str:
    s = _PII["Email"].sub(lambda m: m.group(0)[0] + "***@***", s)
    s = _PII["IBAN"].sub("IBAN****", s)
    s = _PII["Credit card"].sub("**** **** **** ****", s)
    s = _PII["US SSN"].sub("***-**-****", s)
    s = _PII["UK NINO"].sub("********", s)
    s = _PII["Aadhaar"].sub("**** **** ****", s)
    s = _PII["PAN (India)"].sub("XXXXX****X", s)
    s = _PII["Phone"].sub("**********", s)
    return s


def _pseudonym(label: str, n: int) -> str:
    """A consistent fake value for a real one — keeps the data usable (same input
    always maps to the same token, distinct values stay distinct) but anonymous."""
    l = (label or "").lower()
    if "email" in l:  return f"user{n}@example.com"
    if "phone" in l:  return f"+10000{n:06d}"
    if "name" in l:   return f"Person {n}"
    if "iban" in l or "account" in l: return f"ACCT{n:08d}"
    if "card" in l:   return f"4000-0000-0000-{n:04d}"
    return f"{(label or 'ID').replace(' ', '')}-{n:04d}"


@app.post("/scrub")
async def scrub(
    file: UploadFile = File(...),
    identity: bool = Form(True),         # wipe author / company / dates
    hidden_sheets: bool = Form(True),    # remove hidden sheets
    comments: bool = Form(True),         # strip cell comments
    unhide: bool = Form(True),           # unhide rows / columns
    redact_pii: bool = Form(False),      # mask emails / phones / IDs in cells
    remove_columns: str = Form(default=""),  # JSON list of column headers to delete entirely (manual)
    remove_words: str = Form(default=""),    # JSON list of words → blank any cell containing them (manual)
    strip_categories: str = Form(default=""),  # images: JSON list subset of {gps,camera,owner,date}; empty = all
    anonymize: bool = Form(False),       # spreadsheets: replace PII with consistent fake IDs (keeps data usable)
):
    """Return a privacy-clean copy, applying ONLY the options the user selected —
    machine-detected items (identity/hidden/comments/PII) PLUS the client's own
    manual picks (remove whole columns, blank cells containing chosen words)."""
    name = file.filename.lower() if file.filename else ""

    # ── Image path: strip the SELECTED EXIF categories, return a clean photo ──
    if name.endswith(_IMG_EXT):
        contents = await file.read()
        try:
            cats = json.loads(strip_categories) if strip_categories else None
            if not isinstance(cats, list): cats = None
        except Exception:
            cats = None
        try:
            clean, fmt = strip_image(contents, name, cats)
        except Exception as e:
            raise HTTPException(500, f"Could not clean image: {str(e)}")
        ext = "jpg" if fmt == "JPEG" else fmt.lower()
        base = (file.filename or "photo").rsplit(".", 1)[0]
        media = "image/jpeg" if fmt == "JPEG" else f"image/{ext}"
        return StreamingResponse(io.BytesIO(clean), media_type=media,
            headers={"Content-Disposition": f'attachment; filename="{base}_cleaned.{ext}"'})

    # ── PDF path: strip metadata + scripts/attachments + annotations ──
    if name.endswith(_PDF_EXT):
        contents = await file.read()
        try:
            clean = clean_pdf(contents)
        except Exception as e:
            raise HTTPException(500, f"Could not clean PDF: {str(e)}")
        base = (file.filename or "document").rsplit(".", 1)[0]
        return StreamingResponse(io.BytesIO(clean), media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{base}_cleaned.pdf"'})

    # ── Word / PowerPoint path: strip document properties ──
    if name.endswith(_DOC_EXT):
        contents = await file.read()
        ext = name.rsplit(".", 1)[-1]
        try:
            clean = clean_document(contents, name)
        except Exception as e:
            raise HTTPException(500, f"Could not clean document: {str(e)}")
        base = (file.filename or "document").rsplit(".", 1)[0]
        return StreamingResponse(io.BytesIO(clean), media_type=_MEDIA.get("." + ext, "application/octet-stream"),
            headers={"Content-Disposition": f'attachment; filename="{base}_cleaned.{ext}"'})

    # ── Video path: strip metadata with ffmpeg (no re-encode) ──
    if name.endswith(_VIDEO_EXT):
        contents = await file.read()
        if len(contents) > _VIDEO_MAX:
            raise HTTPException(400, "Video too large (max 250 MB).")
        try:
            clean, ext = strip_video(contents, name)
        except Exception as e:
            raise HTTPException(500, f"Could not clean video: {str(e)}")
        base = (file.filename or "video").rsplit(".", 1)[0]
        return StreamingResponse(io.BytesIO(clean), media_type=f"video/{'mp4' if ext in ('mp4','m4v','mov') else ext}",
            headers={"Content-Disposition": f'attachment; filename="{base}_cleaned.{ext}"'})

    # ── Email path: strip technical headers (Received/IP/Message-ID/X-*) + Bcc ──
    if name.endswith(_EMAIL_EXT):
        contents = await file.read()
        try:
            clean, ext = clean_email(contents, name)
        except Exception as e:
            raise HTTPException(500, f"Could not clean email: {str(e)}")
        base = (file.filename or "email").rsplit(".", 1)[0]
        return StreamingResponse(io.BytesIO(clean), media_type="message/rfc822",
            headers={"Content-Disposition": f'attachment; filename="{base}_cleaned.{ext}"'})

    if not name.endswith(".xlsx"):
        raise HTTPException(400, "Footprint cleaning is for Excel .xlsx files.")
    contents = await file.read()
    try:
        rm_cols = json.loads(remove_columns) if remove_columns else []
        if not isinstance(rm_cols, list): rm_cols = []
    except Exception:
        rm_cols = []
    try:
        rm_words = json.loads(remove_words) if remove_words else []
        if not isinstance(rm_words, list): rm_words = []
    except Exception:
        rm_words = []
    rm_cols_l = [str(c).strip().lower() for c in rm_cols if str(c).strip()]
    rm_words_l = [str(w).strip().lower() for w in rm_words if str(w).strip()]
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents))
        if identity:
            from openpyxl.packaging.core import DocumentProperties
            wb.properties = DocumentProperties()
            wb.properties.creator = ""          # openpyxl defaults this to "openpyxl"
            wb.properties.lastModifiedBy = ""
        if hidden_sheets:
            for ws in list(wb.worksheets):
                if ws.sheet_state != "visible" and len(wb.worksheets) > 1:
                    wb.remove(ws)
        # ── Manual: delete whole columns by header (search first 5 rows for the name) ──
        if rm_cols_l:
            for ws in wb.worksheets:
                try:
                    hits = []
                    for r in range(1, min(6, ws.max_row + 1)):
                        for c in range(1, ws.max_column + 1):
                            v = ws.cell(r, c).value
                            if v is not None and str(v).strip().lower() in rm_cols_l:
                                hits.append(c)
                        if hits: break
                    for idx in sorted(set(hits), reverse=True):
                        ws.delete_cols(idx, 1)
                except Exception:
                    continue
        # ── Anonymize: replace PII column values with consistent fake IDs ──
        if anonymize:
            try:
                df_a, _f, _s, _su, _c = _read_and_clean(contents, name)
                pii = scan_pii(df_a)            # {label: [columns]}
                col_label = {}
                for label, cols in pii.items():
                    for c in cols:
                        col_label[str(c).strip().lower()] = label
                if col_label:
                    for ws in wb.worksheets:
                        hdr = None
                        for r in range(1, min(6, ws.max_row + 1)):
                            vals = [str(ws.cell(r, c).value).strip().lower() if ws.cell(r, c).value is not None else "" for c in range(1, ws.max_column + 1)]
                            if any(v in col_label for v in vals):
                                hdr = r; break
                        if not hdr:
                            continue
                        targets = {c: col_label[str(ws.cell(hdr, c).value).strip().lower()]
                                   for c in range(1, ws.max_column + 1)
                                   if ws.cell(hdr, c).value is not None and str(ws.cell(hdr, c).value).strip().lower() in col_label}
                        maps = {c: {} for c in targets}
                        for r in range(hdr + 1, ws.max_row + 1):
                            for c, label in targets.items():
                                v = ws.cell(r, c).value
                                if v is None or str(v).strip() == "":
                                    continue
                                m = maps[c]; key = str(v)
                                if key not in m:
                                    m[key] = _pseudonym(label, len(m) + 1)
                                ws.cell(r, c).value = m[key]
            except Exception:
                pass
        if comments or unhide or redact_pii or rm_words_l:
            cells = 0
            for ws in wb.worksheets:
                try:
                    if unhide:
                        for d in ws.column_dimensions.values(): d.hidden = False
                        for d in ws.row_dimensions.values():    d.hidden = False
                    if comments or redact_pii or rm_words_l:
                        for row in ws.iter_rows():
                            for cell in row:
                                cells += 1
                                if comments and cell.comment:
                                    cell.comment = None
                                if isinstance(cell.value, str):
                                    if rm_words_l and any(w in cell.value.lower() for w in rm_words_l):
                                        cell.value = ""
                                    elif redact_pii:
                                        masked = _mask_pii(cell.value)
                                        if masked != cell.value:
                                            cell.value = masked
                            if cells > 100000: break
                        if cells > 100000: break
                except Exception:
                    continue
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    except Exception as e:
        raise HTTPException(500, f"Could not clean file: {str(e)}")
    base = (file.filename or "file").rsplit(".", 1)[0]
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{base}_safe.xlsx"'})
